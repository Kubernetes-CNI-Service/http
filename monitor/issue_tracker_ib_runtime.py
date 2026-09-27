"""Read-only IB report evidence for REQ7 Stage L.

The standalone validator report has no completed-cycle or write authority.
Only a protected worker attestation can later promote its bytes into a source.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import zipfile

from ztp.config.ib_topology_provenance import ProvenanceError, read_report_provenance


class IbRuntimeHoldError(ValueError):
    """The local IB report is absent, malformed, stale, or unbound."""


@dataclass(frozen=True)
class IbLocalReportWitness:
    report_path: str
    report_sha256: str
    expected_topology_sha256: str
    actual_path: str
    actual_sha256: str
    cvt_path: str
    profile_sha256: str
    qualified: bool = False


@dataclass(frozen=True)
class IbReportLink:
    issue_type: str
    expected_endpoints: tuple[str, str, str, str]
    actual_destination: tuple[str, str]
    source_sheet: str
    report_row: int


@dataclass(frozen=True)
class IbBoundCycleReport:
    """Literal rows from one protected completed cycle, not a qualified set."""

    cycle_id: str
    sequence: int
    completion_sha256: str
    report_sha256: str
    expected_topology_sha256: str
    rows: tuple[IbReportLink, ...]
    qualified: bool = False
    recorded_at_utc: str = ""


@dataclass(frozen=True)
class IbKWindow:
    """A source-bound IB K reduction, never a three-panel qualification."""

    status: str
    reason: str
    observed: int
    required: int
    cycle_ids: tuple[str, ...]
    report_sha256s: tuple[str, ...]
    expected_topology_sha256: str | None
    rows: tuple[IbReportLink, ...]
    qualified: bool = False
    recorded_at_utc: str = ""


@dataclass(frozen=True)
class IbWhitelistSkip:
    row: IbReportLink
    matched_rule: str
    snapshot_sha256: str


@dataclass(frozen=True)
class IbWhitelistedKWindow:
    """C-24 IB preview only; the held workbook is not a C5 authority."""

    window: IbKWindow
    workbook_sha256: str
    whitelist_sha256: str
    kept_rows: tuple[IbReportLink, ...]
    skips: tuple[IbWhitelistSkip, ...]
    qualified: bool = False


@dataclass(frozen=True)
class ProdIbKWindow:
    """Protected completed-cycle IB rows under one durable K and local W1.

    This is source evidence, not a three-panel QualifiedSet or write permit.
    ``recorded_at_utc`` comes only from the protected producer run path,
    never from a caller's cycle ID or from an ETH observation.
    """

    cycle_ids: tuple[str, ...]
    completion_sha256: tuple[str, ...]
    report_sha256: tuple[str, ...]
    expected_topology_sha256: str
    rows: tuple[IbReportLink, ...]
    whitelist_skips: tuple[tuple[tuple[str, str, str, str], str], ...]
    k_event_sha256: str
    whitelist_sha256: str
    template_sha256: str
    qualified: bool = False
    recorded_at_utc: str = ""


_REPORT_NAME = re.compile(r".+-topology-validation\.xlsx\Z", re.IGNORECASE)
_HEX_SHA256 = re.compile(r"[0-9a-f]{64}\Z")
_ROW_COLUMNS = ("SrcDevice", "SrcPort", "Expected_DstDevice", "Expected_DstPort")


def intersect_ib_report_rows(
    rows_by_cycle: tuple[tuple[IbReportLink, ...], ...],
    topology_sha256_by_cycle: tuple[str, ...],
) -> tuple[IbReportLink, ...]:
    """Reduce already verified K cycles without granting publication authority.

    The protected caller, not this reducer, must replay and bind every cycle.
    """
    if (type(rows_by_cycle) is not tuple or not rows_by_cycle
            or type(topology_sha256_by_cycle) is not tuple
            or len(rows_by_cycle) != len(topology_sha256_by_cycle)
            or any(type(digest) is not str or _HEX_SHA256.fullmatch(digest) is None
                   for digest in topology_sha256_by_cycle)
            or len(set(topology_sha256_by_cycle)) != 1):
        raise IbRuntimeHoldError("IB K report topology is absent or unstable")
    windows: list[dict[tuple[str, str, str, str], IbReportLink]] = []
    for rows in rows_by_cycle:
        if type(rows) is not tuple:
            raise IbRuntimeHoldError("IB K report rows are not a tuple")
        active: dict[tuple[str, str, str, str], IbReportLink] = {}
        for row in rows:
            if type(row) is not IbReportLink:
                raise IbRuntimeHoldError("IB K report row type is invalid")
            expected = row.expected_endpoints
            actual = row.actual_destination
            if (type(expected) is not tuple or len(expected) != 4
                    or any(type(part) is not str or not part or "\x00" in part
                           for part in expected)
                    or type(actual) is not tuple or len(actual) != 2
                    or type(row.report_row) is not int or row.report_row < 2):
                raise IbRuntimeHoldError("IB K report row values are invalid")
            if row.issue_type == "Link Down":
                if row.source_sheet != "Missing_Links" or actual != ("", ""):
                    raise IbRuntimeHoldError("IB Missing row invents an actual side")
            elif row.issue_type == "Mis-wiring":
                if (row.source_sheet != "Miswired_Links"
                        or any(type(part) is not str or not part or "\x00" in part
                               for part in actual)):
                    raise IbRuntimeHoldError("IB Miswired row has no actual side")
            else:
                raise IbRuntimeHoldError("IB K report issue type is unknown")
            key = ("ib_cabling", expected[0], expected[1], row.issue_type)
            if key in active:
                raise IbRuntimeHoldError("IB K report has ambiguous activity key")
            active[key] = row
        windows.append(active)
    intersection = set(windows[0])
    for active in windows[1:]:
        intersection.intersection_update(active)
    for key in intersection:
        if len({active[key].expected_endpoints for active in windows}) != 1:
            raise IbRuntimeHoldError("IB K report expected peer changed within topology")
    return tuple(windows[-1][key] for key in sorted(intersection))


def parse_ib_validation_rows(report_bytes: bytes) -> tuple[IbReportLink, ...]:
    """Decode literal Missing/Miswired values without conferring cycle authority."""
    if type(report_bytes) is not bytes or not report_bytes or len(report_bytes) > 64 * 1024 * 1024:
        raise IbRuntimeHoldError("IB validation report bytes are not bounded")
    try:
        from openpyxl import load_workbook

        with zipfile.ZipFile(BytesIO(report_bytes)) as archive:
            members = archive.infolist()
            if (len(members) > 2048
                    or len({item.filename for item in members}) != len(members)
                    or sum(item.file_size for item in members) > 128 * 1024 * 1024
                    or any(item.file_size > 64 * 1024 * 1024 for item in members)
                    or archive.testzip() is not None):
                raise IbRuntimeHoldError("IB validation report ZIP is unsafe")
        book = load_workbook(BytesIO(report_bytes), read_only=True,
                             data_only=False, keep_links=False)
        try:
            if not {"Summary", "Missing_Links", "Miswired_Links"}.issubset(book.sheetnames):
                raise IbRuntimeHoldError("IB validation report required sheets are missing")
            result: list[IbReportLink] = []
            seen: set[tuple[str, str, str, str]] = set()
            for sheet_name, issue_type in (("Missing_Links", "Link Down"),
                                           ("Miswired_Links", "Mis-wiring")):
                sheet = book[sheet_name]
                if sheet.max_row > 100_001 or sheet.max_column > 64:
                    raise IbRuntimeHoldError("IB validation report sheet is unbounded")
                iterator = sheet.iter_rows(values_only=False)
                header_cells = next(iterator, ())
                headers = tuple(cell.value for cell in header_cells)
                if not any(value is not None for value in headers):
                    if sheet.max_row > 1:
                        raise IbRuntimeHoldError("IB validation report has rows without headers")
                    continue
                while headers and headers[-1] is None:
                    headers = headers[:-1]
                if (not all(isinstance(value, str) and value for value in headers)
                        or len(headers) != len(set(headers))):
                    raise IbRuntimeHoldError("IB validation report headers are ambiguous")
                positions = {name: index for index, name in enumerate(headers)}
                needed = set(_ROW_COLUMNS)
                if sheet_name == "Miswired_Links":
                    needed.update(("Actual_DstDevice", "Actual_DstPort"))
                for row_number, cells in enumerate(iterator, start=2):
                    if not any(cell.value not in (None, "") for cell in cells):
                        continue
                    if not needed.issubset(positions):
                        raise IbRuntimeHoldError("IB validation report row lacks required columns")
                    values: dict[str, str] = {}
                    for name in needed:
                        cell = cells[positions[name]]
                        if (cell.data_type == "f" or not isinstance(cell.value, str)
                                or not cell.value or "\x00" in cell.value):
                            raise IbRuntimeHoldError("IB validation report row has unsafe value")
                        values[name] = cell.value
                    expected = tuple(values[name] for name in _ROW_COLUMNS)
                    key = ("ib_cabling", expected[0], expected[1], issue_type)
                    if key in seen:
                        raise IbRuntimeHoldError("IB validation report activity key repeats")
                    seen.add(key)
                    actual = ((values["Actual_DstDevice"], values["Actual_DstPort"])
                              if sheet_name == "Miswired_Links" else ("", ""))
                    result.append(IbReportLink(
                        issue_type, expected, actual, sheet_name, row_number,
                    ))
            return tuple(result)
        finally:
            book.close()
    except (OSError, TypeError, ValueError, zipfile.BadZipFile) as exc:
        if isinstance(exc, IbRuntimeHoldError):
            raise
        raise IbRuntimeHoldError("IB validation report rows are unsafe") from exc


def read_ib_local_report(report_path: Path) -> IbLocalReportWitness:
    """Verify producer sidecars and return a nonqualifying byte witness."""
    report = Path(report_path)
    if not _REPORT_NAME.fullmatch(report.name):
        raise IbRuntimeHoldError("IB report has no validation-report name")
    try:
        record = read_report_provenance(report)
    except (OSError, TypeError, ValueError, ProvenanceError) as exc:
        raise IbRuntimeHoldError("IB report provenance is missing, stale, or unsafe") from exc
    expected = record["expected_topology_sha256"]
    if not isinstance(expected, str) or _HEX_SHA256.fullmatch(expected) is None:
        raise IbRuntimeHoldError("IB expected-topology digest is invalid")
    return IbLocalReportWitness(
        report_path=record["report"]["path"],
        report_sha256=record["report"]["sha256"],
        expected_topology_sha256=expected,
        actual_path=record["actual"]["path"],
        actual_sha256=record["actual"]["sha256"],
        cvt_path=record["cvt"]["path"],
        profile_sha256=record["profile"]["sha256"],
    )


def read_completed_ib_cycle_report(
    store, *, http_root: Path | str, sequence: int,
) -> IbBoundCycleReport:
    """Reopen an independently protected cycle's literal IB report bytes.

    The cycle inspector must replay completion and root-only attestation on
    both sides of the read. A local report witness alone is never sufficient.
    """
    from monitor.collection_ib_cycle_binding import (
        IbCycleRoleSource, inspect_completed_ib_cycle_roles,
    )
    from monitor.issue_tracker_switch_runtime import _read_bound, _safe_child

    try:
        root = Path(http_root).resolve(strict=True)
        source = inspect_completed_ib_cycle_roles(
            store, http_root=root, sequence=sequence,
        )
        if type(source) is not IbCycleRoleSource or source.qualified is not False:
            raise IbRuntimeHoldError("IB completed cycle has no protected report role")
        report = _safe_child(root, source.report_relative_path)
        raw = _read_bound(root, report, maximum=64 * 1024 * 1024)
        if hashlib.sha256(raw).hexdigest() != source.report_sha256:
            raise IbRuntimeHoldError("IB completed report bytes differ from role")
        witness = read_ib_local_report(report)
        if (Path(witness.report_path).resolve(strict=True) != report
                or witness.report_sha256 != source.report_sha256
                or witness.expected_topology_sha256
                != source.expected_topology_sha256):
            raise IbRuntimeHoldError("IB report provenance differs from completed role")
        rows = parse_ib_validation_rows(raw)
        parts = Path(source.report_relative_path).parts
        recorded_at_utc = ""
        project_parts = (Path(store.project_identity).relative_to(root).parts
                         if hasattr(store, "project_identity") else ())
        if (len(parts) >= 7 and parts[:2] == project_parts
                and parts[2:4] == ("99-output-ufm", "runs")):
            run_id = parts[4]
            if re.fullmatch(r"[0-9]{8}-[0-9]{4}-prod-[0-9a-f]{16}", run_id) is None:
                raise IbRuntimeHoldError("IB protected producer run time is malformed")
            recorded_at_utc = datetime.strptime(
                run_id[:13], "%Y%m%d-%H%M",
            ).strftime("%Y-%m-%dT%H:%M:00Z")
        if (inspect_completed_ib_cycle_roles(
                store, http_root=root, sequence=sequence,
            ) != source or _read_bound(root, report, maximum=64 * 1024 * 1024) != raw):
            raise IbRuntimeHoldError("IB completed source changed during report read")
        return IbBoundCycleReport(
            cycle_id=source.cycle_id, sequence=source.sequence,
            completion_sha256=source.completion_sha256,
            report_sha256=source.report_sha256,
            expected_topology_sha256=source.expected_topology_sha256,
            rows=rows, recorded_at_utc=recorded_at_utc,
        )
    except (OSError, TypeError, AttributeError, RuntimeError, ValueError) as exc:
        if isinstance(exc, IbRuntimeHoldError):
            raise
        raise IbRuntimeHoldError("IB completed report source is unsafe") from exc


def read_ib_k_window(store, *, http_root: Path | str, token) -> IbKWindow:
    """Reduce consecutive protected IB cycles under the actual local K owner.

    This is intentionally read-only. Even a history-sufficient IB window is
    not a QualifiedSet and cannot authorize C5, C6, or online publication.
    """
    from monitor.issue_tracker_cycle_source import read_completed_cycle_evidence
    from monitor.issue_tracker_k_chain import assess_k_window
    from monitor.issue_tracker_local_commit import _owner
    from monitor.issue_tracker_state_owner import current_k

    try:
        root = Path(http_root).resolve(strict=True)
        owner = _owner(token)
        if (not root.is_dir() or Path(store.root_status_dir) != root / "monitor/status"
                or store.scope != "prod"
                or Path(store.project_identity).resolve(strict=True) != owner.project):
            raise IbRuntimeHoldError("IB K owner, store, and HTTP root differ")
        settings = current_k(token)
        cycles = read_completed_cycle_evidence(store)
        window = assess_k_window(cycles, settings)
        if window.status != "history_sufficient":
            if current_k(token) != settings or read_completed_cycle_evidence(store) != cycles:
                raise IbRuntimeHoldError("IB K cold-start source changed")
            return IbKWindow(
                window.status, window.reason, window.observed, window.required,
                (), (), None, (),
            )
        tail = cycles[-settings.k:]
        reports = tuple(read_completed_ib_cycle_report(
            store, http_root=root, sequence=cycle.sequence,
        ) for cycle in tail)
        if any(
            report.cycle_id != cycle.cycle_id
            or report.sequence != cycle.sequence
            or report.completion_sha256 != cycle.completion_sha256
            or report.qualified is not False
            for cycle, report in zip(tail, reports)
        ):
            raise IbRuntimeHoldError("IB K report is not the selected completed cycle")
        rows = intersect_ib_report_rows(
            tuple(report.rows for report in reports),
            tuple(report.expected_topology_sha256 for report in reports),
        )
        if current_k(token) != settings or read_completed_cycle_evidence(store) != cycles:
            raise IbRuntimeHoldError("IB K source changed during reduction")
        return IbKWindow(
            window.status, window.reason, window.observed, window.required,
            tuple(cycle.cycle_id for cycle in tail),
            tuple(report.report_sha256 for report in reports),
            reports[-1].expected_topology_sha256, rows,
            recorded_at_utc=reports[-1].recorded_at_utc,
        )
    except (OSError, TypeError, AttributeError, RuntimeError, ValueError) as exc:
        if isinstance(exc, IbRuntimeHoldError):
            raise
        raise IbRuntimeHoldError("IB K owner or completed-cycle history is unsafe") from exc


def read_whitelisted_ib_k_window(
    store, *, http_root: Path | str, token, workbook_path: Path | str,
) -> IbWhitelistedKWindow:
    """Apply W1 to actual expected A/Z nodes and recheck both read sources.

    This remains an IB-only, nonqualifying preview; C5 must separately bind
    the selected workbook to the actual local transaction and three panels.
    """
    from monitor.issue_tracker_whitelist import evaluate_whitelist
    from monitor.issue_tracker_whitelist_workbook import read_whitelist_workbook

    try:
        snapshot = read_whitelist_workbook(workbook_path)
        window = read_ib_k_window(store, http_root=http_root, token=token)
        records = tuple({
            "record_id": "ib:" + hashlib.sha256(json.dumps(
                ("ib_cabling", row.expected_endpoints[0],
                 row.expected_endpoints[1], row.issue_type),
                ensure_ascii=False, separators=(",", ":"),
            ).encode("utf-8")).hexdigest(),
            "source": "Cabling",
            "a_node": row.expected_endpoints[0],
            "z_node": row.expected_endpoints[2],
        } for row in window.rows)
        decisions = evaluate_whitelist(snapshot.whitelist, records).decisions
        kept = tuple(row for row, decision in zip(window.rows, decisions)
                     if not decision.skipped)
        skips = tuple(IbWhitelistSkip(row, decision.matched_rule,
                                      decision.snapshot_sha256)
                      for row, decision in zip(window.rows, decisions)
                      if decision.skipped)
        if (read_whitelist_workbook(workbook_path) != snapshot
                or read_ib_k_window(store, http_root=http_root, token=token) != window):
            raise IbRuntimeHoldError("IB K or local Whitelist changed during preview")
        return IbWhitelistedKWindow(
            window, snapshot.workbook_sha256, snapshot.whitelist.sha256,
            kept, skips,
        )
    except (OSError, TypeError, AttributeError, RuntimeError, ValueError) as exc:
        if isinstance(exc, IbRuntimeHoldError):
            raise
        raise IbRuntimeHoldError("IB K Whitelist source is unsafe") from exc


def read_prod_ib_k_window(
    store, *, http_root: Path | str, settings,
    whitelist_snapshot, whitelist_path: Path | str, writer_token=None,
) -> ProdIbKWindow:
    """Replay genuine protected nine-role reports, durable K and W1 together.

    Neither a standalone report nor a caller-supplied role inventory reaches
    this path: ``read_ib_k_window`` opens the completed-cycle inspector, which
    reopens the fixed protected producer attestation for every selected cycle.
    """
    from monitor.issue_tracker_cycle_source import read_completed_cycle_evidence
    from monitor.issue_tracker_k_chain import KChainState, assess_k_window
    from monitor.issue_tracker_local_commit import _owner
    from monitor.issue_tracker_state_owner import current_k, tracker_writer
    from monitor.issue_tracker_whitelist import evaluate_whitelist
    from monitor.issue_tracker_whitelist_workbook import (
        WorkbookWhitelistSnapshot, read_whitelist_workbook,
    )

    try:
        root = Path(http_root).resolve(strict=True)
        project = Path(store.project_identity)
        publication = project / "99-output-monitor"
        if (not root.is_dir() or store.scope != "prod"
                or Path(store.root_status_dir) != root / "monitor/status"
                or not project.is_absolute() or project.is_symlink()
                or type(settings) is not KChainState
                or type(whitelist_snapshot) is not WorkbookWhitelistSnapshot):
            raise IbRuntimeHoldError("IB K adapter authority is unsupported")

        def observe(token):
            owner = _owner(token)
            if (owner.project != project or owner.publication != publication
                    or current_k(token) != settings
                    or read_whitelist_workbook(whitelist_path) != whitelist_snapshot):
                raise IbRuntimeHoldError("IB K or Whitelist owner changed")
            cycles = read_completed_cycle_evidence(store)
            if assess_k_window(cycles, settings).status != "history_sufficient":
                raise IbRuntimeHoldError("IB protected K history is incomplete")
            tail = cycles[-settings.k:]
            window = read_ib_k_window(store, http_root=root, token=token)
            if (window.status != "history_sufficient"
                    or window.qualified is not False
                    or window.cycle_ids != tuple(cycle.cycle_id for cycle in tail)
                    or window.expected_topology_sha256 is None
                    or not window.recorded_at_utc
                    or len(window.report_sha256s) != settings.k):
                raise IbRuntimeHoldError("IB protected K reports do not bind the selected cycles")
            records = tuple({
                "record_id": "ib:" + hashlib.sha256(json.dumps(
                    ("ib_cabling", row.expected_endpoints[0],
                     row.expected_endpoints[1], row.issue_type),
                    ensure_ascii=False, separators=(",", ":"),
                ).encode("utf-8")).hexdigest(),
                "source": "Cabling", "a_node": row.expected_endpoints[0],
                "z_node": row.expected_endpoints[2],
            } for row in window.rows)
            decisions = evaluate_whitelist(
                whitelist_snapshot.whitelist, records,
            ).decisions
            kept = tuple(row for row, decision in zip(window.rows, decisions)
                         if not decision.skipped)
            skips = tuple((
                ("ib_cabling", row.expected_endpoints[0],
                 row.expected_endpoints[1], row.issue_type),
                decision.matched_rule,
            ) for row, decision in zip(window.rows, decisions) if decision.skipped)
            if (current_k(token) != settings
                    or read_whitelist_workbook(whitelist_path) != whitelist_snapshot
                    or read_completed_cycle_evidence(store) != cycles
                    or read_ib_k_window(store, http_root=root, token=token) != window):
                raise IbRuntimeHoldError("IB protected K source changed during W1 replay")
            return ProdIbKWindow(
                cycle_ids=window.cycle_ids,
                completion_sha256=tuple(cycle.completion_sha256 for cycle in tail),
                report_sha256=window.report_sha256s,
                expected_topology_sha256=window.expected_topology_sha256,
                rows=kept, whitelist_skips=skips,
                k_event_sha256=settings.event_sha256,
                whitelist_sha256=whitelist_snapshot.whitelist.sha256,
                template_sha256=whitelist_snapshot.workbook_sha256,
                recorded_at_utc=window.recorded_at_utc,
            )

        if writer_token is not None:
            return observe(writer_token)
        with tracker_writer(project, publication) as token:
            return observe(token)
    except (OSError, TypeError, KeyError, AttributeError, RuntimeError,
            ValueError) as exc:
        if isinstance(exc, IbRuntimeHoldError):
            raise
        raise IbRuntimeHoldError("IB protected K/W1 source cannot be replayed") from exc


def validate_prod_ib_k_window(
    witness: ProdIbKWindow, store, *, http_root: Path | str, settings,
    whitelist_snapshot, whitelist_path: Path | str, writer_token=None,
) -> ProdIbKWindow:
    """Reopen all protected bytes; a copied source witness is not a permit."""
    if type(witness) is not ProdIbKWindow or witness.qualified is not False:
        raise IbRuntimeHoldError("IB K witness claims unearned authority")
    actual = read_prod_ib_k_window(
        store, http_root=http_root, settings=settings,
        whitelist_snapshot=whitelist_snapshot, whitelist_path=whitelist_path,
        writer_token=writer_token,
    )
    if actual != witness:
        raise IbRuntimeHoldError("IB K/W1 source changed before freeze")
    return actual


def require_completed_ib_cycle(witness: IbLocalReportWitness, cycle: object) -> None:
    """Never turn a caller-supplied cycle claim into protected worker authority."""
    if not isinstance(witness, IbLocalReportWitness) or witness.qualified:
        raise IbRuntimeHoldError("IB local witness is invalid")
    slots = cycle.get("source_slots") if isinstance(cycle, dict) else None
    if (not isinstance(cycle, dict) or cycle.get("scope") != "prod"
            or cycle.get("qualifying") is not True
            or not isinstance(slots, (tuple, list))
            or "infiniband/prod" not in slots):
        raise IbRuntimeHoldError("IB completed prod cycle is not established")
    raise IbRuntimeHoldError("completed collector has no IB analyzer artifact role")
