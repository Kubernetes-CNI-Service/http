#!/usr/bin/env python3
"""Independent, local P2P workbook -> LLDPQ derivation witness.

This module proves a *frozen local input* relation. It cannot establish that a
worker issued a cycle or qualify an issue-tracker operation. The caller must
first bind these paths to its own durable worker start/launch/completion chain.
"""

from __future__ import annotations

from dataclasses import dataclass
import fnmatch
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import stat
from typing import Mapping
import zipfile

from openpyxl import load_workbook


ROLES = frozenset({"topology", "inventory", "port_mapping", "lldpq", "splitter_profiles"})
LIMITS = {"topology": 64 * 1024 * 1024, "inventory": 4 * 1024 * 1024,
          "port_mapping": 4 * 1024 * 1024, "lldpq": 16 * 1024 * 1024,
          "splitter_profiles": 4 * 1024 * 1024}
SHA = re.compile(r"[0-9a-f]{64}\Z")
DOT_EDGE = re.compile(
    r'^\s*"([^"]+)"\s*:\s*"([^"]+)"\s*--\s*'
    r'"([^"]+)"\s*:\s*"([^"]+)"\s*;?\s*$'
)
PLACEHOLDERS = {"", "-", "na", "n/a", "#n/a", "none", "null", "empty"}


class TopologyDerivationHoldError(ValueError):
    """One claimed derivation input is absent, changed, ambiguous, or inconsistent."""


@dataclass(frozen=True)
class DerivationEdge:
    sheet: str
    row: int
    a_node: str
    a_port: str
    z_node: str
    z_port: str


@dataclass(frozen=True)
class TopologyDerivationWitness:
    source_sha256: Mapping[str, str]
    edges: tuple[DerivationEdge, ...]

    @property
    def qualified(self) -> bool:
        return False


def _hold(message: str) -> TopologyDerivationHoldError:
    return TopologyDerivationHoldError(message)


def _identity(value: os.stat_result) -> tuple:
    return (value.st_dev, value.st_ino, value.st_mode, value.st_nlink,
            value.st_size, value.st_mtime_ns, value.st_ctime_ns)


def _held_bytes(path: Path, role: str, digest: str) -> bytes:
    if not isinstance(digest, str) or not SHA.fullmatch(digest):
        raise _hold(f"invalid {role} expected SHA256")
    descriptor = -1
    try:
        named = path.lstat()
        if not stat.S_ISREG(named.st_mode) or named.st_nlink != 1 \
                or named.st_size < 0 or named.st_size > LIMITS[role]:
            raise _hold(f"{role} must be a bounded single-link regular file")
        descriptor = os.open(path, os.O_RDONLY | getattr(os, "O_CLOEXEC", 0)
                             | getattr(os, "O_NOFOLLOW", 0))
        before = os.fstat(descriptor)
        if _identity(before) != _identity(named):
            raise _hold(f"{role} identity changed before read")
        remaining = before.st_size
        chunks = []
        while remaining:
            part = os.read(descriptor, min(remaining, 1024 * 1024))
            if not part:
                raise _hold(f"{role} shortened during read")
            chunks.append(part)
            remaining -= len(part)
        if os.read(descriptor, 1):
            raise _hold(f"{role} grew during read")
        after = os.fstat(descriptor)
        named_after = path.lstat()
        raw = b"".join(chunks)
        if _identity(before) != _identity(after) \
                or _identity(after) != _identity(named_after) \
                or hashlib.sha256(raw).hexdigest() != digest:
            raise _hold(f"{role} frozen source changed")
        return raw
    except OSError as exc:
        raise _hold(f"{role} cannot be read safely: {exc}") from exc
    finally:
        if descriptor >= 0:
            os.close(descriptor)


def _unique_pairs(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise _hold(f"duplicate sidecar key: {key}")
        result[key] = value
    return result


def _reject_constant(value):
    raise _hold(f"invalid sidecar constant: {value}")


def _text(raw: bytes, label: str) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeError as exc:
        raise _hold(f"{label} is not UTF-8") from exc


def _field(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _empty(value: str) -> bool:
    return value.casefold() in PLACEHOLDERS


def _columns(first: tuple, second: tuple, *, legacy: bool) -> tuple[int, ...]:
    a = [_field(cell).casefold() for cell in first]
    b = [_field(cell).casefold() for cell in second]
    start = next((i for i, value in enumerate(a)
                  if "source" in value or "src" in value), None)
    end = next((i for i, value in enumerate(a)
                if "dest" in value or "dst" in value), None)
    if start is not None and end is not None and start < end:
        def find(lo, hi, names):
            return next((i for i in range(lo, min(hi, len(b)))
                         if any(name == b[i] or name in b[i] for name in names)), None)
        cols = (find(start, end, ("name",)), find(start, end, ("port",)),
                find(end, len(b), ("name",)), find(end, len(b), ("port",)))
        if all(value is not None for value in cols):
            return cols
    if legacy:
        return (6, 7, 12, 13)
    raise _hold("P2P TAN/OOB header is not independently recognizable")


def _workbook_rows(raw: bytes, *, legacy: bool) -> tuple[DerivationEdge, ...]:
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 2048 or sum(item.file_size for item in members) > 256 * 1024 * 1024:
                raise _hold("P2P workbook expanded size is not bounded")
            if any(item.file_size > 128 * 1024 * 1024 for item in members):
                raise _hold("P2P workbook has an oversized member")
            names = [item.filename.casefold() for item in members]
            if len(names) != len(set(names)):
                raise _hold("P2P workbook has duplicate ZIP members")
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            sheets = [name for name in book.sheetnames
                      if name.upper().startswith(("TAN", "OOB"))]
            if not sheets:
                raise _hold("P2P workbook has no TAN/OOB source sheet")
            rows = []
            for name in sheets:
                iterator = book[name].iter_rows(values_only=True)
                first, second = next(iterator), next(iterator)
                cols = _columns(first, second, legacy=legacy)
                for number, row in enumerate(iterator, 3):
                    if all(value is None for value in row) or (len(row) > 1 and row[1] == "#N/A"):
                        continue
                    fields = tuple(_field(row[index] if index < len(row) else None)
                                   for index in cols)
                    a_empty = _empty(fields[0]) and _empty(fields[1])
                    z_empty = _empty(fields[2]) and _empty(fields[3])
                    if _empty(fields[0]) != _empty(fields[1]) \
                            or _empty(fields[2]) != _empty(fields[3]):
                        raise _hold(f"P2P {name}:{number} has a partial endpoint")
                    if a_empty and z_empty:
                        continue
                    rows.append(DerivationEdge(name, number, *fields))
            if not rows:
                raise _hold("P2P workbook has no source rows")
            return tuple(rows)
        finally:
            book.close()
    except TopologyDerivationHoldError:
        raise
    except (OSError, ValueError, KeyError, StopIteration, zipfile.BadZipFile) as exc:
        raise _hold(f"P2P workbook cannot be parsed: {exc}") from exc


def _inventory(raw: bytes) -> tuple[tuple[str, tuple[str, ...]], ...]:
    sections = []
    current = None
    metadata = False
    seen = set()
    for number, raw_line in enumerate(_text(raw, "inventory").splitlines(), 1):
        line = raw_line.strip()
        if not line:
            current = None
            metadata = False
        elif line.startswith("#"):
            continue
        elif line.startswith("[[") and line.endswith("]]"):
            current = None
            metadata = True
        elif line.startswith("[") and line.endswith("]"):
            if metadata:
                continue
            name = line[1:-1].strip()
            if not name or name.casefold() in seen:
                raise _hold(f"inventory section {number} is invalid or duplicated")
            seen.add(name.casefold())
            current = []
            sections.append((name, current))
        elif not metadata:
            if current is None:
                raise _hold(f"inventory pattern {number} is orphaned")
            current.append(line)
    if not sections:
        raise _hold("inventory has no sections")
    return tuple((name, tuple(patterns)) for name, patterns in sections)


def _device_type(name: str, sections) -> str:
    for section, patterns in sections:
        if any(fnmatch.fnmatchcase(name.casefold(), pattern.casefold())
               for pattern in patterns):
            return section
    return "unknown"


def _port_map(raw: bytes):
    direct = {}
    switch = {}
    section = None
    for number, raw_line in enumerate(_text(raw, "port mapping").splitlines(), 1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("[") and line.endswith("]"):
            section = line[1:-1]
            continue
        if section is None or "," not in line:
            continue
        pattern, mapped = (part.strip() for part in line.split(",", 1))
        if not pattern or not mapped:
            raise _hold(f"port mapping {number} is empty")
        key = f"{section}:{pattern}"
        target = switch if "#" in pattern else direct
        if key in target:
            raise _hold(f"duplicate port mapping {number}")
        target[key] = mapped
    return direct, switch


def _profiles(rows, sections, direct, switch):
    switch_types = {key.split(":", 1)[0] for key in switch}
    grouped = {}
    unsplit = set()
    for row in rows:
        for node, port in ((row.a_node, row.a_port), (row.z_node, row.z_port)):
            if _empty(node):
                continue
            kind = _device_type(node, sections)
            if kind not in switch_types or f"{kind}:{port}" in direct:
                continue
            base, *suffix = port.split("/")
            key = (node.casefold(), base)
            if suffix:
                grouped.setdefault(key, set()).add(tuple(suffix))
            else:
                unsplit.add(key)
    result = {}
    for key, suffixes in grouped.items():
        if key in unsplit:
            raise _hold("P2P has contradictory split and unsplit port")
        depths = {len(suffix) for suffix in suffixes}
        if len(depths) != 1 or next(iter(depths)) not in {1, 2}:
            raise _hold("P2P breakout depth is ambiguous")
        try:
            numeric = [tuple(int(part) for part in suffix) for suffix in suffixes]
        except ValueError as exc:
            raise _hold("P2P breakout suffix is not numeric") from exc
        if next(iter(depths)) == 1:
            if any(not 1 <= value[0] <= 2 for value in numeric):
                raise _hold("P2P breakout lane is invalid")
            mode = "1to2"
        else:
            if any(not (1 <= value[0] <= 2 and 1 <= value[1] <= 4)
                   for value in numeric):
                raise _hold("P2P breakout lane is invalid")
            mode = "1to8" if any(value[1] > 2 for value in numeric) else "1to4"
        result[key] = mode
    return result


def _resolve(node: str, port: str, sections, direct, switch, profiles) -> str:
    kind = _device_type(node, sections)
    if f"{kind}:{port}" in direct:
        return direct[f"{kind}:{port}"]
    base, separator, suffix = port.partition("/")
    mode = profiles.get((node.casefold(), base))
    if separator and mode is None:
        return port
    pattern = f"{mode or '1to1'}#{separator}{suffix}"
    return switch.get(f"{kind}:{pattern}", port).replace("#", base)


def _edge_key(values: tuple[str, str, str, str]) -> frozenset[tuple[str, str]]:
    a, ap, z, zp = values
    endpoints = frozenset(((a.casefold(), ap.casefold()),
                           (z.casefold(), zp.casefold())))
    if len(endpoints) != 2:
        raise _hold("P2P or DOT has a degenerate edge")
    return endpoints


def _dot_edges(raw: bytes) -> set[frozenset[tuple[str, str]]]:
    edges = set()
    for number, line in enumerate(_text(raw, "LLDPQ DOT").splitlines(), 1):
        match = DOT_EDGE.fullmatch(line)
        if match is None:
            if "--" in line:
                raise _hold(f"LLDPQ DOT has malformed edge at line {number}")
            continue
        key = _edge_key(match.groups())
        if key in edges:
            raise _hold("LLDPQ DOT has duplicate edge")
        edges.add(key)
    if not edges:
        raise _hold("LLDPQ DOT has no physical edges")
    return edges


def verify_lldpq_derivation(*, sources: Mapping[str, Path],
                            expected_sha256: Mapping[str, str],
                            selected_workbook_name: str,
                            legacy_columns: bool = False) -> TopologyDerivationWitness:
    """Recompute physical edges from held bytes and compare the full DOT set.

    The sidecar is checked for agreement, never treated as an edge oracle.
    Neither a caller-provided digest nor this return value proves worker origin.
    """
    if not isinstance(sources, Mapping) or set(sources) != ROLES \
            or not isinstance(expected_sha256, Mapping) \
            or set(expected_sha256) != ROLES:
        raise _hold("P2P derivation source roles are incomplete")
    if not isinstance(selected_workbook_name, str) \
            or not selected_workbook_name.lower().endswith(".xlsx") \
            or Path(selected_workbook_name).name != selected_workbook_name \
            or selected_workbook_name in {".", ".."}:
        raise _hold("selected workbook basename is invalid")
    if type(legacy_columns) is not bool:
        raise _hold("legacy-columns mode must be explicit boolean")
    paths = {}
    raw = {}
    for role in sorted(ROLES):
        path = sources[role]
        if not isinstance(path, Path):
            raise _hold(f"{role} must be an exact Path")
        paths[role] = path
        raw[role] = _held_bytes(path, role, expected_sha256[role])
    # Worker's frozen topology name can be slot-scoped. The selected basename
    # is instead authenticated by the worker setup selection and sidecar.
    rows = _workbook_rows(raw["topology"], legacy=legacy_columns)
    sections = _inventory(raw["inventory"])
    direct, switch = _port_map(raw["port_mapping"])
    profiles = _profiles(rows, sections, direct, switch)
    try:
        sidecar = json.loads(_text(raw["splitter_profiles"], "splitter sidecar"),
                             object_pairs_hook=_unique_pairs,
                             parse_constant=_reject_constant)
    except (TypeError, ValueError) as exc:
        if isinstance(exc, TopologyDerivationHoldError):
            raise
        raise _hold(f"splitter sidecar is malformed: {exc}") from exc
    if not isinstance(sidecar, dict) or set(sidecar) != {
        "schema_version", "source_workbook", "workbook_sha256",
        "inventory_sha256", "port_mapping_sha256", "lldpq_sha256", "profiles"
    } or type(sidecar["schema_version"]) is not int \
            or sidecar["schema_version"] != 1 \
            or sidecar["source_workbook"] != selected_workbook_name:
        raise _hold("splitter sidecar selected workbook identity is invalid")
    for sidecar_key, role in (("workbook_sha256", "topology"),
                              ("inventory_sha256", "inventory"),
                              ("port_mapping_sha256", "port_mapping"),
                              ("lldpq_sha256", "lldpq")):
        if sidecar[sidecar_key] != expected_sha256[role]:
            raise _hold(f"splitter sidecar {sidecar_key} disagrees with frozen source")
    if not isinstance(sidecar["profiles"], list):
        raise _hold("splitter profile list is invalid")
    observed_profiles = set()
    for entry in sidecar["profiles"]:
        if not isinstance(entry, dict) or set(entry) != {"device", "parent", "profile"} \
                or not all(isinstance(value, str) for value in entry.values()):
            raise _hold("splitter profile entry is malformed")
        key = (entry["device"].casefold(), entry["parent"].removeprefix("swp"))
        if entry["parent"] != "swp" + key[1] or key in observed_profiles \
                or profiles.get(key) != entry["profile"]:
            raise _hold("splitter sidecar profile was not independently inferred")
        observed_profiles.add(key)
    if observed_profiles != set(profiles):
        raise _hold("splitter sidecar omits inferred profile")
    physical = {}
    for row in rows:
        if _empty(row.a_node) or _empty(row.z_node):
            continue
        resolved = (row.a_node, _resolve(row.a_node, row.a_port, sections, direct, switch, profiles),
                    row.z_node, _resolve(row.z_node, row.z_port, sections, direct, switch, profiles))
        key = _edge_key(resolved)
        if key in physical:
            raise _hold("P2P physical edge has duplicate workbook row")
        physical[key] = row
    if not physical or set(physical) != _dot_edges(raw["lldpq"]):
        raise _hold("LLDPQ physical edge set disagrees with P2P derivation")
    return TopologyDerivationWitness(
        {role: expected_sha256[role] for role in sorted(ROLES)},
        tuple(physical.values()),
    )
