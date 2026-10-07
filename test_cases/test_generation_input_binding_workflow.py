#!/usr/bin/env python3
"""Real setup/load/c2 handoff when a project input changes after validation."""

from __future__ import annotations

import base64
import copy
import csv
from contextlib import redirect_stdout
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import signal
import stat
import struct
import subprocess
import sys
import tempfile
import time
from types import SimpleNamespace
import unittest
from unittest import mock

import openpyxl
import yaml

from test_cases.module_loader import load_script
from test_cases.test_v2_generation_flow import (
    BASE_HEADER,
    EVPN_HEADER,
    FIXED_HEADER,
    V2MlagSetupPreflightTests,
    v2_mlag_preflight_global,
)
from test_cases.test_mlag_evpn_generation import border_globals
from test_cases import test_flow_release_platform_matrix as PLATFORM_FLOW
from test_cases import test_load_release_transaction as RELEASE_FIXTURE


ROOT = Path(__file__).resolve().parents[1]
SETUP = load_script(
    "r22_generation_binding_workflow_setup", ROOT / "DAY0-Prepare/01-a-setup.py",
)
LOAD = load_script(
    "r22_generation_binding_workflow_load", ROOT / "DAY0-Prepare/11-load.py",
)
C2 = load_script(
    "r22_generation_binding_workflow_c2",
    ROOT / "ztp/config/cumulus/template/90-c2-generate_configs.py",
)
ACTIVATE = load_script(
    "r22_generation_binding_workflow_activate", ROOT / "infra/docker/activate.py",
)
PUBLISHER = load_script(
    "r22_generation_binding_workflow_publisher",
    ROOT / "ztp/config/cumulus/d-hostname2mac.py",
)
SUBNET_HEADER = (
    "shared_network,subnet,netmask,range_start,range_end,routers,"
    "ztp_service_ip,cumulus_profile,nvos_ztp\n"
)

REAL_CHILD_SCRIPTS = (
    "ztp/config/cumulus/template/P2P/b-xlsx_to_dot.py",
    "ztp/config/isc-dhcp-server/c1-generate_dhcp.py",
    "ztp/config/cumulus/template/90-c2-generate_configs.py",
    "ztp/config/cumulus/d-hostname2mac.py",
)
DHCP_OUTPUTS = (
    "dhcpd.conf", "dhcpd_eth.hosts", "dhcpd_ib.hosts",
    "dhcpd_nvl.hosts", "dhcp-release-manifest.json",
)


def _write_legal_project(project: Path, rows: list[list[str]]) -> None:
    project.mkdir()
    document = v2_mlag_preflight_global()
    document["common"]["mgmt"].update({
        "dhcp-server": {"status": "disabled", "package": "isc-dhcp-server"},
        "http": {
            "status": "disabled", "http_root": "/var/www/html",
            "package": "apache2", "port": 80,
        },
        "ztp": {"status": "disabled", "ztp_url_prefix": "/ztp"},
    })
    document["common"]["switch"] = {
        "system": {"dns": {}, "ntp": {}, "date-time": {}},
    }
    document["switches"][0]["eth"]["system"]["aaa"] = {
        "user": {"cumulus": {"hashed-password": "$6$fixture$safe-value"}},
    }
    (project / "01-global.yaml").write_text(
        yaml.safe_dump(document, sort_keys=False), encoding="utf-8",
    )
    with (project / "02-devices_config.csv").open(
        "w", newline="", encoding="utf-8",
    ) as stream:
        csv.writer(stream).writerows([BASE_HEADER + FIXED_HEADER + EVPN_HEADER, *rows])
    (project / "02-dhcp-subnet_config.csv").write_text(
        SUBNET_HEADER, encoding="utf-8",
    )
    workbook = openpyxl.Workbook()
    try:
        workbook.save(project / "p2p.xlsx")
    finally:
        workbook.close()
    # A literal, structurally valid OpenSSH public key keeps LOAD's real
    # workstation key preflight on the path without touching the real HOME.
    key_type = b"ssh-ed25519"
    wire = (
        struct.pack(">I", len(key_type)) + key_type
        + struct.pack(">I", 32) + b"\x11" * 32
    )
    (project / "operator.pub").write_text(
        "ssh-ed25519 " + base64.b64encode(wire).decode("ascii") + " fixture\n",
        encoding="utf-8",
    )
    (project / ".management-pubkeys").write_text(
        "mgmt-server.pub\n", encoding="utf-8",
    )
    (project / "mgmt-server.pub").write_bytes(b"")


def _private_real_chain(
    root: Path, dns_ip: str, *, http_enabled: bool = False,
) -> tuple[Path, Path]:
    """Build only the files real LOAD/b/c1/c2/d read under a private root."""
    project = root / "DAY0-Prepare/project"
    project.parent.mkdir(parents=True)
    _write_legal_project(project, V2MlagSetupPreflightTests().pair())
    global_file = project / "01-global.yaml"
    document = yaml.safe_load(global_file.read_text(encoding="utf-8"))
    document["common"]["mgmt"]["dhcp-server"]["status"] = "enabled"
    document["common"]["mgmt"]["ztp"]["status"] = "enabled"
    if http_enabled:
        document["common"]["mgmt"]["http"]["status"] = "enabled"
    complete = border_globals()
    eth = document["switches"][0]["eth"]
    eth["bridge"] = complete["bridge"]
    eth["version"] = complete["version"]
    eth["vrf"] = complete["vrf"]
    eth["mlag"].update(complete["mlag"])
    eth["mlag"]["priority"] = [100, 200]
    eth["system"]["date-time"] = complete["system"]["date-time"]
    eth["system"]["ntp"] = complete["system"]["ntp"]
    eth["system"]["dns"] = {"server": [dns_ip], "vrf": "mgmt"}
    eth["system"]["aaa"]["user"]["cumulus"]["full-name"] = "cumulus,,,"
    global_file.write_text(yaml.safe_dump(document, sort_keys=False), encoding="utf-8")
    (project / "02-dhcp-subnet_config.csv").write_text(
        SUBNET_HEADER
        + "mgmt,192.0.2.0,255.255.255.0,192.0.2.100,192.0.2.120,"
        "192.0.2.1,192.0.2.20,oob,no\n",
        encoding="utf-8",
    )
    workbook = openpyxl.Workbook()
    try:
        sheet = workbook.active
        sheet.title = "OOB Server-LF"
        sheet.cell(1, 5, "Source")
        sheet.cell(1, 11, "Dest")
        for col, value in ((7, "name"), (8, "HCA/port"), (13, "name"), (14, "port")):
            sheet.cell(2, col, value)
        sheet.cell(3, 7, "Border01")
        sheet.cell(3, 8, "swp53")
        sheet.cell(3, 13, "Border02")
        sheet.cell(3, 14, "swp53")
        workbook.save(project / "selected-p2p.xlsx")
    finally:
        workbook.close()
    (project / "p2p.xlsx").unlink()
    (project / "p2p.xlsx").symlink_to("selected-p2p.xlsx")

    ztp = root / "ztp"
    shutil.copytree(ROOT / "tools", root / "tools")
    resources = (
        "ztp/nvue_normalizer.py",
        "ztp/config/topology_rules.py",
        "ztp/config/cumulus/template/P2P/lldpq-template.dot",
        "ztp/config/cumulus/template/P2P/air-template-no-oob.json",
        "ztp/config/cumulus/default.yaml",
        "ztp/optimize/sample_links.py",
        *REAL_CHILD_SCRIPTS,
    )
    for resource in resources:
        destination = root / resource
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(ROOT / resource, destination)
    shutil.copytree(
        ROOT / "ztp/config/cumulus/template/03-templates-j2",
        ztp / "config/cumulus/template/03-templates-j2",
    )

    def link(path: Path, target: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.symlink_to(os.path.relpath(target, path.parent))

    p2p = ztp / "config/cumulus/template/P2P"
    link(p2p / "p2p.xlsx", project / "p2p.xlsx")
    (p2p / "01-inventory.log").write_text("[Eth-SW]\nBorder01\nBorder02\n")
    (p2p / "02-port-mapping.log").write_text("[Eth-SW]\nswp53,swp53\n")
    (p2p / "output-p2p").mkdir()
    for directory in (ztp / "config/isc-dhcp-server", ztp / "config/cumulus/template"):
        link(directory / "01-global.yaml", global_file)
        link(directory / "02-devices_config.csv", project / "02-devices_config.csv")
    link(
        ztp / "config/isc-dhcp-server/02-subnet_config.csv",
        project / "02-dhcp-subnet_config.csv",
    )
    # A scoped ETH generation preserves the untouched families' existing
    # host files. Seed empty private files exactly as an installed tree does.
    for filename in ("dhcpd_eth.hosts", "dhcpd_ib.hosts", "dhcpd_nvl.hosts", "dhcpd.conf"):
        (ztp / "config/isc-dhcp-server" / filename).write_bytes(b"")
    output = project / "99-output-eth"
    output.mkdir()
    link(ztp / "config/cumulus/template/99-output", output)
    (ztp / "config/cumulus/latest_yaml").symlink_to("template/99-output/latest")
    return project, ztp


def _artifact_tree(root: Path) -> dict[str, tuple[str, str]]:
    """Include file bytes and symlink destinations, not volatile directory mtimes."""
    result = {}
    for item in sorted(root.rglob("*")):
        relative = str(item.relative_to(root))
        if item.is_symlink():
            result[relative] = ("link", os.readlink(item))
        elif item.is_file():
            result[relative] = ("file", hashlib.sha256(item.read_bytes()).hexdigest())
        elif item.is_dir():
            result[relative] = ("dir", "")
        else:
            raise AssertionError(f"unexpected artifact type: {item}")
    return result


def _publish_private_dhcp_outputs(project: Path, ztp: Path) -> None:
    """Model setup's canonical project-owned five-output publication."""
    dhcp_dir = ztp / "config/isc-dhcp-server"
    project_dhcp = project / "99-output-dhcp"
    project_dhcp.mkdir()
    for name in DHCP_OUTPUTS:
        public_output = dhcp_dir / name
        project_output = project_dhcp / name
        shutil.copy2(public_output, project_output)
        public_output.unlink()
        public_output.symlink_to(os.path.relpath(project_output, dhcp_dir))


def _write_same_inode(path: Path, data: bytes, original: os.stat_result) -> None:
    descriptor = os.open(path, os.O_RDWR | os.O_NOFOLLOW)
    try:
        current = os.fstat(descriptor)
        if not stat.S_ISREG(current.st_mode) or current.st_nlink != 1:
            raise AssertionError("mutation target must be a single-link regular file")
        if current.st_ino != original.st_ino or len(data) != original.st_size:
            raise AssertionError("A/B mutation must preserve inode and size")
        os.lseek(descriptor, 0, os.SEEK_SET)
        written = 0
        while written < len(data):
            written += os.write(descriptor, data[written:])
        os.ftruncate(descriptor, len(data))
        os.fsync(descriptor)
    finally:
        os.close(descriptor)
    os.utime(path, ns=(original.st_atime_ns, original.st_mtime_ns), follow_symlinks=False)
    after = path.stat()
    if (after.st_ino, after.st_size, after.st_mtime_ns) != (
        original.st_ino, original.st_size, original.st_mtime_ns,
    ):
        raise AssertionError("A/B mutation unexpectedly changed inode, size or mtime")


class GenerationInputBindingWorkflowTests(unittest.TestCase):
    """LOAD's frozen A authority must survive a real downstream c2 read."""

    def test_real_setup_death_disk_only_a_runtime_restore_uses_wal_not_argv_context(self):
        self._assert_real_persisted_a_runtime_switch(False)

    def test_real_setup_q_link_registration_holds_replaced_a_runtime_receipt(self):
        self._assert_real_persisted_a_runtime_switch(True)

    def test_real_setup_manifest_holds_a_receipt_replaced_during_wal_stage_sync(self):
        fixture = RELEASE_FIXTURE.ReleaseTransactionTests()
        fixture.setUp()
        try:
            state, context = fixture._prepare_persisted_a_runtime_switch()
            root, q, common, old_only = fixture.root.resolve(), state[1], state[2], state[3]
            journal, descriptor, _backend = state[-3:]
            shutil.copytree(ROOT / "tools", root / "tools")
            (root / "ztp/optimize").mkdir()
            shutil.copy2(ROOT / "ztp/optimize/sample_links.py", root / "ztp/optimize/sample_links.py")
            for name in ("11-load.py", "01-a-setup.py"):
                shutil.copy2(ROOT / "DAY0-Prepare" / name, root / "DAY0-Prepare" / name)
            RELEASE_FIXTURE.LOAD.release_deployment_lock(descriptor)
            fixture._cleanups = [item for item in fixture._cleanups
                                 if item[0] != RELEASE_FIXTURE.LOAD.release_deployment_lock]
            evidence, manifest = journal.read_bytes(), state[7].read_bytes()
            writer = """
import importlib.util, json, os, stat, sys
from pathlib import Path
from types import SimpleNamespace
root, q, common, old_only = map(Path, sys.argv[1:5])
def script(name):
    spec = importlib.util.spec_from_file_location('late_manifest_' + name, root / 'DAY0-Prepare' / (name + '.py'))
    api = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = api
    spec.loader.exec_module(api)
    return api
load, setup = script('11-load'), script('01-a-setup')
load.active_managed_services = lambda *args, **kwargs: ()
load.service_runtime_backend = lambda: SimpleNamespace(name='supervisor')
fd = load.acquire_deployment_lock(exclusive=True)
setup._SETUP_SWITCH_WAL_CONTEXT = (vars(load), fd)
setup.ZTP = str(root / 'ztp')
setup.MANIFEST_FILE = str(root / 'ztp/.setup_manifest')
# Limit the independent low-level fixture to the one actual WAL-owned link;
# the manifest writer, WAL registration and fsync/replace remain real scripts.
setup._collect_expected_links = lambda project: [str(common)]
common.unlink()
common.symlink_to(os.path.relpath(q / '01-global.yaml', common.parent))
old_only.unlink()
journal = root / 'ztp/.setup-switch-journal.json'
record = json.loads(journal.read_bytes())['body']['a_runtime']
receipt = Path(record['receipt_directory']['path']) / 'receipt.json'
original_owner = receipt.stat().st_ino
real_sync = os.fsync
injected = []
def sync(descriptor):
    result = real_sync(descriptor)
    owner = os.fstat(descriptor)
    if not injected and stat.S_ISREG(owner.st_mode):
        for staged in journal.parent.glob('.setup-switch-journal.*'):
            info = staged.lstat()
            if (info.st_dev, info.st_ino) == (owner.st_dev, owner.st_ino):
                receipt.rename(root / 'held-original-a-receipt')
                receipt.write_bytes((root / 'held-original-a-receipt').read_bytes())
                receipt.chmod(0o644)
                assert original_owner != receipt.stat().st_ino
                injected.append(True)
                break
    return result
os.fsync = sync
try:
    setup._write_manifest(str(q))
except ValueError as exc:
    assert injected == [True] and 'runtime' in str(exc), str(exc)
    print('SETUP_LATE_A_RECEIPT_MANIFEST_HELD')
    raise SystemExit(7)
else:
    print('SETUP_LATE_A_RECEIPT_MANIFEST_WRONGLY_PUBLISHED')
    raise SystemExit(8)
finally:
    os.fsync = real_sync
    load.release_deployment_lock(fd)
"""
            result = subprocess.run(
                [sys.executable, "-B", "-c", writer, str(root), str(q), str(common), str(old_only)],
                capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(7, result.returncode, result.stdout + result.stderr)
            self.assertIn("SETUP_LATE_A_RECEIPT_MANIFEST_HELD", result.stdout)
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(manifest, state[7].read_bytes(), "SETUP replaced P manifest after A authority drift")
            self.assertEqual(os.path.relpath(state[5], common.parent), os.readlink(common))
            self.assertFalse(os.path.lexists(old_only))
            self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
            self.assertEqual([journal], list(journal.parent.glob(".setup-switch-journal.*")),
                             "HOLD must preserve the authority and clean only its own intent stage")
        finally:
            fixture.doCleanups()
            fixture.tearDown()

    def test_real_setup_q_local_prewrite_holds_late_a_stage_sync(self):
        self._assert_real_latest_prewrite_holds_late_a("q-local")

    def test_real_setup_materialized_prewrite_holds_late_a_stage_sync(self):
        self._assert_real_latest_prewrite_holds_late_a("materialized")

    def _assert_real_latest_prewrite_holds_late_a(self, kind):
        fixture = RELEASE_FIXTURE.ReleaseTransactionTests()
        fixture.setUp()
        try:
            state, context = fixture._prepare_persisted_a_runtime_switch()
            root, q = fixture.root.resolve(), state[1].resolve()
            journal, descriptor, backend = state[-3:]
            # The real generic cleanup also selects this fixture's P latest
            # export. Freeze its independently specified row before Q starts.
            public_latest = root / "ztp/config/cumulus/latest_yaml"
            public_latest.unlink()
            public_latest.symlink_to(os.path.relpath(state[0] / "99-output-eth/latest", public_latest.parent))
            physical_links = {path.parent.resolve() / path.name: raw for path, raw in state[9].items()}
            state[7].write_text(f"# setup manifest — proj: {state[0]}\n"
                                + "".join(f"{path}\n" for path in physical_links)
                                + f"{public_latest}\n")
            journal.unlink()  # untouched private fixture authority only
            context = RELEASE_FIXTURE.LOAD.CrossProjectRollbackContext(
                state[0], context.selector_target, context.runtime_receipt,
                RELEASE_FIXTURE.LOAD.snapshot_active_setup_state(expected_project=state[0]),
            )
            with mock.patch.object(RELEASE_FIXTURE.LOAD, "active_managed_services", return_value=()):
                RELEASE_FIXTURE.LOAD.prepare_setup_switch_journal(
                    state[0], q, expected_new_links={state[2].parent.resolve() / state[2].name: state[5]},
                    deployment_lock_descriptor=descriptor, services_quiesced=True,
                    runtime_backend=backend, runtime_rollback_context=context,
                )
            latest, releases = fixture._q_local_latest_fixture(q)
            if kind == "materialized":
                latest.unlink()
                shutil.copytree(releases[1], latest)
            owner = latest.lstat().st_ino
            receipt = context.runtime_receipt / "receipt.json"
            receipt_owner, receipt_bytes = receipt.stat().st_ino, receipt.read_bytes()
            shutil.copytree(ROOT / "tools", root / "tools")
            (root / "ztp/optimize").mkdir()
            shutil.copy2(ROOT / "ztp/optimize/sample_links.py", root / "ztp/optimize/sample_links.py")
            for name in ("11-load.py", "01-a-setup.py"):
                shutil.copy2(ROOT / "DAY0-Prepare" / name, root / "DAY0-Prepare" / name)
            RELEASE_FIXTURE.LOAD.release_deployment_lock(descriptor)
            fixture._cleanups = [item for item in fixture._cleanups
                                 if item[0] != RELEASE_FIXTURE.LOAD.release_deployment_lock]
            evidence, manifest = journal.read_bytes(), state[7].read_bytes()
            links = {path: os.readlink(path) for path in physical_links}
            links[public_latest] = os.readlink(public_latest)
            writer = r'''
import importlib.util, json, os, stat, sys
from pathlib import Path
from types import SimpleNamespace
root, q = map(Path, sys.argv[1:3])
def script(name):
    spec = importlib.util.spec_from_file_location('late_latest_' + name, root / 'DAY0-Prepare' / (name + '.py'))
    api = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = api
    spec.loader.exec_module(api)
    return api
load, setup = script('11-load'), script('01-a-setup')
load.active_managed_services = lambda *args, **kwargs: ()
load.service_runtime_backend = lambda: SimpleNamespace(name='supervisor')
# Supply the actual copied LOAD module with its stopped service backend;
# every inventory, resolver, prewrite guard, WAL writer and fsync is real.
real_run_path = setup.runpy.run_path
setup.runpy.run_path = lambda path, *args, **kwargs: (vars(load) if Path(path).name == '11-load.py'
    else real_run_path(path, *args, **kwargs))
setup.ZTP = str(root / 'ztp')
setup.MANIFEST_FILE = str(root / 'ztp/.setup_manifest')
fd = load.acquire_deployment_lock(exclusive=True)
journal = root / 'ztp/.setup-switch-journal.json'
record = json.loads(journal.read_bytes())['body']['a_runtime']
receipt = Path(record['receipt_directory']['path']) / 'receipt.json'
real_sync, injected = os.fsync, []
def sync(descriptor):
    result = real_sync(descriptor)
    owner = os.fstat(descriptor)
    if not injected and stat.S_ISREG(owner.st_mode):
        for staged in journal.parent.glob('.setup-switch-journal.*'):
            info = staged.lstat()
            if staged != journal and (info.st_dev, info.st_ino) == (owner.st_dev, owner.st_ino):
                original = receipt.stat().st_ino
                receipt.rename(root / 'held-original-latest-a-receipt')
                receipt.write_bytes((root / 'held-original-latest-a-receipt').read_bytes())
                receipt.chmod(0o644)
                assert original != receipt.stat().st_ino
                injected.append(True)
                break
    return result
os.fsync = sync
try:
    setup._require_setup_switch_wal_prewrite(str(q), fd)
except ValueError as exc:
    assert injected == [True] and 'runtime' in str(exc), str(exc)
    print('SETUP_LATEST_LATE_A_PREWRITE_HELD')
    raise SystemExit(7)
else:
    assert injected == [True], 'late A injection did not reach the real prewrite'
    print('SETUP_LATEST_LATE_A_PREWRITE_WRONGLY_AUTHORIZED')
    raise SystemExit(8)
finally:
    os.fsync = real_sync
    load.release_deployment_lock(fd)
'''
            result = subprocess.run(
                [sys.executable, "-B", "-c", writer, str(root), str(q)],
                capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(7, result.returncode, result.stdout + result.stderr)
            self.assertIn("SETUP_LATEST_LATE_A_PREWRITE_HELD", result.stdout)
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(manifest, state[7].read_bytes())
            self.assertEqual(links, {path: os.readlink(path) for path in links})
            self.assertEqual(owner, latest.lstat().st_ino)
            if kind == "materialized":
                self.assertEqual(b"hostname: q-1\n", (latest / "leaf.yaml").read_bytes())
            else:
                self.assertEqual(releases[0].name, os.readlink(latest))
            self.assertFalse(list(latest.parent.glob(".latest.setup-switch-wal-*")))
            self.assertNotEqual(receipt_owner, receipt.stat().st_ino)
            self.assertEqual(receipt_bytes, receipt.read_bytes())
            self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
            self.assertEqual([journal], list(journal.parent.glob(".setup-switch-journal.*")))
            backend.start.assert_not_called()
        finally:
            fixture.doCleanups()
            fixture.tearDown()

    def test_real_setup_death_disk_only_a_runtime_late_wal_stops_remaining_files(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="wal")

    def test_real_setup_death_disk_only_a_runtime_late_receipt_stops_remaining_files(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="a")

    def test_real_setup_death_disk_only_a_runtime_late_q_commit_stops_remaining_files(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="q")

    def test_real_setup_death_disk_only_a_runtime_missing_latest_sync_retry_retains_wal(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="latest-sync")

    def test_real_setup_death_runtime_orphan_holds_cold_entry_and_retirement(self):
        for kind in ("file", "dangling", "late"):
            with self.subTest(kind=kind):
                self._assert_real_persisted_a_runtime_switch(False, late_authority="runtime-orphan-" + kind)

    def test_real_setup_death_managed_key_orphan_holds_before_flat_write_and_retirement(self):
        for kind in ("file", "dangling", "late"):
            with self.subTest(kind=kind):
                self._assert_real_persisted_a_runtime_switch(False, late_authority="key-orphan-" + kind)

    def test_real_setup_death_prefix_marker_foreign_stage_is_preserved(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="prefix-stage")

    def test_real_setup_death_prefix_marker_sync_failure_retains_wal_until_retry(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="prefix-sync")

    def test_real_setup_death_runtime_surface_orphans_hold_before_p_write(self):
        for surface in ("yaml-link", "yaml-missing", "yaml-link-old", "yaml-missing-old", "prefix-link-new", "prefix-link-old", "prefix-marker"):
            with self.subTest(surface=surface):
                self._assert_real_persisted_a_runtime_switch(False, late_authority="surface-orphan-" + surface + ":cold")

    def test_real_setup_death_runtime_surface_orphans_hold_after_runtime_verify(self):
        for surface in ("yaml-link", "yaml-missing", "yaml-link-old", "yaml-missing-old", "prefix-link-new", "prefix-link-old", "prefix-marker"):
            with self.subTest(surface=surface):
                self._assert_real_persisted_a_runtime_switch(False, late_authority="surface-orphan-" + surface + ":late")

    def test_real_setup_death_yaml_parent_rebind_holds_before_p_write(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="yaml-parent-rebind:cold")

    def test_real_setup_death_yaml_parent_rebind_holds_after_runtime_verify(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="yaml-parent-rebind:late")

    def test_real_setup_death_missing_a_yaml_restore_does_not_create_parent(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="missing-yaml-parent")

    def test_real_setup_death_explicit_recovery_cli_restores_a_and_stays_stopped(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="explicit-recovery")

    def test_real_setup_death_explicit_recovery_cli_wrong_q_holds_original_evidence(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="explicit-recovery:wrong-q")

    def test_real_setup_death_explicit_recovery_cli_running_service_holds_original_evidence(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="explicit-recovery:running")

    def test_real_setup_death_explicit_recovery_cli_committed_b_holds_original_evidence(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="explicit-recovery:committed-b")

    def test_real_setup_death_explicit_recovery_cli_late_dhcp_retains_wal_after_a_restore(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="explicit-recovery:late-dhcp")

    def test_real_setup_death_yaml_parent_legacy_wal_holds_without_inventing_identity(self):
        self._assert_real_persisted_a_runtime_switch(False, late_authority="yaml-parent-rebind:legacy")

    def test_real_setup_death_runtime_file_restore_preserves_foreign_stage(self):
        for kind in ("inode", "bytes", "hold"):
            with self.subTest(kind=kind):
                self._assert_real_persisted_a_runtime_switch(False, late_authority="file-stage-" + kind)

    def _assert_real_persisted_a_runtime_switch(self, replace_a_receipt, *, late_authority=None):
        fixture = RELEASE_FIXTURE.ReleaseTransactionTests()
        fixture.setUp()
        try:
            managed_key = bool(late_authority and late_authority.startswith("key-orphan-"))
            prefix = bool(late_authority and late_authority.startswith(("prefix-", "surface-orphan-prefix-")))
            state, context = fixture._prepare_persisted_a_runtime_switch(managed_key=managed_key, prefix=prefix)
            prefix_before = ((fixture.root / ".ztp-prefix-publication.json").read_bytes()
                             if prefix else None)
            root = fixture.root.resolve()
            p, q, common = state[:3]
            journal, descriptor, _backend = state[-3:]
            shutil.copytree(ROOT / "tools", root / "tools")
            (root / "ztp/optimize").mkdir()
            shutil.copy2(ROOT / "ztp/optimize/sample_links.py", root / "ztp/optimize/sample_links.py")
            for name in ("11-load.py", "01-a-setup.py"):
                shutil.copy2(ROOT / "DAY0-Prepare" / name, root / "DAY0-Prepare" / name)
            # The next process acquires its own lock; no parent context or
            # receipt path is serialized into its command line.
            RELEASE_FIXTURE.LOAD.release_deployment_lock(descriptor)
            fixture._cleanups = [item for item in fixture._cleanups
                                 if item[0] != RELEASE_FIXTURE.LOAD.release_deployment_lock]
            before_writer = journal.read_bytes()
            writer = """
import importlib.util, os, signal, sys
from pathlib import Path
from types import SimpleNamespace
root, q, common = map(Path, sys.argv[1:4])
def script(name):
    spec = importlib.util.spec_from_file_location('persisted_' + name, root / 'DAY0-Prepare' / (name + '.py'))
    api = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = api
    spec.loader.exec_module(api)
    return api
load, setup = script('11-load'), script('01-a-setup')
load.active_managed_services = lambda *args, **kwargs: ()
load.service_runtime_backend = lambda: SimpleNamespace(name='supervisor')
fd = load.acquire_deployment_lock(exclusive=True)
setup._SETUP_SWITCH_WAL_CONTEXT = (vars(load), fd)
if sys.argv[4] == 'replace':
    import json
    record = json.loads((root / 'ztp/.setup-switch-journal.json').read_bytes())['body']['a_runtime']
    path = Path(record['receipt_directory']['path']) / 'receipt.json'
    path.rename(root / 'held-original-a-receipt')
    path.write_bytes((root / 'held-original-a-receipt').read_bytes())
    path.chmod(0o644)
try:
    setup._register_setup_switch_link_before_write(str(common), str(q / '01-global.yaml'))
except ValueError as exc:
    assert sys.argv[4] == 'replace' and 'runtime' in str(exc), str(exc)
    print('SETUP_REPLACED_A_RECEIPT_HELD')
    load.release_deployment_lock(fd)
    raise SystemExit(7)
common.unlink()
common.symlink_to(os.path.relpath(q / '01-global.yaml', common.parent))
(root / 'ztp/ztp-bootstrap_oob.sh').write_bytes(b'partial B runtime\\n')
if sys.argv[4] == 'late' and sys.argv[5].startswith('key-orphan-'):
    key = root / 'ztp/config/publickey/laptop.pub'
    key.unlink()
    key.write_bytes(b'partial B flat key\\n')
if sys.argv[4] == 'late':
    (root / 'ztp/ztp-bootstrap_oobofoob.sh').write_bytes(b'partial B runtime\\n')
    (root / 'ztp/ztp.json').write_bytes(b'partial B runtime\\n')
    if sys.argv[5].startswith(('prefix-', 'surface-orphan-prefix-')):
        (root / '.ztp-prefix-publication.json').write_bytes(b'partial B prefix marker\\n')
    if sys.argv[5] == 'latest-sync':
        import json
        p = Path(json.loads((root / 'ztp/.setup-switch-journal.json').read_bytes())['body']['old_project']['path'])
        (p / '99-output-ib_nvl').mkdir()
        (p / '99-output-ib_nvl/latest').symlink_to('independent-partial-B-target')
os.kill(os.getpid(), signal.SIGKILL)
"""
            writer_command = [sys.executable, "-B", "-c", writer, str(root), str(q), str(common),
                              "replace" if replace_a_receipt else "late" if late_authority else "exact"]
            if late_authority:
                writer_command.append(late_authority)
            first = subprocess.run(writer_command,
                                   capture_output=True, text=True, check=False, timeout=30)
            if replace_a_receipt:
                self.assertEqual(7, first.returncode, first.stdout + first.stderr)
                self.assertIn("SETUP_REPLACED_A_RECEIPT_HELD", first.stdout)
                self.assertEqual(before_writer, journal.read_bytes())
                self.assertEqual(context.setup_state.links, {p: os.readlink(p) for p in context.setup_state.links})
                self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                return
            self.assertEqual(-9, first.returncode, first.stdout + first.stderr)
            evidence = journal.read_bytes()
            mixed = os.readlink(common)
            if late_authority and late_authority.startswith("explicit-recovery"):
                runtime_before = (root / "ztp/ztp-bootstrap_oob.sh").read_bytes()
                recovery = r'''
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root, project = map(Path, sys.argv[1:3])
spec = importlib.util.spec_from_file_location('explicit_recovery_cold_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
def forbidden(*args, **kwargs):
    raise AssertionError('explicit recovery must not regenerate, stop or start services')
load.runtime_os = lambda: 'Linux'
load.active_managed_services = lambda *args, **kwargs: ('apache2',) if sys.argv[3] == 'explicit-recovery:running' else ()
load.service_runtime_backend = lambda: SimpleNamespace(name='supervisor', start=forbidden, stop=forbidden)
load.dhcp_file_mappings = lambda *args, **kwargs: {root / 'unused-source': root / 'isolated-etc-dhcp/dhcpd.conf'}
if sys.argv[3] == 'explicit-recovery:late-dhcp':
    original_restore = load.restore_runtime_rollback_receipt
    def inject_after_runtime(*args, **kwargs):
        result = original_restore(*args, **kwargs)
        (root / 'isolated-etc-dhcp/.load-dhcp-transaction-independent').mkdir(parents=True)
        return result
    load.restore_runtime_rollback_receipt = inject_after_runtime
for name in ('initialize_from_template', 'activate_project', 'render_ztp_runtime', 'start_services',
             'quiesce_services', 'validate_inputs', 'prepare_pubkeys'):
    setattr(load, name, forbidden)
raise SystemExit(load.main([str(project), '--host-role', 'management-server', '--recover-setup-switch']))
'''
                wrong_q = late_authority.endswith(":wrong-q")
                if late_authority.endswith(":committed-b"):
                    parent = q / "99-output-ztp/current-release.json"
                    parent.parent.mkdir()
                    parent.write_bytes(b'{"release_id":"independent-committed-B"}\n')
                holds = ":" in late_authority
                result = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), str(p if wrong_q else q), late_authority],
                    capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(1 if holds else 0, result.returncode, result.stdout + result.stderr)
                if holds:
                    self.assertEqual(evidence, journal.read_bytes())
                    if late_authority.endswith(":late-dhcp"):
                        self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                        self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                        self.assertTrue((root / "isolated-etc-dhcp/.load-dhcp-transaction-independent").is_dir())
                    else:
                        self.assertEqual(mixed, os.readlink(common))
                        self.assertEqual(runtime_before, (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                    self.assertIn("HOLD", result.stderr)
                else:
                    self.assertFalse(os.path.lexists(journal))
                    self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                    self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                    self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                self.assertFalse((root / "service-started").exists())
                return
            if late_authority == "missing-yaml-parent":
                missing = p / "99-output-ib_nvl/latest"
                self.assertFalse(os.path.lexists(missing.parent))
                recovery = r'''
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('missing_yaml_parent_cold_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
backend = SimpleNamespace(name='supervisor', start=lambda *args: (_ for _ in ()).throw(AssertionError('service start forbidden')))
fd = load.acquire_deployment_lock(exclusive=True)
try:
    load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=fd, services_quiesced=True, runtime_backend=backend,
        restore_runtime_from_wal=True)
finally:
    load.release_deployment_lock(fd)
'''
                result = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                    capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(0, result.returncode, result.stdout + result.stderr)
                self.assertFalse(os.path.lexists(missing.parent), "cold recovery created an unjournaled A output directory")
                self.assertFalse(os.path.lexists(missing))
                self.assertFalse(os.path.lexists(journal))
                self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                self.assertFalse((root / "service-started").exists())
                return
            if late_authority and late_authority.startswith("yaml-parent-rebind:"):
                phase = late_authority.split(":")[1]
                latest = p / "99-output-eth/latest"
                if phase == "legacy":
                    envelope = json.loads(journal.read_bytes())
                    envelope["body"]["a_runtime"].pop("yaml_link_parents", None)
                    payload = json.dumps(envelope["body"], sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode()
                    envelope["sha256"] = hashlib.sha256(payload).hexdigest()
                    journal.write_bytes(json.dumps(envelope, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode() + b"\n")
                def snapshot(path):
                    return (path.lstat().st_ino, os.readlink(path) if path.is_symlink() else path.read_bytes())
                before_paths = [common, state[7], root / "ztp/ztp-bootstrap_oob.sh"]
                before = [snapshot(path) for path in before_paths]
                wal_before = snapshot(journal)
                recovery = r'''
import importlib.util, json, os, sys
from pathlib import Path
from types import SimpleNamespace
root, latest, phase = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3]
spec = importlib.util.spec_from_file_location('yaml_parent_cold_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
def rebind():
    parent, retained = latest.parent, root / 'held-original-a-yaml-parent'
    owner, leaf = parent.lstat().st_ino, (latest.lstat().st_ino, os.readlink(latest))
    parent.rename(retained)
    parent.mkdir()
    for child in retained.iterdir():
        child.rename(parent / child.name)
    assert owner != parent.lstat().st_ino
    assert leaf == (latest.lstat().st_ino, os.readlink(latest))
    (root / 'yaml-parent-injection.json').write_text(json.dumps({'parent': parent.lstat().st_ino, 'leaf': leaf}))
if phase == 'cold':
    rebind()
elif phase == 'late':
    restore = load.restore_runtime_rollback_receipt
    def late_rebind(*args, **kwargs):
        result = restore(*args, **kwargs)
        rebind()
        return result
    load.restore_runtime_rollback_receipt = late_rebind
backend = SimpleNamespace(name='supervisor', start=lambda *args: (_ for _ in ()).throw(AssertionError('service start forbidden')))
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=fd, services_quiesced=True, runtime_backend=backend,
            restore_runtime_from_wal=True)
    except load.LoadError as exc:
        print('EXPECTED YAML PARENT HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(fd)
'''
                result = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), str(latest), phase],
                    capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(7, result.returncode, result.stdout + result.stderr)
                self.assertEqual(wal_before, snapshot(journal))
                if phase == "legacy":
                    self.assertEqual(before, [snapshot(path) for path in before_paths])
                    self.assertFalse((root / "yaml-parent-injection.json").exists())
                    self.assertFalse((root / "held-original-a-yaml-parent").exists())
                    return
                held = json.loads((root / "yaml-parent-injection.json").read_text())
                self.assertEqual(held["parent"], latest.parent.lstat().st_ino)
                self.assertEqual(tuple(held["leaf"]), snapshot(latest))
                self.assertTrue((root / "held-original-a-yaml-parent").is_dir())
                if phase == "cold":
                    self.assertEqual(before, [snapshot(path) for path in before_paths])
                else:
                    self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                    self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                self.assertFalse((root / "service-started").exists())
                return
            if late_authority and late_authority.startswith("surface-orphan-"):
                surface, phase = late_authority.removeprefix("surface-orphan-").split(":")
                targets = {
                    "yaml-link": p / "99-output-eth/latest",
                    "yaml-missing": p / "99-output-ib_nvl/latest",
                    "yaml-link-old": p / "99-output-eth/latest",
                    "yaml-missing-old": p / "99-output-ib_nvl/latest",
                    "prefix-link-new": root / "r22-prefix", "prefix-link-old": root / "r22-prefix",
                    "prefix-marker": root / ".ztp-prefix-publication.json",
                }
                target = targets[surface]
                target.parent.mkdir(exist_ok=True)
                stem = ".rollback." if surface in {"yaml-link-old", "yaml-missing-old", "prefix-link-old", "prefix-marker"} else ".setup-rollback."
                orphan = target.parent / ("." + target.name + stem + "independent-orphan")
                if phase == "cold":
                    orphan.symlink_to("independent-surface-stage-evidence")
                def snapshot(path):
                    if not os.path.lexists(path):
                        return None
                    return (path.lstat().st_ino, os.readlink(path) if path.is_symlink() else path.read_bytes())
                paths = [common, state[7], root / "ztp/ztp-bootstrap_oob.sh", target]
                before = [snapshot(path) for path in paths]
                wal_before = snapshot(journal)
                orphan_before = snapshot(orphan)
                recovery = r'''
import importlib.util, json, os, sys
from pathlib import Path
from types import SimpleNamespace
root, orphan, phase = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3]
spec = importlib.util.spec_from_file_location('surface_orphan_cold_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
if phase == 'late':
    restore = load.restore_runtime_rollback_receipt
    def late_orphan(*args, **kwargs):
        result = restore(*args, **kwargs)
        orphan.symlink_to('independent-surface-stage-evidence')
        (root / 'surface-orphan-injection.json').write_text(json.dumps({'ino': orphan.lstat().st_ino}))
        return result
    load.restore_runtime_rollback_receipt = late_orphan
backend = SimpleNamespace(name='supervisor', start=lambda *args: (_ for _ in ()).throw(AssertionError('service start forbidden')))
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=fd, services_quiesced=True, runtime_backend=backend,
            restore_runtime_from_wal=True)
    except load.LoadError as exc:
        print('EXPECTED SURFACE ORPHAN HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(fd)
'''
                result = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), str(orphan), phase],
                    capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(7, result.returncode, result.stdout + result.stderr)
                self.assertEqual(wal_before, snapshot(journal))
                self.assertEqual("independent-surface-stage-evidence", os.readlink(orphan))
                if phase == "cold":
                    self.assertEqual(orphan_before, snapshot(orphan))
                    self.assertEqual(before, [snapshot(path) for path in paths], "orphan detected after P/flat mutation")
                else:
                    held = json.loads((root / "surface-orphan-injection.json").read_text())
                    self.assertEqual(held["ino"], orphan.lstat().st_ino)
                    self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                    if prefix:
                        self.assertEqual(prefix_before, (root / ".ztp-prefix-publication.json").read_bytes())
                    self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                    self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                self.assertFalse((root / "service-started").exists())
                return
            if prefix:
                marker = root / ".ztp-prefix-publication.json"
                marker_before = (marker.lstat().st_ino, marker.read_bytes())
                wal_before = (journal.lstat().st_ino, evidence)
                recovery = r'''
import importlib.util, json, os, stat, sys
from pathlib import Path
from types import SimpleNamespace
root, kind, phase = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
spec = importlib.util.spec_from_file_location('prefix_cold_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
marker, journal = load.ZTP_PREFIX_MARKER, root / 'ztp/.setup-switch-journal.json'
body = json.loads(journal.read_bytes())['body']
receipt = Path(body['a_runtime']['receipt_directory']['path']) / 'receipt.json'
import base64
expected = base64.b64decode(json.loads(receipt.read_bytes())['prefix']['marker_b64'])
sync, parent_sync, injected = os.fsync, load._fsync_parent_directory, []
def foreign_stage(fd):
    result = sync(fd)
    info = os.fstat(fd)
    if not injected and stat.S_ISREG(info.st_mode):
        for stage in marker.parent.glob('.' + marker.name + '.rollback.*'):
            if stage.lstat().st_ino == info.st_ino:
                stage.rename(root / 'held-original-prefix-stage')
                stage.write_bytes(b'independent foreign prefix stage\n')
                stage.chmod(0o640)
                injected.append(True)
                (root / 'prefix-stage-evidence.json').write_text(json.dumps({'path': str(stage), 'ino': stage.lstat().st_ino}))
    return result
def marker_parent_sync(parent):
    if parent == marker.parent and marker.read_bytes() == expected:
        assert journal.is_file(), 'WAL retired before prefix namespace sync'
        injected.append(True)
        if phase == 'fault':
            raise OSError('independent prefix marker parent sync')
    return parent_sync(parent)
if kind == 'prefix-stage':
    os.fsync = foreign_stage
else:
    load._fsync_parent_directory = marker_parent_sync
backend = SimpleNamespace(name='supervisor', start=lambda *args: (_ for _ in ()).throw(AssertionError('service start forbidden')))
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal, deployment_lock_descriptor=fd,
            services_quiesced=True, runtime_backend=backend, restore_runtime_from_wal=True)
    except load.LoadError as exc:
        assert injected, 'prefix fault injection was not reached'
        print('EXPECTED PREFIX HOLD:', exc)
        raise SystemExit(7)
    assert phase == 'retry' and injected, 'prefix namespace was not synchronized'
    assert not journal.exists()
finally:
    load.release_deployment_lock(fd)
'''
                result = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), late_authority, "fault"],
                    capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(7, result.returncode, result.stdout + result.stderr)
                self.assertEqual(wal_before, (journal.lstat().st_ino, journal.read_bytes()))
                if late_authority == "prefix-stage":
                    self.assertEqual(marker_before, (marker.lstat().st_ino, marker.read_bytes()))
                    held = json.loads((root / "prefix-stage-evidence.json").read_text())
                    stage = Path(held["path"])
                    self.assertTrue(os.path.lexists(stage), "foreign prefix stage was cleaned")
                    self.assertEqual((held["ino"], b"independent foreign prefix stage\n", 0o640),
                                     (stage.lstat().st_ino, stage.read_bytes(), stat.S_IMODE(stage.stat().st_mode)))
                    self.assertEqual(prefix_before, (root / "held-original-prefix-stage").read_bytes())
                else:
                    self.assertEqual(prefix_before, marker.read_bytes())
                    retry = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), late_authority, "retry"],
                        capture_output=True, text=True, check=False, timeout=30)
                    self.assertEqual(0, retry.returncode, retry.stdout + retry.stderr)
                    self.assertFalse(os.path.lexists(journal))
                    self.assertEqual(prefix_before, marker.read_bytes())
                self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                self.assertFalse((root / "service-started").exists())
                return
            if late_authority and late_authority.startswith(("runtime-orphan-", "key-orphan-")):
                kind = late_authority.removeprefix("key-orphan-" if managed_key else "runtime-orphan-")
                target = root / ("ztp/config/publickey/laptop.pub" if managed_key else "ztp/ztp-bootstrap_oob.sh")
                orphan = target.parent / (f".{target.name}.setup-rollback.independent-orphan" if managed_key
                                          else f".{target.name}.rollback-independent-orphan")
                if kind == "file":
                    orphan.write_bytes(b"independent unknown earlier runtime stage\n")
                elif kind == "dangling":
                    orphan.symlink_to("independent-absent-runtime-evidence")
                before = (common.lstat().st_ino, mixed, state[7].lstat().st_ino,
                          state[7].read_bytes(), target.lstat().st_ino, target.read_bytes())
                runtime = root / "ztp/ztp-bootstrap_oob.sh"
                runtime_before = (runtime.lstat().st_ino, runtime.read_bytes())
                wal_before = (journal.lstat().st_ino, journal.read_bytes())
                recovery = r'''
import importlib.util, json, os, sys
from pathlib import Path
from types import SimpleNamespace
root, kind, managed_key = Path(sys.argv[1]), sys.argv[2], sys.argv[3] == 'key'
spec = importlib.util.spec_from_file_location('runtime_orphan_cold_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
target = root / ('ztp/config/publickey/laptop.pub' if managed_key else 'ztp/ztp-bootstrap_oob.sh')
orphan = target.parent / ('.' + target.name + ('.setup-rollback.independent-orphan' if managed_key
                                            else '.rollback-independent-orphan'))
if kind == 'late':
    restore = load.restore_runtime_rollback_receipt
    def late_orphan(*args, **kwargs):
        result = restore(*args, **kwargs)
        orphan.write_bytes(b'independent unknown late runtime stage\n')
        (root / 'runtime-orphan-injection.json').write_text(json.dumps({'ino': orphan.lstat().st_ino}))
        return result
    load.restore_runtime_rollback_receipt = late_orphan
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=fd, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'), restore_runtime_from_wal=True)
    except load.LoadError as exc:
        print('EXPECTED UNKNOWN RUNTIME STAGE HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(fd)
'''
                orphan_inode = orphan.lstat().st_ino if kind != "late" else None
                result = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), kind,
                                         "key" if managed_key else "file"],
                    capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(7, result.returncode, result.stdout + result.stderr)
                self.assertEqual(wal_before, (journal.lstat().st_ino, journal.read_bytes()))
                if kind == "late":
                    marker = json.loads((root / "runtime-orphan-injection.json").read_text())
                    self.assertEqual(marker["ino"], orphan.lstat().st_ino)
                    self.assertEqual(b"independent unknown late runtime stage\n", orphan.read_bytes())
                    self.assertEqual(b"A laptop key\n" if managed_key else b"A oob bootstrap\n", target.read_bytes())
                    if managed_key:
                        self.assertTrue(target.is_symlink())
                        self.assertEqual(os.path.relpath(context.project / "laptop.pub", target.parent), os.readlink(target))
                    self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                    self.assertEqual(context.setup_state.links, {p: os.readlink(p) for p in context.setup_state.links})
                else:
                    self.assertEqual(runtime_before, (runtime.lstat().st_ino, runtime.read_bytes()))
                    self.assertEqual(before, (common.lstat().st_ino, os.readlink(common), state[7].lstat().st_ino,
                                      state[7].read_bytes(), target.lstat().st_ino, target.read_bytes()))
                    self.assertEqual(orphan_inode, orphan.lstat().st_ino)
                    if kind == "file":
                        self.assertEqual(b"independent unknown earlier runtime stage\n", orphan.read_bytes())
                    else:
                        self.assertEqual("independent-absent-runtime-evidence", os.readlink(orphan))
                self.assertFalse((root / "service-started").exists())
                return
            if late_authority and late_authority.startswith("file-stage-"):
                recovery = r'''
import importlib.util, json, os, stat, sys
from pathlib import Path
from types import SimpleNamespace
root, kind = Path(sys.argv[1]), sys.argv[2]
spec = importlib.util.spec_from_file_location('foreign_runtime_stage_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
target, journal = root / 'ztp/ztp-bootstrap_oob.sh', root / 'ztp/.setup-switch-journal.json'
sync, injected = os.fsync, []
def foreign_stage(fd):
    result = sync(fd)
    info = os.fstat(fd)
    if not injected and stat.S_ISREG(info.st_mode):
        for stage in target.parent.glob('.' + target.name + '.rollback-*'):
            if stage.lstat().st_ino == info.st_ino:
                held = root / 'held-original-runtime-stage'
                if kind != 'bytes':
                    stage.rename(held)
                stage.write_bytes(b'independent foreign runtime stage\n')
                stage.chmod(0o640)
                if kind == 'hold':
                    original = journal.read_bytes()
                    journal.rename(root / 'held-original-runtime-stage-wal')
                    journal.write_bytes(original)
                    journal.chmod(0o600)
                injected.append(True)
                (root / 'runtime-stage-injection.json').write_text(json.dumps({
                    'path': str(stage), 'ino': stage.lstat().st_ino, 'wal_ino': journal.lstat().st_ino}))
    return result
os.fsync = foreign_stage
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal, deployment_lock_descriptor=fd,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
            restore_runtime_from_wal=True)
    except load.LoadError as exc:
        assert injected == [True], str(exc)
        print('EXPECTED FOREIGN RUNTIME STAGE HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(fd)
'''
                target = root / "ztp/ztp-bootstrap_oob.sh"
                target_before = (target.lstat().st_ino, target.read_bytes())
                kind = late_authority.removeprefix("file-stage-")
                held = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), kind],
                    capture_output=True, text=True, check=False, timeout=30)
                marker = json.loads((root / "runtime-stage-injection.json").read_text())
                self.assertEqual(7, held.returncode, held.stdout + held.stderr)
                self.assertEqual(target_before, (target.lstat().st_ino, target.read_bytes()))
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(marker["wal_ino"], journal.lstat().st_ino)
                stage = Path(marker["path"])
                self.assertTrue(os.path.lexists(stage), "cold recovery cleaned a foreign runtime stage")
                self.assertEqual(marker["ino"], stage.lstat().st_ino)
                self.assertEqual(b"independent foreign runtime stage\n", stage.read_bytes())
                if kind != "bytes":
                    self.assertEqual(b"A oob bootstrap\n", (root / "held-original-runtime-stage").read_bytes())
                if kind == "hold":
                    self.assertEqual(evidence, (root / "held-original-runtime-stage-wal").read_bytes())
                self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                self.assertFalse((root / "service-started").exists())
                return
            if late_authority == "latest-sync":
                recovery = """
import importlib.util, json, os, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('missing_latest_runtime_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
journal = root / 'ztp/.setup-switch-journal.json'
body = json.loads(journal.read_bytes())['body']
latest = Path(body['old_project']['path']) / '99-output-ib_nvl/latest'
sync = load._fsync_parent_directory
seen = []
def missing_latest_parent_sync(parent):
    if parent == latest.parent and not os.path.lexists(latest):
        assert journal.is_file(), 'WAL retired before missing latest parent sync'
        seen.append(True)
        if sys.argv[2] == 'fault':
            raise OSError('independent missing latest parent sync')
    return sync(parent)
load._fsync_parent_directory = missing_latest_parent_sync
backend = SimpleNamespace(name='supervisor', start=lambda *args: (_ for _ in ()).throw(AssertionError('service start forbidden')))
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal, deployment_lock_descriptor=fd,
            services_quiesced=True, runtime_backend=backend, restore_runtime_from_wal=True)
    except load.LoadError as error:
        assert sys.argv[2] == 'fault' and 'independent missing latest parent sync' in str(error), str(error)
        assert seen and not os.path.lexists(latest)
        raise SystemExit(7)
    assert sys.argv[2] == 'retry' and seen, 'reentry failed to repay missing latest sync'
    assert not journal.exists() and not os.path.lexists(latest)
finally:
    load.release_deployment_lock(fd)
"""
                for phase, code in (("fault", 7), ("retry", 0)):
                    result = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), phase],
                                            capture_output=True, text=True, check=False, timeout=30)
                    self.assertEqual(code, result.returncode, result.stdout + result.stderr)
                    self.assertFalse(os.path.lexists(p / "99-output-ib_nvl/latest"))
                    if phase == "fault":
                        self.assertEqual(evidence, journal.read_bytes())
                    else:
                        self.assertFalse(os.path.lexists(journal))
                    self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                    self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                    self.assertFalse((root / "service-started").exists())
                return
            if late_authority:
                recovery = """
import importlib.util, json, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
kind = sys.argv[2]
spec = importlib.util.spec_from_file_location('late_a_runtime_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
journal = root / 'ztp/.setup-switch-journal.json'
body = json.loads(journal.read_bytes())['body']
authority = (journal if kind == 'wal' else Path(body['a_runtime']['receipt_directory']['path']) / 'receipt.json'
             if kind == 'a' else Path(body['new_project']['path']) / '99-output-ztp/current-release.json')
prior = authority.read_bytes() if kind != 'q' else None
restore = load._restore_receipt_file
injected = []
def restore_then_drift(path, record, payload, **guard):
    result = restore(path, record, payload, **guard)
    if path.name == 'ztp-bootstrap_oob.sh' and not injected:
        if kind == 'q':
            authority.parent.mkdir(parents=True, exist_ok=True)
            authority.write_bytes(b'independent Q commit marker\\n')
        else:
            owner = authority.stat().st_ino
            authority.rename(root / 'held-original-a-runtime-authority')
            authority.write_bytes(prior)
            authority.chmod(0o600 if kind == 'wal' else 0o644)
            assert owner != authority.stat().st_ino
        injected.append(authority.stat().st_ino)
    return result
load._restore_receipt_file = restore_then_drift
backend = SimpleNamespace(name='supervisor', start=lambda *args: (_ for _ in ()).throw(AssertionError('service start forbidden')))
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal, deployment_lock_descriptor=fd,
            services_quiesced=True, runtime_backend=backend, restore_runtime_from_wal=True)
    except load.LoadError as error:
        expected = 'journal changed' if kind == 'wal' else 'runtime' if kind == 'a' else 'Q parent changed'
        assert expected in str(error), str(error)
        assert len(injected) == 1, 'runtime authority injection not reached'
        assert (root / 'ztp/ztp-bootstrap_oobofoob.sh').read_bytes() == b'partial B runtime\\n', 'authority drift licensed later runtime write'
        assert (root / 'ztp/ztp.json').read_bytes() == b'partial B runtime\\n', 'authority drift licensed later JSON write'
        assert authority.stat().st_ino == injected[0]
        assert authority.read_bytes() == (prior if kind != 'q' else b'independent Q commit marker\\n')
        raise SystemExit(7)
    raise AssertionError('A runtime accepted late authority drift')
finally:
    load.release_deployment_lock(fd)
"""
                held = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), late_authority],
                                      capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(7, held.returncode, held.stdout + held.stderr)
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
                self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
                self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                self.assertEqual(b"partial B runtime\n", (root / "ztp/ztp-bootstrap_oobofoob.sh").read_bytes())
                self.assertEqual(b"partial B runtime\n", (root / "ztp/ztp.json").read_bytes())
                self.assertFalse((root / "service-started").exists())
                return
            recovery = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('disk_only_runtime_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
backend = SimpleNamespace(name='supervisor', start=lambda *args: (_ for _ in ()).throw(AssertionError('service start forbidden')))
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=fd, services_quiesced=True, runtime_backend=backend,
            restore_runtime_from_wal=(sys.argv[2] == 'explicit'))
    except load.LoadError as exc:
        assert sys.argv[2] == 'bare' and 'runtime' in str(exc), str(exc)
        print('BARE_RUNTIME_RECOVERY_HELD')
    else:
        assert sys.argv[2] == 'explicit'
        print('DISK_ONLY_A_RUNTIME_RECOVERED')
finally:
    load.release_deployment_lock(fd)
"""
            held = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), "bare"],
                                  capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(0, held.returncode, held.stdout + held.stderr)
            self.assertIn("BARE_RUNTIME_RECOVERY_HELD", held.stdout)
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            restored = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), "explicit"],
                                      capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(0, restored.returncode, restored.stdout + restored.stderr)
            self.assertIn("DISK_ONLY_A_RUNTIME_RECOVERED", restored.stdout)
            self.assertFalse(os.path.lexists(journal))
            self.assertEqual(context.setup_state.manifest, state[7].read_bytes())
            self.assertEqual(context.setup_state.links, {path: os.readlink(path) for path in context.setup_state.links})
            self.assertEqual(b"A oob bootstrap\n", (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
        finally:
            fixture.doCleanups()
            fixture.tearDown()

    def _private_setup_switch_fixture(self, root: Path):
        """Copy the real LOAD/SETUP scripts into a disposable HTTP root."""
        root = root.resolve()
        day0 = root / "DAY0-Prepare"
        day0.mkdir()
        for name in ("11-load.py", "01-a-setup.py"):
            source = ROOT / "DAY0-Prepare" / name
            destination = day0 / name
            shutil.copy2(source, destination)
            self.assertEqual(source.read_bytes(), destination.read_bytes())
        shutil.copytree(ROOT / "DAY0-Prepare/template", day0 / "template")
        shutil.copytree(ROOT / "tools", root / "tools")
        optimize = root / "ztp/optimize"
        optimize.mkdir(parents=True)
        shutil.copy2(ROOT / "ztp/optimize/sample_links.py", optimize / "sample_links.py")
        # Link-only WAL does not cover shared directory creation.  Model an
        # initialized directory skeleton for these link crash witnesses.
        for link_rel, _project_rel, _kind in SETUP.MAPPINGS:
            (root / "ztp" / link_rel).parent.mkdir(parents=True, exist_ok=True)
        for link_rel, _project_rel, _kind in SETUP.WORKSPACE_INPUT_MAPPINGS:
            (root / link_rel).parent.mkdir(parents=True, exist_ok=True)
        for link_path in (*SETUP.P2P_INPUT_LINKS, *SETUP.P2P_OUTPUT_LINKS):
            (root / Path(link_path).relative_to(ROOT)).parent.mkdir(
                parents=True, exist_ok=True,
            )
        for link_path, _target in SETUP.BRINGUP_OUTPUT_MAPPINGS:
            (root / Path(link_path).relative_to(ROOT)).parent.mkdir(
                parents=True, exist_ok=True,
            )
        for link_path, _target in (*SETUP.ANALYZER_INPUT_MAPPINGS,
                                   *SETUP.ANALYZER_OUTPUT_MAPPINGS):
            (root / Path(link_path).relative_to(ROOT)).parent.mkdir(
                parents=True, exist_ok=True,
            )
        for link_path, _target in SETUP._NET_CSV_LINKS:
            (root / Path(link_path).relative_to(ROOT)).parent.mkdir(
                parents=True, exist_ok=True,
            )
        for link_path, _target in SETUP._monitor_inventory_link_pairs():
            (root / Path(link_path).relative_to(ROOT)).parent.mkdir(
                parents=True, exist_ok=True,
            )
        (root / "infra").mkdir(exist_ok=True)
        (root / "image").mkdir(exist_ok=True)
        p, q = day0 / "p", day0 / "q"
        for project in (p, q):
            _write_legal_project(project, V2MlagSetupPreflightTests().pair())
        common = root / "ztp/config/cumulus/template/01-global.yaml"
        p_only = root / "ztp/config/publickey/p-only.pub"
        q_only = root / "ztp/config/publickey/operator.pub"
        for path in (common, p_only, q_only):
            path.parent.mkdir(parents=True, exist_ok=True)
        common.symlink_to(os.path.relpath(p / "01-global.yaml", common.parent))
        p_only.symlink_to(os.path.relpath(p / "operator.pub", p_only.parent))
        manifest = root / "ztp/.setup_manifest"
        manifest.write_bytes(
            (f"# setup manifest — proj: {p}\n{common}\n{p_only}\n").encode("utf-8")
        )
        before = {
            "manifest": manifest.read_bytes(),
            "common": os.readlink(common),
            "p_only": os.readlink(p_only),
        }
        hook = root / "test-hook"
        hook.mkdir()
        barrier = root / "q-first-link.observed"
        (hook / "sitecustomize.py").write_text(
            """
import os
from pathlib import Path
import signal

original_symlink = os.symlink
def kill_after_real_q_link(source, destination, *args, **kwargs):
    result = original_symlink(source, destination, *args, **kwargs)
    prefix = os.environ.get('HTTP_TEST_FIRST_Q_LINK_PREFIX')
    if (os.path.abspath(destination) == os.environ.get('HTTP_TEST_FIRST_Q_LINK')
            or (prefix and os.path.abspath(destination).startswith(prefix))):
        barrier = Path(os.environ['HTTP_TEST_Q_LINK_BARRIER'])
        with barrier.open('wb') as stream:
            stream.write((os.path.abspath(destination) + '\\n').encode()
                         if os.environ.get('HTTP_TEST_BARRIER_RECORD_DEST')
                         else b'first Q managed link written\\n')
            stream.flush()
            os.fsync(stream.fileno())
        os.kill(os.getpid(), signal.SIGKILL)
    return result
os.symlink = kill_after_real_q_link
original_replace = os.replace
def kill_after_infra_repoint(source, destination, *args, **kwargs):
    result = original_replace(source, destination, *args, **kwargs)
    if os.path.abspath(destination) == os.environ.get('HTTP_TEST_KILL_AFTER_REPLACE'):
        barrier = Path(os.environ['HTTP_TEST_Q_LINK_BARRIER'])
        with barrier.open('wb') as stream:
            stream.write(b'infra logs pointer replaced\\n')
            stream.flush()
            os.fsync(stream.fileno())
        os.kill(os.getpid(), signal.SIGKILL)
    return result
os.replace = kill_after_infra_repoint
original_link = os.link
def kill_after_infra_hardlink(source, destination, *args, **kwargs):
    result = original_link(source, destination, *args, **kwargs)
    if os.path.abspath(destination) == os.environ.get('HTTP_TEST_KILL_AFTER_HARDLINK'):
        barrier = Path(os.environ['HTTP_TEST_Q_LINK_BARRIER'])
        with barrier.open('wb') as stream:
            stream.write(b'infra logs first hardlink written\\n')
            stream.flush()
            os.fsync(stream.fileno())
        os.kill(os.getpid(), signal.SIGKILL)
    return result
os.link = kill_after_infra_hardlink
""", encoding="utf-8",
        )
        home = root / "private-home"
        home.mkdir()
        return root, p, q, common, p_only, q_only, manifest, before, hook, barrier, home

    @staticmethod
    def _run_private_load_setup_switch(
        fixture, *, prepare_journal: bool, corrupt_journal: bool = False,
        auto_prepare_journal: bool = False,
        first_link: Path | None = None, first_link_prefix: str | None = None,
        kill_after_replace: Path | None = None,
        kill_after_hardlink: Path | None = None,
        omit_common_from_wal: bool = False,
        publish_images: bool = False,
    ):
        root, p, q, common, _p_only, q_only, _manifest, _before, hook, barrier, home = fixture
        driver = """
import importlib.util
import os
import sys
from pathlib import Path
from types import SimpleNamespace

root, p, q, common, q_only, prepared, omit_common = sys.argv[1:8]
root, p, q, common, q_only = map(Path, (root, p, q, common, q_only))
spec = importlib.util.spec_from_file_location('private_setup_switch_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    if prepared == 'yes':
        expected = {common: q / '01-global.yaml', q_only: q / 'operator.pub'}
        if omit_common == 'yes':
            del expected[common]
        logs = root / 'infra/logs'
        if logs.is_symlink():
            expected[logs] = q / '99-output-infra'
        load.prepare_setup_switch_journal(
            p, q, expected_new_links=expected,
            deployment_lock_descriptor=descriptor, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'),
        )
    if prepared == 'corrupt':
        (root / 'ztp/.setup-switch-journal.json').write_bytes(b'{"schema_version":')
    if sys.argv[8] == 'images':
        os.umask(0o022)
    try:
        activation_options = (
            {'prepare_switch_wal': True,
             'runtime_backend': SimpleNamespace(name='supervisor')}
            if prepared == 'auto' else {}
        )
        load.activate_project(
            q, q / 'p2p.xlsx', strict=False, no_upgrade=sys.argv[8] != 'images',
            deployment_lock_descriptor=descriptor, host_role='workstation',
            **activation_options,
        )
    except load.LoadError as exc:
        if 'exit=-9' in str(exc):
            raise SystemExit(23)
        print('UNEXPECTED LOAD ERROR:', exc, file=sys.stderr)
        raise SystemExit(24)
    raise SystemExit(25)
finally:
    load.release_deployment_lock(descriptor)
"""
        env = os.environ.copy()
        env["HOME"] = str(home)
        # HOME is private so setup cannot inspect or generate the operator's
        # real key. Preserve the interpreter's pre-existing dependency paths
        # because macOS user-site discovery otherwise follows the new HOME.
        dependency_paths = [
            item for item in sys.path
            if item and item.endswith("site-packages") and Path(item).is_dir()
        ]
        env["PYTHONPATH"] = os.pathsep.join(
            item for item in (str(hook), env.get("PYTHONPATH", ""), *dependency_paths)
            if item
        )
        if first_link_prefix is not None:
            env.pop("HTTP_TEST_FIRST_Q_LINK", None)
            env["HTTP_TEST_FIRST_Q_LINK_PREFIX"] = first_link_prefix
            env["HTTP_TEST_BARRIER_RECORD_DEST"] = "1"
        elif kill_after_replace is not None or kill_after_hardlink is not None:
            env.pop("HTTP_TEST_FIRST_Q_LINK", None)
            env.pop("HTTP_TEST_FIRST_Q_LINK_PREFIX", None)
            env.pop("HTTP_TEST_BARRIER_RECORD_DEST", None)
        else:
            env["HTTP_TEST_FIRST_Q_LINK"] = str(first_link or common)
            env.pop("HTTP_TEST_FIRST_Q_LINK_PREFIX", None)
            env.pop("HTTP_TEST_BARRIER_RECORD_DEST", None)
        if kill_after_replace is not None:
            env["HTTP_TEST_KILL_AFTER_REPLACE"] = str(kill_after_replace)
        else:
            env.pop("HTTP_TEST_KILL_AFTER_REPLACE", None)
        if kill_after_hardlink is not None:
            env["HTTP_TEST_KILL_AFTER_HARDLINK"] = str(kill_after_hardlink)
        else:
            env.pop("HTTP_TEST_KILL_AFTER_HARDLINK", None)
        env["HTTP_TEST_Q_LINK_BARRIER"] = str(barrier)
        return subprocess.run(
            [sys.executable, "-B", "-c", driver, str(root), str(p), str(q),
             str(common), str(q_only),
             "corrupt" if corrupt_journal else "auto" if auto_prepare_journal
             else "yes" if prepare_journal else "no",
             "yes" if omit_common_from_wal else "no",
             "images" if publish_images else "no-images"],
            cwd=root, env=env, capture_output=True, text=True, check=False,
        )

    def test_real_setup_then_fresh_load_rejects_manifest_read_aba(self):
        """Cold recovery cannot authorize P writes from a read spanning A→B→A."""
        with tempfile.TemporaryDirectory(prefix="r22-manifest-read-aba-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, _before, _hook, barrier, _home = fixture
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            mixed = os.readlink(common)
            evidence = manifest.read_bytes()
            journal = root / "ztp/.setup-switch-journal.json"
            wal = journal.read_bytes()
            reader = """
import importlib.util
import os
import sys
from pathlib import Path
from types import SimpleNamespace
root, q = map(Path, sys.argv[1:3])
spec = importlib.util.spec_from_file_location('fresh_manifest_aba_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
path = root / 'ztp/.setup_manifest'
evidence = path.read_bytes()
original = path.stat()
real_read = os.read
observed = []
def read_with_aba(descriptor, count):
    result = real_read(descriptor, count)
    opened = os.fstat(descriptor)
    if not observed and (opened.st_dev, opened.st_ino) == (original.st_dev, original.st_ino):
        observed.append(True)
        path.write_bytes(b'!' + evidence[1:])
        path.write_bytes(evidence)
        os.utime(path, ns=(original.st_atime_ns, original.st_mtime_ns))
        changed = path.stat()
        assert (changed.st_ino, changed.st_size, changed.st_mtime_ns) == (original.st_ino, original.st_size, original.st_mtime_ns)
        assert changed.st_ctime_ns != original.st_ctime_ns
    return result
os.read = read_with_aba
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    )
except load.LoadError as exc:
    assert 'changed while reading' in str(exc), str(exc)
    assert observed and path.read_bytes() == evidence
    print('R22_MANIFEST_READ_ABA_REJECTED')
else:
    raise SystemExit('manifest read ABA was accepted')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", reader, str(root), str(q)], cwd=root,
                capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertIn("R22_MANIFEST_READ_ABA_REJECTED", second.stdout)
            self.assertEqual(evidence, manifest.read_bytes())
            self.assertEqual(wal, journal.read_bytes())
            self.assertEqual(mixed, os.readlink(common), "untrusted manifest authorized P writes")
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_death_fresh_load_holds_unknown_manifest_stage(self):
        for kind in ("file", "dangling-link"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(prefix="r22-manifest-orphan-cold-") as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, _q, common, p_only, q_only, manifest, _before, _hook, barrier, _home = fixture
                first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
                self.assertEqual(23, first.returncode, first.stdout + first.stderr)
                self.assertEqual(b"first Q managed link written\n", barrier.read_bytes())
                journal = root / "ztp/.setup-switch-journal.json"
                orphan = manifest.parent / f".{manifest.name}.rollback.independent-orphan"
                if kind == "file":
                    orphan.write_bytes(b"unknown prior manifest stage\n")
                else:
                    orphan.symlink_to("independent-absent-operator-object")
                before = (common.lstat().st_ino, os.readlink(common), manifest.lstat().st_ino,
                          manifest.read_bytes(), journal.lstat().st_ino, journal.read_bytes(), orphan.lstat().st_ino)
                recovery = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_manifest_orphan_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as exc:
        print('EXPECTED ORPHAN HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(descriptor)
"""
                second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                        cwd=root, capture_output=True, text=True, check=False)
                self.assertEqual(before, (common.lstat().st_ino, os.readlink(common), manifest.lstat().st_ino,
                                         manifest.read_bytes(), journal.lstat().st_ino if journal.exists() else None,
                                         journal.read_bytes() if journal.exists() else None, orphan.lstat().st_ino),
                                 "cold recovery must not mutate P/Q or retire WAL over an unknown stage")
                self.assertEqual(7, second.returncode, second.stdout + second.stderr)
                self.assertIn("manifest", second.stdout)
                if kind == "file":
                    self.assertEqual(b"unknown prior manifest stage\n", orphan.read_bytes())
                else:
                    self.assertEqual("independent-absent-operator-object", os.readlink(orphan))
                self.assertFalse(os.path.lexists(p_only))
                self.assertFalse(os.path.lexists(q_only))
                self.assertFalse((root / "service-started").exists())

    def test_real_setup_death_fresh_load_holds_unknown_link_stage(self):
        for kind in ("file", "dangling-link", "late"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(prefix="r22-link-orphan-cold-") as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, _q, common, p_only, q_only, manifest, before, *_rest = fixture
                first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
                self.assertEqual(23, first.returncode, first.stdout + first.stderr)
                journal = root / "ztp/.setup-switch-journal.json"
                wal_before = (journal.lstat().st_ino, journal.read_bytes())
                q_before = (common.lstat().st_ino, os.readlink(common), manifest.lstat().st_ino, manifest.read_bytes())
                orphan = common.parent / f".{common.name}.setup-rollback.independent-orphan"
                if kind == "file":
                    orphan.write_bytes(b"unknown prior link stage\n")
                elif kind == "dangling-link":
                    orphan.symlink_to("independent-absent-operator-object")
                orphan_before = orphan.lstat().st_ino if kind != "late" else None
                recovery = r'''
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root, orphan, kind = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3]
spec = importlib.util.spec_from_file_location('fresh_link_orphan_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
if kind == 'late':
    snapshot = load.snapshot_active_setup_state
    def late_orphan(*args, **kwargs):
        state = snapshot(*args, **kwargs)
        orphan.write_bytes(b'unknown prior link stage\n')
        (root / 'late-stage-inode').write_text(str(orphan.lstat().st_ino))
        return state
    load.snapshot_active_setup_state = late_orphan
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=fd, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as exc:
        assert 'link stage' in str(exc), str(exc)
        print('EXPECTED LINK ORPHAN HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(fd)
'''
                second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), str(orphan), kind],
                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(7, second.returncode, second.stdout + second.stderr)
                self.assertEqual(wal_before, (journal.lstat().st_ino, journal.read_bytes()))
                if kind == "late":
                    self.assertEqual(int((root / "late-stage-inode").read_text()), orphan.lstat().st_ino)
                    self.assertEqual(before["common"], os.readlink(common))
                    self.assertEqual(before["manifest"], manifest.read_bytes())
                    self.assertEqual(before["p_only"], os.readlink(p_only))
                else:
                    self.assertEqual(orphan_before, orphan.lstat().st_ino)
                    self.assertEqual(q_before, (common.lstat().st_ino, os.readlink(common), manifest.lstat().st_ino, manifest.read_bytes()))
                    self.assertFalse(os.path.lexists(p_only))
                if kind == "dangling-link":
                    self.assertEqual("independent-absent-operator-object", os.readlink(orphan))
                else:
                    self.assertEqual(b"unknown prior link stage\n", orphan.read_bytes())
                self.assertFalse(os.path.lexists(q_only))
                self.assertFalse((root / "service-started").exists())

    def test_real_load_setup_child_sigkill_then_fresh_load_restores_p(self):
        """The real setup child dies after Q's first link, before Q manifest."""
        with tempfile.TemporaryDirectory(prefix="r22-real-setup-kill-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertEqual(b"first Q managed link written\n", barrier.read_bytes())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertNotEqual(before["common"], os.readlink(common))
            self.assertFalse(os.path.lexists(p_only))
            self.assertFalse(os.path.lexists(q_only))
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertTrue(journal.is_file())

            # A fresh interpreter and new deployment lock must recover from
            # disk alone; no first-LOAD or setup-child Python objects survive.
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_setup_switch_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(os.path.lexists(journal))
            self.assertEqual(p / "01-global.yaml", common.resolve())
            self.assertFalse((root / "service-started").exists())

    def test_real_load_auto_prepares_wal_before_setup_child_first_q_link(self):
        """Real LOAD prepares P-only WAL and SETUP fills Q values before writing."""
        with tempfile.TemporaryDirectory(prefix="r22-real-load-auto-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=False, auto_prepare_journal=True,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertEqual(b"first Q managed link written\n", barrier.read_bytes())
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertTrue(journal.is_file())
            rows = json.loads(journal.read_bytes())["body"]["links"]
            row = next(row for row in rows if row["path"] == str(common))
            self.assertEqual(before["common"], row["before"])
            self.assertIsNotNone(row["after"])
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertNotEqual(before["common"], os.readlink(common))
            self.assertFalse(os.path.lexists(p_only))
            self.assertFalse(os.path.lexists(q_only))

            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location(
    'fresh_auto_wal_load', root / 'DAY0-Prepare/11-load.py',
)
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(journal.exists())
            self.assertEqual(p / "01-global.yaml", common.resolve())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_new_shared_key_directory_cold_recovery_restores_absence(self):
        """A fresh LOAD must undo SETUP's owned shared directory after child death."""
        with tempfile.TemporaryDirectory(prefix="r22-created-dir-cold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, _common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            key_root = root / "infiniband/publickey"
            key_root.parent.mkdir(exist_ok=True)
            self.assertFalse(os.path.lexists(key_root))
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=str(key_root) + os.sep,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            self.assertTrue(key_root.is_dir())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_created_dir_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)], cwd=root,
                capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(key_root), "SETUP-created shared directory survived cold rollback")
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())
            self.assertFalse((root / "service-started").exists())

    def _assert_real_setup_shared_directory_registration_drift_holds(self, drift):
        with tempfile.TemporaryDirectory(prefix="r22-shared-directory-drift-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, _common, _p_only, _q_only, _manifest, _before, hook, barrier, _home = fixture
            directory = root / "infiniband/publickey"
            directory.parent.mkdir(exist_ok=True)
            site = hook / "sitecustomize.py"
            site.write_text(site.read_text(encoding="utf-8") + """
import json
previous_shared_replace = os.replace
def drift_registered_shared_directory(source, destination, *args, **kwargs):
    result = previous_shared_replace(source, destination, *args, **kwargs)
    journal = Path(os.environ['HTTP_TEST_SHARED_WAL'])
    marker = journal.parent / 'shared-directory-drifted'
    if Path(destination) == journal and not marker.exists():
        evidence = journal.read_bytes()
        rows = json.loads(evidence)['body'].get('created_directories', [])
        selected = [row for row in rows if row['path'] == os.environ['HTTP_TEST_SHARED_TARGET']]
        if selected:
            staged = Path(selected[0]['staged'])
            drift = os.environ['HTTP_TEST_SHARED_DRIFT']
            if drift == 'owner':
                staged.rename(staged.with_name('retained-shared-stage'))
                staged.mkdir()
            elif drift == 'parent':
                parent = staged.parent
                retained = parent.with_name('retained-shared-parent')
                parent.rename(retained)
                parent.mkdir()
                (retained / staged.name).rename(staged)
            else:
                (staged / 'unknown.pub').write_bytes(b'independent unregistered key\\n')
            marker.write_bytes(evidence)
    return result
os.replace = drift_registered_shared_directory
""", encoding="utf-8")
            journal = root / "ztp/.setup-switch-journal.json"
            with mock.patch.dict(os.environ, {"HTTP_TEST_SHARED_WAL": str(journal),
                                             "HTTP_TEST_SHARED_TARGET": str(directory),
                                             "HTTP_TEST_SHARED_DRIFT": drift}):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True, first_link_prefix=str(directory) + os.sep,
                )
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertFalse(os.path.lexists(directory), "unknown shared stage was installed before HOLD")
            self.assertFalse(barrier.exists(), "shared directory drift crossed key publication")
            evidence = (journal.parent / "shared-directory-drifted").read_bytes()
            self.assertEqual(evidence, journal.read_bytes())
            row = next(row for row in json.loads(evidence)["body"]["created_directories"]
                       if row["path"] == str(directory))
            staged = Path(row["staged"])
            self.assertTrue(staged.is_dir())
            if drift == "owner":
                self.assertEqual(row["ino"], staged.with_name("retained-shared-stage").stat().st_ino)
                self.assertNotEqual(row["ino"], staged.stat().st_ino)
            else:
                self.assertEqual(row["ino"], staged.stat().st_ino)
            if drift == "parent":
                self.assertNotEqual(row["parent"]["ino"], directory.parent.stat().st_ino)
                self.assertEqual(row["parent"]["ino"], directory.parent.with_name("retained-shared-parent").stat().st_ino)
            if drift == "contents":
                self.assertEqual(b"independent unregistered key\n", (staged / "unknown.pub").read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_shared_directory_registered_owner_drift_blocks_install(self):
        self._assert_real_setup_shared_directory_registration_drift_holds("owner")

    def test_real_setup_shared_directory_registered_parent_drift_blocks_install(self):
        self._assert_real_setup_shared_directory_registration_drift_holds("parent")

    def test_real_setup_shared_directory_registered_contents_drift_blocks_install(self):
        self._assert_real_setup_shared_directory_registration_drift_holds("contents")

    def test_real_setup_created_directory_recovery_parent_drift_blocks_rmdir(self):
        with tempfile.TemporaryDirectory(prefix="r22-recovery-created-parent-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, _p_only, _q_only, _manifest, _before, _hook, barrier, _home = fixture
            parent = root / "infiniband"
            parent.mkdir(exist_ok=True)
            directory = parent / "publickey"
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=str(directory) + os.sep,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            original = directory.stat()
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root, common, directory = map(Path, sys.argv[1:4])
spec = importlib.util.spec_from_file_location('fresh_created_parent_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
restore = load._atomic_restore_setup_link
injected = []
def rebind_parent_after_known_p_write(path, raw_target, **guard):
    restore(path, raw_target, **guard)
    if path == common and not injected:
        parent = directory.parent
        retained = parent.with_name('retained-recovery-parent')
        parent.rename(retained)
        parent.mkdir()
        (retained / directory.name).rename(directory)
        injected.append(True)
load._atomic_restore_setup_link = rebind_parent_after_known_p_write
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        assert injected == [True], 'parent drift injection did not run'
        print('CREATED_DIRECTORY_PARENT_DRIFT_HOLD=' + str(exc))
    else:
        raise AssertionError('created-directory recovery accepted a rebound parent')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root), str(common), str(directory)],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertIn("CREATED_DIRECTORY_PARENT_DRIFT_HOLD=", second.stdout)
            self.assertTrue(directory.is_dir(), "fresh recovery deleted the moved inode before parent HOLD")
            self.assertEqual((original.st_dev, original.st_ino),
                             (directory.stat().st_dev, directory.stat().st_ino))
            row = next(row for row in json.loads(evidence)["body"]["created_directories"]
                       if row["path"] == str(directory))
            self.assertEqual(row["parent"]["ino"], parent.with_name("retained-recovery-parent").stat().st_ino)
            self.assertNotEqual(row["parent"]["ino"], parent.stat().st_ino)
            self.assertEqual(evidence, journal.read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_legacy_comparison_transfer_death_cold_restores_reports(self):
        self._exercise_real_comparison_transfer_death("pointer")

    def test_real_setup_publickey_directory_death_preserves_public_mode_and_cold_absence(self):
        self._exercise_publickey_directory_death(stage_death=False)

    def test_real_setup_publickey_unregistered_stage_death_cold_holds(self):
        self._exercise_publickey_directory_death(stage_death=True)

    def _exercise_publickey_directory_death(self, *, stage_death):
        with tempfile.TemporaryDirectory(prefix="r22-real-publickey-directory-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, before, hook, barrier, _home = fixture
            parent = root / "infiniband"
            parent.mkdir(exist_ok=True)
            directory = parent / "publickey"
            self.assertFalse(os.path.lexists(directory))
            hook_file = hook / "sitecustomize.py"
            hook_file.write_text(hook_file.read_text() + f"""
os.umask(0o022)
publickey_sync = os.fsync
def kill_before_publickey_stage_row(fd):
    stages = list(Path({str(parent)!r}).glob('.publickey.setup-*'))
    if {stage_death!r} and any(stage.is_dir() for stage in stages):
        Path(os.environ['HTTP_TEST_Q_LINK_BARRIER']).write_bytes(b'unregistered publickey stage observed\\n')
        os.kill(os.getpid(), signal.SIGKILL)
    return publickey_sync(fd)
os.fsync = kill_before_publickey_stage_row
""")
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link=directory / "operator.pub",
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            held_manifest, held_common = manifest.read_bytes(), os.readlink(common)
            if stage_death:
                stages = list(parent.glob(".publickey.setup-*"))
                self.assertEqual(1, len(stages))
                owner = stages[0].stat()
                self.assertFalse(os.path.lexists(directory))
                self.assertNotIn(str(directory), {row["path"] for row in
                    json.loads(evidence)["body"].get("created_directories", [])})
            else:
                self.assertEqual(0o755, stat.S_IMODE(directory.stat().st_mode))
                self.assertEqual((q / "operator.pub").read_bytes(), (directory / "operator.pub").read_bytes())
            recovery = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_publickey_directory_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
fd = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=fd, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor')) == 'restored-p'
finally:
    load.release_deployment_lock(fd)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            if stage_death:
                self.assertEqual(1, second.returncode, second.stdout + second.stderr)
                self.assertIn("unregistered", second.stderr)
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(held_manifest, manifest.read_bytes())
                self.assertEqual(held_common, os.readlink(common))
                self.assertEqual(owner.st_ino, stages[0].stat().st_ino)
            else:
                self.assertEqual(0, second.returncode, second.stdout + second.stderr)
                self.assertEqual(before["manifest"], manifest.read_bytes())
                self.assertEqual(before["common"], os.readlink(common))
                self.assertFalse(os.path.lexists(journal))
            self.assertFalse(os.path.lexists(directory))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_shared_image_directory_death_cold_restores_absence(self):
        self._exercise_shared_image_directory_death(recovery_death=False)

    def test_real_setup_shared_image_nested_directory_recovery_death_reenters(self):
        self._exercise_shared_image_directory_death(recovery_death=True)

    def test_real_setup_shared_nvos_image_directory_death_cold_restores_absence(self):
        self._exercise_shared_image_directory_death(
            recovery_death=False, image_name="nvos-amd64-25.2.1.bin", platform="nvos",
        )

    def test_real_setup_shared_image_unregistered_stage_death_cold_holds(self):
        self._exercise_shared_image_directory_death(recovery_death=False, stage_death=True)

    def test_real_setup_shared_image_absent_parent_reentry_replacement_holds(self):
        self._exercise_shared_image_directory_death(recovery_death=True, replace_parent=True)

    def _exercise_shared_image_directory_death(
            self, *, recovery_death, stage_death=False, replace_parent=False,
            image_name="cumulus-linux-5.18.1-mlx-amd64.bin", platform="cumulus"):
        with tempfile.TemporaryDirectory(prefix="r22-real-image-directories-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            image = root / "image" / image_name
            payload = b"independent shared image payload\n"
            image.write_bytes(payload)
            directory = root / "ztp/image" / platform
            self.assertFalse(os.path.lexists(directory.parent))
            if stage_death:
                hook = fixture[8] / "sitecustomize.py"
                hook.write_text(hook.read_text() + "\n" + f"""
image_sync = os.fsync
def kill_before_image_stage_row(fd):
    stages = list(Path({str(root / 'ztp')!r}).glob('.image.setup-*'))
    if any(stage.is_dir() for stage in stages):
        Path(os.environ['HTTP_TEST_Q_LINK_BARRIER']).write_bytes(b'unregistered image stage observed\\n')
        os.kill(os.getpid(), signal.SIGKILL)
    return image_sync(fd)
os.fsync = kill_before_image_stage_row
""")
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, publish_images=True,
                first_link=directory / image.name,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            if not stage_death:
                self.assertTrue((directory / image.name).is_symlink())
                for path in (directory.parent, directory):
                    self.assertEqual(0o755, stat.S_IMODE(path.stat().st_mode))
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            held_manifest, held_common = manifest.read_bytes(), os.readlink(common)
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_shared_image_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
if len(sys.argv) > 2:
    import os, signal
    sync = load._fsync_parent_directory
    def kill_before_outer_parent_sync(path):
        if Path(path) == root / 'ztp' and not os.path.lexists(root / 'ztp/image'):
            (root / 'image-recovery-death.observed').write_bytes(b'image parent removal observed\\n')
            os.kill(os.getpid(), signal.SIGKILL)
        return sync(path)
    load._fsync_parent_directory = kill_before_outer_parent_sync
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
            if recovery_death:
                interrupted = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), "kill"],
                                             cwd=root, capture_output=True, text=True, check=False, timeout=30)
                self.assertEqual(-signal.SIGKILL, interrupted.returncode, interrupted.stdout + interrupted.stderr)
                self.assertTrue((root / "image-recovery-death.observed").is_file())
                self.assertEqual(evidence, journal.read_bytes())
                self.assertFalse(os.path.lexists(directory.parent))
            if replace_parent:
                directory.parent.mkdir()
                (directory.parent / "unknown-operator.txt").write_bytes(b"unrelated replacement parent\n")
                replacement = directory.parent.stat()
                held_manifest, held_common = manifest.read_bytes(), os.readlink(common)
            if stage_death:
                stages = list((root / "ztp").glob(".image.setup-*"))
                self.assertEqual(1, len(stages))
                stage_owner = stages[0].stat()
                self.assertEqual([], json.loads(evidence)["body"].get("created_directories", []))
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            if stage_death or replace_parent:
                self.assertEqual(1, second.returncode, second.stdout + second.stderr)
                self.assertIn("LoadError", second.stderr)
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(held_manifest, manifest.read_bytes())
                self.assertEqual(held_common, os.readlink(common))
                self.assertEqual(payload, image.read_bytes())
                self.assertFalse((root / "service-started").exists())
                if stage_death:
                    self.assertIn("unregistered shared image stage", second.stderr)
                    self.assertEqual(stage_owner.st_ino, stages[0].stat().st_ino)
                    self.assertFalse(os.path.lexists(directory.parent))
                else:
                    self.assertEqual(replacement.st_ino, directory.parent.stat().st_ino)
                    self.assertEqual(b"unrelated replacement parent\n",
                                     (directory.parent / "unknown-operator.txt").read_bytes())
                return
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(directory))
            self.assertFalse(os.path.lexists(directory.parent))
            self.assertEqual(payload, image.read_bytes())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_legacy_comparison_transfer_rename_death_cold_restores_reports(self):
        self._exercise_real_comparison_transfer_death("rename")

    def test_real_setup_comparison_transfer_changed_reports_cold_holds(self):
        self._exercise_real_comparison_transfer_death("pointer", mutation="report")

    def test_real_setup_comparison_transfer_unknown_target_cold_holds(self):
        self._exercise_real_comparison_transfer_death("pointer", mutation="target")

    def _exercise_real_comparison_transfer_death(self, boundary, mutation=None):
        with tempfile.TemporaryDirectory(prefix="r22-real-comparison-transfer-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            source = root / "ztp/optimize/q-sample/comparison"
            source.mkdir(parents=True)
            report = b"independent original comparison report\n"
            (source / "operator-report.csv").write_bytes(report)
            original = source.stat()
            (q / "99-output-ztp").mkdir(exist_ok=True)
            destination = q / "99-output-ztp/optimize"
            if boundary == "rename":
                hook = fixture[8] / "sitecustomize.py"
                hook.write_text(hook.read_text() + "\n" + f"""
original_fsync = os.fsync
def kill_after_comparison_rename(fd):
    source = Path({str(source)!r})
    destination = Path({str(destination)!r})
    if (not os.path.lexists(source) and destination.is_dir()
            and destination.stat().st_ino == {original.st_ino}):
        Path(os.environ['HTTP_TEST_Q_LINK_BARRIER']).write_bytes(b'comparison rename observed\\n')
        os.kill(os.getpid(), signal.SIGKILL)
    return original_fsync(fd)
os.fsync = kill_after_comparison_rename
""")
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link=root / "never-created-link",
                kill_after_replace=source if boundary == "pointer" else root / "never-replaced",
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            if boundary == "pointer":
                self.assertTrue(source.is_symlink())
            else:
                self.assertFalse(os.path.lexists(source))
            self.assertEqual(original.st_ino, destination.stat().st_ino)
            self.assertEqual(report, (destination / "operator-report.csv").read_bytes())
            journal = root / "ztp/.setup-switch-journal.json"
            evidence, q_common = journal.read_bytes(), os.readlink(common)
            if mutation == "report":
                (destination / "operator-report.csv").write_bytes(b"unknown changed report\n")
            elif mutation == "target":
                retained = destination.with_name("retained-comparison")
                destination.rename(retained)
                destination.mkdir()
                (destination / "unknown.csv").write_bytes(b"unknown operator data\n")
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_comparison_transfer_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            if mutation is not None:
                self.assertNotEqual(0, second.returncode, second.stdout + second.stderr)
                self.assertIn("comparison transfer is unknown", second.stderr)
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(q_common, os.readlink(common), "HOLD changed a previously written Q link")
                self.assertTrue(source.is_symlink())
                if mutation == "report":
                    self.assertEqual(b"unknown changed report\n", (destination / "operator-report.csv").read_bytes())
                    self.assertEqual(original.st_ino, destination.stat().st_ino)
                else:
                    self.assertEqual(original.st_ino, retained.stat().st_ino)
                    self.assertEqual(report, (retained / "operator-report.csv").read_bytes())
                    self.assertEqual(b"unknown operator data\n", (destination / "unknown.csv").read_bytes())
                self.assertFalse((root / "service-started").exists())
                return
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(source.is_symlink())
            self.assertEqual(original.st_ino, source.stat().st_ino)
            self.assertEqual(report, (source / "operator-report.csv").read_bytes())
            self.assertFalse(os.path.lexists(destination))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_new_sample_directory_cold_recovery_restores_absence(self):
        """Sample's empty directory itself belongs to the crash rollback write set."""
        with tempfile.TemporaryDirectory(prefix="r22-sample-dir-cold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            sample = root / "ztp/optimize/q-sample"
            self.assertFalse(os.path.lexists(sample))
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=str(sample) + os.sep,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            self.assertTrue(sample.is_dir())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_sample_directory_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)], cwd=root,
                capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(sample), "sample mkdir survived durable cold rollback")
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_comparison_directory_death_cold_restores_absence(self):
        """Q-local comparison mkdir is a persistent write, not a disposable alias."""
        self._assert_real_setup_q_comparison_directory_recovery(replace_target=False)

    def test_real_setup_q_comparison_directory_replacement_cold_holds_unchanged(self):
        self._assert_real_setup_q_comparison_directory_recovery(replace_target=True)

    def test_real_setup_q_comparison_death_before_wal_row_holds_orphan(self):
        with tempfile.TemporaryDirectory(prefix="r22-q-stage-before-row-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, _before, hook, barrier, _home = fixture
            output = q / "99-output-ztp"
            output.mkdir()
            site = hook / "sitecustomize.py"
            site.write_text(site.read_text(encoding="utf-8") + """
real_mkdir = Path.mkdir
def interrupted_q_stage_mkdir(path, *args, **kwargs):
    result = real_mkdir(path, *args, **kwargs)
    if path.parent == Path(os.environ['HTTP_TEST_Q_OUTPUT']) and path.name.startswith('.optimize.setup-'):
        Path(os.environ['HTTP_TEST_Q_LINK_BARRIER']).write_bytes(b'Q stage before WAL row\\n')
        os.kill(os.getpid(), signal.SIGKILL)
    return result
Path.mkdir = interrupted_q_stage_mkdir
""", encoding="utf-8")
            with mock.patch.dict(os.environ, {"HTTP_TEST_Q_OUTPUT": str(output)}):
                first = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True, first_link=root / "never-created-link",
                )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertEqual(b"Q stage before WAL row\n", barrier.read_bytes())
            stages = list(output.glob(".optimize.setup-*"))
            self.assertEqual(1, len(stages))
            identity = stages[0].stat().st_ino
            journal = root / "ztp/.setup-switch-journal.json"
            evidence, manifest_before, mixed = journal.read_bytes(), manifest.read_bytes(), os.readlink(common)
            self.assertEqual([], json.loads(evidence)["body"].get("q_created_directories", []))
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_q_stage_orphan', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        print('ORPHAN_HOLD:', exc)
        raise SystemExit(29)
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, timeout=30)
            self.assertEqual(29, second.returncode, second.stdout + second.stderr)
            self.assertIn("unregistered Q comparison stage", second.stdout)
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(manifest_before, manifest.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            self.assertEqual(identity, stages[0].stat().st_ino)
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_comparison_rmdir_death_reentry_syncs_before_wal_retirement(self):
        with tempfile.TemporaryDirectory(prefix="r22-q-dir-reentry-sync-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, _common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            (q / "99-output-ztp").mkdir()
            target = q / "99-output-ztp/optimize"
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=str(root / "ztp/optimize/q-sample") + os.sep,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            recovery = """
import importlib.util
import os
import signal
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
target = root / 'DAY0-Prepare/q/99-output-ztp/optimize'
journal = root / 'ztp/.setup-switch-journal.json'
spec = importlib.util.spec_from_file_location('fresh_q_rmdir_sync', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
if sys.argv[2] == 'kill':
    real_rmdir = Path.rmdir
    def interrupted_rmdir(path):
        real_rmdir(path)
        if path == target:
            (root / 'q-rmdir-observed').write_bytes(b'Q directory removed\\n')
            os.kill(os.getpid(), signal.SIGKILL)
    Path.rmdir = interrupted_rmdir
else:
    real_sync = load._fsync_parent_directory
    def observed_sync(path):
        if Path(path) == target.parent:
            assert journal.is_file(), 'Q absence sync happened after WAL retirement'
            assert not os.path.lexists(target)
            real_sync(path)
            (root / 'q-absence-synced').write_bytes(b'parent synced before retirement\\n')
        else:
            real_sync(path)
    load._fsync_parent_directory = observed_sync
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        journal, deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), "kill"],
                                    cwd=root, capture_output=True, text=True, timeout=30)
            self.assertEqual(-9, second.returncode, second.stdout + second.stderr)
            self.assertEqual(b"Q directory removed\n", (root / "q-rmdir-observed").read_bytes())
            self.assertFalse(os.path.lexists(target))
            self.assertEqual(evidence, journal.read_bytes())
            third = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), "resume"],
                                   cwd=root, capture_output=True, text=True, timeout=30)
            self.assertEqual(0, third.returncode, third.stdout + third.stderr)
            self.assertTrue((root / "q-absence-synced").is_file(),
                            "cold reentry retired WAL without syncing Q's known absence")
            self.assertEqual(b"parent synced before retirement\n", (root / "q-absence-synced").read_bytes())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertFalse(os.path.lexists(journal))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_comparison_registered_stage_rebind_blocks_install(self):
        with tempfile.TemporaryDirectory(prefix="r22-q-stage-registration-rebind-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, _common, _p_only, _q_only, _manifest, _before, hook, barrier, _home = fixture
            output = q / "99-output-ztp"
            output.mkdir()
            site = hook / "sitecustomize.py"
            site.write_text(site.read_text(encoding="utf-8") + """
import json
previous_q_replace = os.replace
def replace_registered_q_stage(source, destination, *args, **kwargs):
    result = previous_q_replace(source, destination, *args, **kwargs)
    journal = Path(os.environ['HTTP_TEST_Q_STAGE_WAL'])
    marker = journal.parent / 'q-stage-rebound'
    if Path(destination) == journal and not marker.exists():
        evidence = journal.read_bytes()
        rows = json.loads(evidence)['body'].get('q_created_directories', [])
        if rows:
            staged = Path(rows[0]['staged'])
            staged.rename(staged.with_name('original-owned-stage'))
            staged.mkdir()
            marker.write_bytes(evidence)
    return result
os.replace = replace_registered_q_stage
""", encoding="utf-8")
            journal = root / "ztp/.setup-switch-journal.json"
            with mock.patch.dict(os.environ, {"HTTP_TEST_Q_STAGE_WAL": str(journal)}):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True,
                    first_link_prefix=str(root / "ztp/optimize/q-sample") + os.sep,
                )
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertFalse(os.path.lexists(output / "optimize"),
                             "unrelated staged inode was installed before the later HOLD")
            self.assertRegex(result.stdout + result.stderr, "Q comparison.*owner changed")
            self.assertFalse(barrier.exists(), "unrelated stage escaped into shared sample publication")
            self.assertFalse(os.path.lexists(output / "optimize"))
            self.assertTrue((output / "original-owned-stage").is_dir())
            evidence = (journal.parent / "q-stage-rebound").read_bytes()
            self.assertEqual(evidence, journal.read_bytes())
            staged = Path(json.loads(evidence)["body"]["q_created_directories"][0]["staged"])
            self.assertTrue(staged.is_dir(), "unknown staged replacement was deleted")
            self.assertNotEqual(json.loads(evidence)["body"]["q_created_directories"][0]["ino"],
                                staged.stat().st_ino)
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_comparison_parent_rebind_with_same_stage_blocks_install(self):
        with tempfile.TemporaryDirectory(prefix="r22-q-parent-registration-rebind-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, _common, _p_only, _q_only, _manifest, _before, hook, barrier, _home = fixture
            output = q / "99-output-ztp"
            output.mkdir()
            site = hook / "sitecustomize.py"
            site.write_text(site.read_text(encoding="utf-8") + """
import json
previous_q_parent_replace = os.replace
def replace_registered_q_parent(source, destination, *args, **kwargs):
    result = previous_q_parent_replace(source, destination, *args, **kwargs)
    journal = Path(os.environ['HTTP_TEST_Q_STAGE_WAL'])
    marker = journal.parent / 'q-parent-rebound'
    if Path(destination) == journal and not marker.exists():
        evidence = journal.read_bytes()
        rows = json.loads(evidence)['body'].get('q_created_directories', [])
        if rows:
            staged = Path(rows[0]['staged'])
            output = staged.parent
            retained = output.with_name('retained-output-parent')
            output.rename(retained)
            output.mkdir()
            (retained / staged.name).rename(staged)
            marker.write_bytes(evidence)
    return result
os.replace = replace_registered_q_parent
""", encoding="utf-8")
            journal = root / "ztp/.setup-switch-journal.json"
            with mock.patch.dict(os.environ, {"HTTP_TEST_Q_STAGE_WAL": str(journal)}):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True,
                    first_link_prefix=str(root / "ztp/optimize/q-sample") + os.sep,
                )
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertFalse(os.path.lexists(output / "optimize"),
                             "WAL-owned staged inode was installed into an unowned parent")
            self.assertRegex(result.stdout + result.stderr, "Q comparison.*parent.*changed")
            self.assertFalse(barrier.exists(), "parent rebind reached sample publication")
            evidence = (journal.parent / "q-parent-rebound").read_bytes()
            self.assertEqual(evidence, journal.read_bytes())
            row = json.loads(evidence)["body"]["q_created_directories"][0]
            staged = Path(row["staged"])
            self.assertEqual((row["dev"], row["ino"]), (staged.stat().st_dev, staged.stat().st_ino))
            self.assertNotEqual(row["parent"]["ino"], output.stat().st_ino)
            self.assertEqual(row["parent"]["ino"], output.with_name("retained-output-parent").stat().st_ino)
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_comparison_registered_stage_unknown_child_blocks_install(self):
        with tempfile.TemporaryDirectory(prefix="r22-q-stage-registration-child-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, _common, _p_only, _q_only, _manifest, _before, hook, barrier, _home = fixture
            output = q / "99-output-ztp"
            output.mkdir()
            site = hook / "sitecustomize.py"
            site.write_text(site.read_text(encoding="utf-8") + """
import json
previous_q_child_replace = os.replace
def populate_registered_q_stage(source, destination, *args, **kwargs):
    result = previous_q_child_replace(source, destination, *args, **kwargs)
    journal = Path(os.environ['HTTP_TEST_Q_STAGE_WAL'])
    marker = journal.parent / 'q-stage-populated'
    if Path(destination) == journal and not marker.exists():
        evidence = journal.read_bytes()
        rows = json.loads(evidence)['body'].get('q_created_directories', [])
        if rows:
            staged = Path(rows[0]['staged'])
            (staged / 'unknown-report.csv').write_bytes(b'independent unregistered report\\n')
            marker.write_bytes(evidence)
    return result
os.replace = populate_registered_q_stage
""", encoding="utf-8")
            journal = root / "ztp/.setup-switch-journal.json"
            with mock.patch.dict(os.environ, {"HTTP_TEST_Q_STAGE_WAL": str(journal)}):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True,
                    first_link_prefix=str(root / "ztp/optimize/q-sample") + os.sep,
                )
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertFalse(os.path.lexists(output / "optimize"),
                             "unregistered report moved into Q before a later HOLD")
            self.assertRegex(result.stdout + result.stderr, "Q comparison.*unregistered children")
            self.assertFalse(barrier.exists(), "unknown stage content reached sample publication")
            evidence = (journal.parent / "q-stage-populated").read_bytes()
            self.assertEqual(evidence, journal.read_bytes())
            row = json.loads(evidence)["body"]["q_created_directories"][0]
            staged = Path(row["staged"])
            self.assertEqual((row["dev"], row["ino"]), (staged.stat().st_dev, staged.stat().st_ino))
            self.assertEqual(b"independent unregistered report\n", (staged / "unknown-report.csv").read_bytes())
            self.assertFalse((root / "service-started").exists())

    def _assert_real_setup_q_comparison_directory_recovery(self, *, replace_target):
        with tempfile.TemporaryDirectory(prefix="r22-q-comparison-dir-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            (q / "99-output-ztp").mkdir()
            target = q / "99-output-ztp/optimize"
            sample = root / "ztp/optimize/q-sample"
            self.assertFalse(os.path.lexists(target))
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=str(sample) + os.sep,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            self.assertTrue(target.is_dir())
            journal = root / "ztp/.setup-switch-journal.json"
            evidence, mixed = journal.read_bytes(), os.readlink(common)
            manifest_before = manifest.read_bytes()
            if replace_target:
                target.rename(target.with_name("original-owned-optimize"))
                target.mkdir()
                (target / "unknown-report.csv").write_bytes(b"unrelated owner\n")
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_q_comparison_directory', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        assert load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        ) == 'restored-p'
    except load.LoadError as exc:
        print('Q_DIRECTORY_HOLD:', exc)
        raise SystemExit(29)
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)], cwd=root,
                capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(29 if replace_target else 0, second.returncode,
                             second.stdout + second.stderr)
            if replace_target:
                self.assertIn("created directory", second.stdout)
                self.assertEqual(mixed, os.readlink(common))
                self.assertEqual(manifest_before, manifest.read_bytes())
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(b"unrelated owner\n", (target / "unknown-report.csv").read_bytes())
                self.assertTrue(sample.is_dir())
            else:
                self.assertFalse(os.path.lexists(target), "cold recovery lost Q's original absence")
                self.assertFalse(os.path.lexists(sample))
                self.assertEqual(before["manifest"], manifest.read_bytes())
                self.assertFalse(os.path.lexists(journal))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_inprocess_error_retains_unknown_replacement_and_wal(self):
        with tempfile.TemporaryDirectory(prefix="r22-setup-rollback-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, _p_only, q_only, manifest, before, hook, barrier, _home = fixture
            original = os.path.relpath(p / "operator.pub", q_only.parent)
            hook_file = hook / "sitecustomize.py"
            hook_file.write_text(hook_file.read_text(encoding="utf-8") + """
import hashlib
original_for_rollback = os.symlink
def replace_at_setup_write(source, destination, *args, **kwargs):
    result = original_for_rollback(source, destination, *args, **kwargs)
    if os.path.abspath(destination) == os.environ['HTTP_TEST_UNKNOWN_Q_LINK']:
        os.unlink(destination)
        # Bypass the failure/kill hooks while installing the unrelated object.
        original_symlink(os.environ['HTTP_TEST_UNKNOWN_Q_TARGET'], destination)
        wal = Path(os.environ['HTTP_TEST_WAL_NAME'])
        Path(os.environ['HTTP_TEST_Q_LINK_BARRIER']).write_text(hashlib.sha256(wal.read_bytes()).hexdigest())
        raise ValueError('injected setup failure after unknown replacement')
    return result
os.symlink = replace_at_setup_write
""", encoding="utf-8")
            with mock.patch.dict(os.environ, {
                "HTTP_TEST_UNKNOWN_Q_LINK": str(q_only),
                "HTTP_TEST_UNKNOWN_Q_TARGET": original,
                "HTTP_TEST_WAL_NAME": str(root / "ztp/.setup-switch-journal.json"),
            }):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True, first_link=root / "never-created-link",
                )
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertTrue(barrier.exists())
            self.assertEqual(original, os.readlink(q_only), "real SETUP rollback deleted the unknown link")
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertEqual(barrier.read_text(), hashlib.sha256(journal.read_bytes()).hexdigest())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(q / "01-global.yaml", common.resolve(), "unknown state did not fence P writes")
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_failure_context_restores_durable_sample_before_runtime_phase(self):
        """Parent's failure API must use disk WAL, not only its detached P links."""
        with tempfile.TemporaryDirectory(prefix="r22-parent-wal-restore-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            sample = root / "ztp/optimize/q-sample"
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=str(sample) + os.sep,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            recovery = """
import importlib.util
import json
import sys
from pathlib import Path
from types import SimpleNamespace
root, p, q, common, p_only = map(Path, sys.argv[1:6])
before = json.loads(sys.argv[6])
spec = importlib.util.spec_from_file_location('fresh_parent_failure_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
wal = root / 'ztp/.setup-switch-journal.json'
body = json.loads(wal.read_bytes())['body']
state = load.ActiveSetupState(p, before['manifest'].encode(), body['manifest_mode'],
                             {common: before['common'], p_only: before['p_only']})
context = load.CrossProjectRollbackContext(p, 'generations/aaaaaaaaaaaaaaaaaaaa',
                                          p / '99-output-ztp/rollback-receipts/aaaaaaaaaaaaaaaaaaaa', state)
def runtime_phase(*args, **kwargs):
    assert not (root / 'ztp/optimize/q-sample').exists()
    assert wal.is_file(), 'A runtime recovery must still hold the original WAL'
    assert (root / 'ztp/.setup_manifest').read_bytes() == state.manifest
    assert load.snapshot_active_setup_state(expected_project=p).links == state.links
    (root / 'runtime-phase-reached').write_bytes(b'after-durable-P-restore\\n')
    if len(sys.argv) > 7:
        import os
        import signal
        os.kill(os.getpid(), signal.SIGKILL)
load.restore_runtime_rollback_receipt = runtime_phase
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    load.restore_cross_project_rollback_context(
        context, current_project=q, deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    )
finally:
    load.release_deployment_lock(descriptor)
"""
            encoded = json.dumps({**before, "manifest": before["manifest"].decode()})
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root), str(p), str(q),
                 str(common), str(p_only), encoded, "kill-at-runtime"],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(-9, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(sample))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(b"after-durable-P-restore\n", (root / "runtime-phase-reached").read_bytes())
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertTrue(journal.is_file(), "death before A verification lost recovery authority")
            self.assertEqual(evidence, journal.read_bytes())
            third = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root), str(p), str(q),
                 str(common), str(p_only), encoded],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, third.returncode, third.stdout + third.stderr)
            self.assertFalse(os.path.lexists(journal))
            self.assertTrue(evidence, "failure WAL was empty")
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_manifest_then_cold_restore_syncs_p_mode_and_retains_failed_wal(self):
        """Real Q publication must restore durable P metadata before WAL retirement."""
        with tempfile.TemporaryDirectory(prefix="r22-p-manifest-mode-sync-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            manifest.chmod(0o640)
            for relative in (
                "tools/lldp-analyze-tool/99-output-p2p",
                "tools/lldp-analyze-tool/99-output-monitor",
                "tools/ibdiagnet-analyze-tool/99-output-p2p",
            ):
                leaf = root / relative
                if leaf.is_dir() and not leaf.is_symlink():
                    shutil.rmtree(leaf)
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=False, auto_prepare_journal=True,
                kill_after_replace=manifest,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            q_manifest = manifest.read_bytes()
            self.assertNotEqual(before["manifest"], q_manifest)
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            recovery = """
import importlib.util, os, stat, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('cold_p_manifest_sync_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
real_create, real_sync, real_replace = load.tempfile.mkstemp, os.fsync, os.replace
held = {}
events = []
def create(*args, **kwargs):
    descriptor, name = real_create(*args, **kwargs)
    if kwargs.get('prefix', '').startswith('..setup_manifest.rollback.'):
        held['fd'], held['path'] = descriptor, Path(name)
    return descriptor, name
def sync(descriptor):
    metadata = os.fstat(descriptor)
    if descriptor == held.get('fd') and stat.S_ISREG(metadata.st_mode):
        assert stat.S_IMODE(metadata.st_mode) == 0o640, 'P final mode was not synchronized'
        assert held['path'].read_bytes().startswith(b'# setup manifest')
        if sys.argv[2] == 'fault':
            raise OSError('independent P manifest file sync fault')
        events.append('file-synced')
    elif events == ['file-synced', 'replaced'] and stat.S_ISDIR(metadata.st_mode):
        events.append('parent-synced')
    return real_sync(descriptor)
def replace(source, destination, *args, **kwargs):
    if Path(destination) == root / 'ztp/.setup_manifest':
        assert events == ['file-synced'], 'P manifest replace preceded final-mode file sync'
        events.append('replaced')
    return real_replace(source, destination, *args, **kwargs)
load.tempfile.mkstemp, os.fsync, os.replace = create, sync, replace
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        assert load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        ) == 'restored-p'
    except OSError as error:
        assert sys.argv[2] == 'fault' and 'manifest file sync fault' in str(error), str(error)
        raise SystemExit(7)
    assert events == ['file-synced', 'replaced', 'parent-synced'], events
    print('P_MANIFEST_FINAL_MODE_SYNC_VERIFIED')
finally:
    load.release_deployment_lock(descriptor)
"""
            argv = [sys.executable, "-B", "-c", recovery, str(root)]
            failed = subprocess.run(argv + ["fault"], cwd=root, capture_output=True,
                                    text=True, check=False, timeout=30)
            self.assertEqual(7, failed.returncode, failed.stdout + failed.stderr)
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(q_manifest, manifest.read_bytes())
            restored = subprocess.run(argv + ["sync"], cwd=root, capture_output=True,
                                      text=True, check=False, timeout=30)
            self.assertEqual(0, restored.returncode, restored.stdout + restored.stderr)
            self.assertIn("P_MANIFEST_FINAL_MODE_SYNC_VERIFIED", restored.stdout)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(0o640, stat.S_IMODE(manifest.stat().st_mode))
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(os.path.lexists(journal))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_runtime_parent_sync_failure_retains_wal_until_fresh_retry(self):
        """Real cold P recovery must not retire WAL after an unsynced runtime write."""
        with tempfile.TemporaryDirectory(prefix="r22-runtime-parent-sync-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            recovery = """
import hashlib, importlib.util, json, os, stat, sys
from pathlib import Path
from types import SimpleNamespace
root, p, q, common, p_only = map(Path, sys.argv[1:6])
before = json.loads(sys.argv[6])
spec = importlib.util.spec_from_file_location('fresh_runtime_sync_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
wal = root / 'ztp/.setup-switch-journal.json'
body = json.loads(wal.read_bytes())['body']
state = load.ActiveSetupState(p, before['manifest'].encode(), body['manifest_mode'],
                             {common: before['common'], p_only: before['p_only']})
context = load.CrossProjectRollbackContext(p, 'generations/aaaaaaaaaaaaaaaaaaaa',
                                          p / '99-output-ztp/rollback-receipts/aaaaaaaaaaaaaaaaaaaa', state)
target = root / 'ztp/runtime-durability-witness'
payload = b'independent A runtime witness\\n'
replaced = False
real_replace, real_fsync = os.replace, os.fsync
def replace(source, destination):
    global replaced
    result = real_replace(source, destination)
    if Path(destination) == target:
        replaced = True
    return result
def sync(fd):
    metadata = os.fstat(fd)
    parent = target.parent.stat()
    if replaced and (metadata.st_dev, metadata.st_ino) == (parent.st_dev, parent.st_ino):
        assert wal.is_file(), 'WAL retired before runtime parent synchronization'
        if sys.argv[7] == 'fault':
            raise OSError('runtime parent fsync fault')
        (root / 'runtime-parent-synced').write_bytes(b'parent sync observed\\n')
    return real_fsync(fd)
def runtime_phase(*args, **kwargs):
    assert wal.is_file()
    os.replace, os.fsync = replace, sync
    try:
        load._restore_receipt_file(target, {'state': 'file', 'sha256': hashlib.sha256(payload).hexdigest(), 'mode': 0o600}, payload)
    finally:
        os.replace, os.fsync = real_replace, real_fsync
load.restore_runtime_rollback_receipt = runtime_phase
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.restore_cross_project_rollback_context(context, current_project=q,
            deployment_lock_descriptor=descriptor, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'))
    except OSError as error:
        assert 'runtime parent fsync fault' in str(error)
        sys.exit(7)
finally:
    load.release_deployment_lock(descriptor)
"""
            encoded = json.dumps({**before, "manifest": before["manifest"].decode()})
            argv = [sys.executable, "-B", "-c", recovery, str(root), str(p), str(q),
                    str(common), str(p_only), encoded]
            failed = subprocess.run(argv + ["fault"], cwd=root, capture_output=True,
                                    text=True, check=False, timeout=30)
            self.assertEqual(7, failed.returncode, failed.stdout + failed.stderr)
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertFalse((root / "runtime-parent-synced").exists())
            recovered = subprocess.run(argv + ["sync"], cwd=root, capture_output=True,
                                       text=True, check=False, timeout=30)
            self.assertEqual(0, recovered.returncode, recovered.stdout + recovered.stderr)
            self.assertEqual(b"parent sync observed\n", (root / "runtime-parent-synced").read_bytes())
            self.assertEqual(b"independent A runtime witness\n", (root / "ztp/runtime-durability-witness").read_bytes())
            self.assertFalse(os.path.lexists(journal))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_manifest_then_load_receipt_syncs_all_authority_before_publish(self):
        """SETUP's real manifest selects A; LOAD must durably publish its receipt."""
        fixture = RELEASE_FIXTURE.ReleaseTransactionTests(
            methodName="test_failed_b_generation_can_restore_and_verify_complete_a_runtime_receipt")
        fixture.setUp()
        self.addCleanup(fixture.tearDown)
        api = RELEASE_FIXTURE.LOAD
        root, project, ztp = fixture.root.resolve(), fixture.project.resolve(), fixture.ztp.resolve()
        settings = yaml.safe_load(fixture.global_file.read_text())
        settings["common"]["mgmt"]["http"]["http_root"] = str(root)
        fixture.global_file.write_text(yaml.safe_dump(settings, sort_keys=False))
        manifest = ztp / ".setup_manifest"
        link = ztp / "config/cumulus/template/01-global.yaml"
        link.parent.mkdir(parents=True)
        SETUP._SetupLinkTransaction._restore_link(str(link), os.path.relpath(fixture.global_file, link.parent))
        with mock.patch.object(SETUP, "MANIFEST_FILE", str(manifest)), \
                mock.patch.object(SETUP, "_collect_expected_links", return_value=[str(link)]):
            SETUP._write_manifest(str(project.resolve()))
        preimage = manifest.read_bytes()
        with mock.patch.multiple(api, HTTP_ROOT=root, MANIFEST=manifest,
                                 ZTP_PREFIX_MARKER=root / ".prefix-marker"):
            _parent, generation, _receipt = fixture._publish_a_with_runtime_rollback_receipt(prepare_receipt=False)
            selector = project / "99-output-ztp/current"
            expected = os.readlink(selector)
            receipts = project / "99-output-ztp/rollback-receipts"
            real_fsync, real_rename = os.fsync, os.rename
            synced = set()
            fault_seen = []

            def sync(descriptor):
                metadata = os.fstat(descriptor)
                if stat.S_ISREG(metadata.st_mode):
                    stage = next(receipts.glob(".rollback-receipt-*"))
                    record = stage / "receipt.json"
                    if record.exists() and record.stat().st_ino == metadata.st_ino and not fault_seen:
                        fault_seen.append(True)
                        raise OSError("independent receipt metadata sync fault")
                    synced.add((metadata.st_dev, metadata.st_ino, stat.S_IMODE(metadata.st_mode)))
                return real_fsync(descriptor)

            def publish(source, destination):
                source, destination = Path(source), Path(destination)
                if destination == receipts / generation.name:
                    expected_files = {"receipt.json", "files/ztp-bootstrap_oob.sh", "files/ztp-bootstrap_oobofoob.sh",
                                      "files/ztp.json", "publickeys/laptop.pub", "publickeys/mgmt-server.pub"}
                    actual = {p.relative_to(source).as_posix() for p in source.rglob("*") if p.is_file()}
                    self.assertEqual(expected_files, actual)
                    for name in expected_files:
                        metadata = (source / name).stat()
                        self.assertIn((metadata.st_dev, metadata.st_ino, stat.S_IMODE(metadata.st_mode)), synced,
                                      "receipt publication preceded durable file authority: " + name)
                return real_rename(source, destination)

            with mock.patch.object(api.os, "fsync", side_effect=sync), \
                    mock.patch.object(api.os, "rename", side_effect=publish):
                with self.assertRaisesRegex(OSError, "receipt metadata sync fault"):
                    api.prepare_active_runtime_rollback_receipt(project, expected)
                self.assertTrue(fault_seen)
                self.assertEqual([], list(receipts.iterdir()))
                self.assertEqual(preimage, manifest.read_bytes())
                self.assertEqual(expected, os.readlink(selector))
                synced.clear()
                receipt = api.prepare_active_runtime_rollback_receipt(project, expected)
            self.assertEqual(receipts / generation.name, receipt)
            self.assertEqual(preimage, manifest.read_bytes())
            self.assertEqual(expected, os.readlink(selector))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_publickey_links_then_load_receipt_restores_exact_publication(self):
        self._assert_real_setup_publickey_receipt_restore(rebind_plain_parent=False)

    def test_real_setup_plain_keys_then_load_receipt_parent_rebind_holds_before_write(self):
        self._assert_real_setup_publickey_receipt_restore(rebind_plain_parent=True)

    def _assert_real_setup_publickey_receipt_restore(self, *, rebind_plain_parent):
        fixture = RELEASE_FIXTURE.ReleaseTransactionTests(
            methodName="test_failed_b_generation_can_restore_and_verify_complete_a_runtime_receipt")
        fixture.setUp()
        self.addCleanup(fixture.tearDown)
        api = RELEASE_FIXTURE.LOAD
        root, project, ztp = fixture.root.resolve(), fixture.project.resolve(), fixture.ztp.resolve()
        settings = yaml.safe_load(fixture.global_file.read_text())
        settings["common"]["mgmt"]["http"]["http_root"] = str(root)
        fixture.global_file.write_text(yaml.safe_dump(settings, sort_keys=False))
        manifest = ztp / ".setup_manifest"
        with mock.patch.multiple(api, HTTP_ROOT=root, MANIFEST=manifest,
                                 DEPLOYMENT_LOCK=root / ".deployment.lock", ZTP_PREFIX_MARKER=root / ".prefix-marker"), \
                mock.patch.multiple(SETUP, HTTP_BASE=str(root), ZTP=str(ztp), MANIFEST_FILE=str(manifest),
                                    _DRY_RUN=False, _LINK_ERRORS=0, _LINK_TRANSACTION=None,
                                    _SETUP_SWITCH_WAL_CONTEXT=None):
            _parent, generation, _receipt = fixture._publish_a_with_runtime_rollback_receipt(prepare_receipt=False)
            for name, payload in (("laptop.pub", b"A laptop key\n"), ("mgmt-server.pub", b"A management key\n")):
                (project / name).write_bytes(payload)
                (ztp / "config/publickey" / name).unlink()
            SETUP._process_pubkeys(str(project))
            keys = [root / base / name for base in ("ztp/config/publickey", "infiniband/publickey")
                    for name in ("laptop.pub", "mgmt-server.pub")]
            targets = {path: os.readlink(path) for path in keys}
            if rebind_plain_parent:
                # Plain public keys are a supported receipt input, independent
                # of SETUP's usual managed-link publication.
                for path in keys[:2]:
                    payload = path.read_bytes()
                    path.unlink()
                    path.write_bytes(payload)
            with mock.patch.object(SETUP, "_collect_expected_links", return_value=list(map(str, keys))):
                SETUP._write_manifest(str(project))
            before = manifest.read_bytes()
            selector = f"generations/{generation.name}"
            receipt = api.prepare_active_runtime_rollback_receipt(project, selector)
            self.assertIsNotNone(receipt)
            receipt_bytes = (receipt / "receipt.json").read_bytes()
            live = ztp / "ztp.json"
            live.write_bytes(b"partial B runtime\n")
            for path in keys[:2]:
                path.unlink()
                path.write_bytes(b"partial B flat key\n")
            if rebind_plain_parent:
                parent = ztp / "config/publickey"
                original = parent.stat().st_ino
                parent.rename(parent.with_name("original-publickey"))
                parent.mkdir()
                (parent / "operator.pub").write_bytes(b"unrelated replacement key\n")
                self.assertNotEqual(original, parent.stat().st_ino)
            with mock.patch.object(api, "active_managed_services", return_value=()):
                descriptor = api.acquire_deployment_lock(exclusive=True)
                try:
                    def restore():
                        api.restore_runtime_rollback_receipt(project, receipt, expected_current=selector,
                            deployment_lock_descriptor=descriptor, services_quiesced=True,
                            runtime_backend=SimpleNamespace(name="supervisor"))
                    if rebind_plain_parent:
                        with self.assertRaisesRegex(api.LoadError, "publickey parent|public key parent"):
                            restore()
                    else:
                        restore()
                finally:
                    api.release_deployment_lock(descriptor)
            self.assertEqual(before, manifest.read_bytes())
            self.assertEqual(receipt_bytes, (receipt / "receipt.json").read_bytes())
            if rebind_plain_parent:
                self.assertEqual(b"partial B runtime\n", live.read_bytes())
                self.assertEqual(["operator.pub"], sorted(path.name for path in parent.iterdir()))
                self.assertEqual(b"unrelated replacement key\n", (parent / "operator.pub").read_bytes())
                self.assertEqual(selector, os.readlink(project / "99-output-ztp/current"))
                self.assertFalse((root / "service-started").exists())
                return
            for path, target in targets.items():
                self.assertTrue(path.is_symlink())
                self.assertEqual(target, os.readlink(path))
            self.assertEqual(b'{"release":"A"}\n', live.read_bytes())
            self.assertEqual(selector, os.readlink(project / "99-output-ztp/current"))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_success_then_main_retains_wal_and_refuses_whole_b_completion(self):
        self._assert_real_setup_main_pending_wal_holds(start_services=False)

    def test_real_setup_pending_wal_blocks_main_before_requested_service_start(self):
        self._assert_real_setup_main_pending_wal_holds(start_services=True)

    def _assert_real_setup_main_pending_wal_holds(self, *, start_services):
        with tempfile.TemporaryDirectory(prefix="r22-main-pending-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, _p_only, _q_only, manifest, _before, _hook, _barrier, _home = fixture
            for relative in ("tools/lldp-analyze-tool/99-output-p2p",
                             "tools/lldp-analyze-tool/99-output-monitor",
                             "tools/ibdiagnet-analyze-tool/99-output-p2p"):
                leaf = root / relative
                if leaf.is_dir() and not leaf.is_symlink():
                    shutil.rmtree(leaf)
            journal = root / "ztp/.setup-switch-journal.json"
            setup_state = []

            def real_setup_action():
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=False, auto_prepare_journal=True,
                    kill_after_replace=root / "never-replaced",
                )
                self.assertEqual(25, result.returncode, result.stdout + result.stderr)
                self.assertTrue(journal.is_file())
                setup_state.append((manifest.read_bytes(), os.readlink(common), journal.read_bytes()))

            harness = RELEASE_FIXTURE.ReleaseTransactionTests()
            harness.setUp()
            try:
                harness._assert_main_pending_wal_completion_holds(
                    activation_action=real_setup_action, journal=journal,
                    start_services=start_services,
                )
            finally:
                harness.tearDown()
                harness.doCleanups()
            self.assertEqual([setup_state[0]],
                             [(manifest.read_bytes(), os.readlink(common), journal.read_bytes())])
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_materialized_registration_rejects_wal_write_aba(self):
        with tempfile.TemporaryDirectory(prefix="r22-registration-wal-aba-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, _q_only, manifest, before, hook, barrier, _home = fixture
            output = q / "99-output-eth"
            latest, release = output / "latest", output / "20261007_080000_combine"
            for path in (latest, release):
                path.mkdir(parents=True)
                (path / ".published-complete").write_bytes(b"complete\n")
                (path / "leaf.yaml").write_bytes(b"independent Q bytes\n")
            original_latest = latest.stat().st_ino
            addition = r'''
import json
import stat
import sys
original_fsync_for_wal_aba = os.fsync
def inject_registration_wal_aba(fd):
    result = original_fsync_for_wal_aba(fd)
    root = Path(__file__).resolve().parent.parent
    journal = root / 'ztp/.setup-switch-journal.json'
    witness = root / 'registration-wal-aba.json'
    held = os.fstat(fd)
    if (Path(sys.argv[0]).name == '01-a-setup.py' and not witness.exists()
            and journal.exists() and stat.S_ISREG(held.st_mode)
            and any(item.name.startswith('.setup-switch-journal.')
                    and item.name != journal.name and item.stat().st_ino == held.st_ino
                    for item in journal.parent.iterdir())):
        frozen, before = journal.read_bytes(), journal.stat()
        journal.write_bytes(b'!' + frozen[1:])
        journal.write_bytes(frozen)
        os.utime(journal, ns=(before.st_atime_ns, before.st_mtime_ns))
        after = journal.stat()
        (root / 'registration-wal-before.bin').write_bytes(frozen)
        witness.write_text(json.dumps({
            'before': [before.st_dev, before.st_ino, before.st_size, before.st_mtime_ns],
            'after': [after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns],
            'ctime_changed': before.st_ctime_ns != after.st_ctime_ns,
        }))
    return result
os.fsync = inject_registration_wal_aba
'''
            site = hook / "sitecustomize.py"
            site.write_text(site.read_text() + addition)
            result = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            witness = json.loads((root / "registration-wal-aba.json").read_text())
            self.assertEqual(witness["before"], witness["after"])
            self.assertTrue(witness["ctime_changed"])
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertIn("journal changed before directory registration", result.stdout + result.stderr)
            self.assertFalse(barrier.exists(), "changed WAL history authorized a shared Q write")
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertEqual((root / "registration-wal-before.bin").read_bytes(), journal.read_bytes())
            self.assertEqual(witness["before"][1], journal.stat().st_ino)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertEqual(original_latest, latest.stat().st_ino)
            self.assertEqual(b"independent Q bytes\n", (latest / "leaf.yaml").read_bytes())
            self.assertFalse(list(output.glob(".latest.setup-backup.*")))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_legacy_sample_comparison_holds_before_first_shared_write(self):
        with tempfile.TemporaryDirectory(prefix="r22-sample-legacy-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            legacy = root / "ztp/optimize/q-sample/comparison"
            legacy.mkdir(parents=True)
            (legacy / "operator-report.csv").write_bytes(b"original comparison bytes\n")
            identity = (legacy.stat().st_dev, legacy.stat().st_ino)
            result = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertRegex(result.stdout + result.stderr, "sample.*comparison|comparison.*WAL")
            self.assertFalse(barrier.exists(), "uncovered comparison did not fence first Q link")
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertEqual(identity, (legacy.stat().st_dev, legacy.stat().st_ino))
            self.assertEqual(b"original comparison bytes\n", (legacy / "operator-report.csv").read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_created_directory_recovery_death_resumes_in_fresh_load(self):
        """A second fresh LOAD must finish after recovery itself dies at rmdir."""
        with tempfile.TemporaryDirectory(prefix="r22-created-dir-reentry-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, _common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            key_root = root / "infiniband/publickey"
            key_root.parent.mkdir(exist_ok=True)
            self.assertFalse(os.path.lexists(key_root))
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=str(key_root) + os.sep,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            self.assertTrue(key_root.is_dir())
            recovery = """
import importlib.util
import os
import signal
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_created_dir_reentry', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
if len(sys.argv) > 2:
    real_rmdir = Path.rmdir
    def interrupted_rmdir(path):
        real_rmdir(path)
        if path == root / 'infiniband/publickey':
            (root / 'recovery-rmdir-barrier').write_bytes(b'owned-directory-removed\\n')
            os.kill(os.getpid(), signal.SIGKILL)
    Path.rmdir = interrupted_rmdir
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root), "kill-after-rmdir"],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(-9, second.returncode, second.stdout + second.stderr)
            self.assertEqual(b"owned-directory-removed\n", (root / "recovery-rmdir-barrier").read_bytes())
            self.assertFalse(os.path.lexists(key_root))
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertTrue(journal.is_file())
            third = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)], cwd=root,
                capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, third.returncode, third.stdout + third.stderr)
            self.assertFalse(os.path.lexists(key_root))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertFalse(os.path.lexists(journal))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_death_unknown_systemd_state_blocks_cold_restore(self):
        """Fresh LOAD must not infer stopped services from a failed status query."""
        for mode in ("query-error", "missing-command"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory(prefix="r22-service-unknown-") as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, _q, common, p_only, q_only, manifest, _before, _hook, barrier, _home = fixture
                first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
                self.assertEqual(23, first.returncode, first.stdout + first.stderr)
                self.assertTrue(barrier.is_file())
                journal = root / "ztp/.setup-switch-journal.json"
                evidence = journal.read_bytes()
                manifest_before = manifest.read_bytes()
                links_before = {path: os.readlink(path) if path.is_symlink() else None
                                for path in (common, p_only, q_only)}
                recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
mode = sys.argv[2]
spec = importlib.util.spec_from_file_location('fresh_unknown_service', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.runtime_os = lambda: 'Linux'
load.shutil.which = lambda name: None if mode == 'missing-command' else '/fixture/systemctl'
def query(command, **kwargs):
    assert command[1] == 'is-active', command
    with (root / 'service-queries').open('a') as stream:
        stream.write(command[-1] + '\\n')
    if mode == 'missing-command':
        raise FileNotFoundError('systemctl unavailable')
    return SimpleNamespace(returncode=4, stdout='', stderr='manager unavailable')
load._run_subprocess = query
backend = load.service_runtime_backend({'HTTP_ZTP_RUNTIME_BACKEND': 'systemd'}, command_runner=query)
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=backend,
        )
    except load.LoadError as exc:
        assert '服务' in str(exc), str(exc)
        print('R22_SERVICE_UNKNOWN_HOLD=' + str(exc))
    else:
        raise AssertionError('unknown service state authorized P rollback')
finally:
    load.release_deployment_lock(descriptor)
"""
                second = subprocess.run(
                    [sys.executable, "-B", "-c", recovery, str(root), mode], cwd=root,
                    capture_output=True, text=True, check=False, timeout=30,
                )
                self.assertEqual(0, second.returncode, second.stdout + second.stderr)
                self.assertIn("R22_SERVICE_UNKNOWN_HOLD=", second.stdout)
                self.assertEqual(b"isc-dhcp-server\n", (root / "service-queries").read_bytes())
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(manifest_before, manifest.read_bytes())
                self.assertEqual(links_before, {path: os.readlink(path) if path.is_symlink() else None
                                              for path in (common, p_only, q_only)})
                self.assertFalse((root / "service-started").exists())

    def test_real_setup_death_pending_wal_blocks_new_load_before_template(self):
        """A real SETUP crash must fence a fresh LOAD before any new project writes."""
        with tempfile.TemporaryDirectory(prefix="r22-pending-wal-main-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, q_only, manifest, _before, _hook, barrier, home = fixture
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            journal = root / "ztp/.setup-switch-journal.json"
            wal_bytes = journal.read_bytes()
            manifest_bytes = manifest.read_bytes()
            common_target = os.readlink(common)
            driver = """
import importlib.util
import sys
from pathlib import Path
root, q = map(Path, sys.argv[1:3])
spec = importlib.util.spec_from_file_location('fresh_pending_wal_main', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.runtime_os = lambda: 'Linux'
original = load.initialize_from_template
def observed_template(project, dry_run):
    (root / 'unsafe-template-entry').write_bytes(b'new LOAD entered template mutation\\n')
    return original(project, dry_run)
load.initialize_from_template = observed_template
raise SystemExit(load.main([str(q), '--host-role=workstation', '--no-upgrade']))
"""
            environment = dict(os.environ)
            environment["HOME"] = str(home)
            dependencies = [item for item in sys.path
                            if item and item.endswith("site-packages") and Path(item).is_dir()]
            environment["PYTHONPATH"] = os.pathsep.join(dependencies)
            second = subprocess.run(
                [sys.executable, "-B", "-c", driver, str(root), str(q)],
                cwd=root, env=environment, capture_output=True, text=True,
                check=False, timeout=30,
            )
            self.assertEqual(1, second.returncode, second.stdout + second.stderr)
            self.assertFalse((root / "unsafe-template-entry").exists())
            self.assertIn("pending setup switch WAL", second.stderr)
            self.assertEqual(wal_bytes, journal.read_bytes())
            self.assertEqual(manifest_bytes, manifest.read_bytes())
            self.assertEqual(common_target, os.readlink(common))
            self.assertFalse(os.path.lexists(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_wal_cold_recovery_rejects_rebound_link_parent(self):
        """A killed SETUP must not let fresh LOAD restore P into a new parent."""
        with tempfile.TemporaryDirectory(prefix="r22-parent-rebind-cold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            journal = root / "ztp/.setup-switch-journal.json"
            frozen_journal = journal.read_bytes()
            original_parent = common.parent.with_name(common.parent.name + "-original")
            common.parent.rename(original_parent)
            common.parent.mkdir()
            common.symlink_to(os.readlink(original_parent / common.name))
            replacement_link = os.readlink(common)

            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location(
    'fresh_rebound_setup_switch_load', root / 'DAY0-Prepare/11-load.py',
)
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=descriptor, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        if any(word in str(exc) for word in ('parent', 'directory', 'identity')):
            raise SystemExit(27)
        raise
    raise SystemExit(28)
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(27, second.returncode, second.stdout + second.stderr)
            self.assertEqual(replacement_link, os.readlink(common))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertFalse(os.path.lexists(p_only))
            self.assertEqual(frozen_journal, journal.read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_prepare_only_keeps_p_publication_unchanged(self):
        """A separate Q-local preparation pass must not switch P's links."""
        with tempfile.TemporaryDirectory(prefix="r22-setup-prepare-only-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, q_only, manifest, before, _hook, barrier, home = fixture
            environment = os.environ.copy()
            environment["HOME"] = str(home)
            dependency_paths = [
                item for item in sys.path
                if item and item.endswith("site-packages") and Path(item).is_dir()
            ]
            environment["PYTHONPATH"] = os.pathsep.join(
                item for item in (environment.get("PYTHONPATH", ""), *dependency_paths)
                if item
            )
            result = subprocess.run(
                [sys.executable, "-B", str(root / "DAY0-Prepare/01-a-setup.py"),
                 "--prepare-only", "--no-upgrade", "--host-role=workstation", str(q)],
                cwd=root, env=environment, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, result.returncode, result.stdout + result.stderr)
            self.assertFalse(barrier.exists())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())
            for output in ("99-output-dhcp", "99-output-eth", "99-output-infra"):
                self.assertTrue((q / output).is_dir(), output)

    def test_real_load_setup_child_requires_durable_wal_before_first_q_link(self):
        """Without a WAL, LOAD/SETUP must not create even one Q managed link."""
        with tempfile.TemporaryDirectory(prefix="r22-real-setup-no-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            result = self._run_private_load_setup_switch(fixture, prepare_journal=False)
            self.assertNotEqual(0, result.returncode)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))

    def test_real_load_setup_child_rejects_corrupt_wal_before_first_q_link(self):
        """A malformed WAL is not authority to mutate the first Q link."""
        with tempfile.TemporaryDirectory(prefix="r22-real-setup-bad-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, corrupt_journal=True,
            )
            self.assertNotEqual(0, result.returncode)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertEqual(
                b'{"schema_version":',
                (root / "ztp/.setup-switch-journal.json").read_bytes(),
            )

    def test_real_setup_rejects_forged_same_project_manifest_before_q_link(self):
        """A forged Q header over P links must not bypass the WAL fence."""
        with tempfile.TemporaryDirectory(prefix="r22-real-setup-forged-header-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, q_only, manifest, _before, _hook, barrier, _home = fixture
            forged = manifest.read_text(encoding="utf-8").replace(
                f"# setup manifest — proj: {p}\n",
                f"# setup manifest — proj: {q}\n", 1,
            )
            manifest.write_text(forged, encoding="utf-8")
            result = self._run_private_load_setup_switch(fixture, prepare_journal=False)
            self.assertNotEqual(0, result.returncode)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(forged, manifest.read_text(encoding="utf-8"))
            self.assertEqual(p / "01-global.yaml", common.resolve())
            self.assertEqual(p / "operator.pub", p_only.resolve())
            self.assertFalse(os.path.lexists(q_only))

    def test_direct_setup_cli_without_wal_rejects_before_first_q_link(self):
        """Direct setup cannot bypass LOAD's cross-project write fence."""
        with tempfile.TemporaryDirectory(prefix="r22-direct-setup-no-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, q_only, manifest, before, hook, barrier, home = fixture
            environment = os.environ.copy()
            environment["HOME"] = str(home)
            dependency_paths = [
                item for item in sys.path
                if item and item.endswith("site-packages") and Path(item).is_dir()
            ]
            environment["PYTHONPATH"] = os.pathsep.join(
                item for item in (str(hook), environment.get("PYTHONPATH", ""),
                                  *dependency_paths) if item
            )
            environment["HTTP_TEST_FIRST_Q_LINK"] = str(common)
            environment["HTTP_TEST_Q_LINK_BARRIER"] = str(barrier)
            result = subprocess.run(
                [sys.executable, "-B", str(root / "DAY0-Prepare/01-a-setup.py"),
                 "-y", "--confirm-project-switch", "--no-upgrade",
                 "--host-role=workstation", str(q)],
                cwd=root, env=environment, capture_output=True, text=True, check=False,
            )
            self.assertNotEqual(0, result.returncode)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())

    def test_real_setup_wal_refuses_unjournaled_materialized_latest_before_q_link(self):
        """A link-only WAL cannot authorize a real latest-directory rename."""
        with tempfile.TemporaryDirectory(prefix="r22-real-latest-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            latest = q / "99-output-eth/latest"
            latest.mkdir(parents=True)
            marker = latest / "marker.yaml"
            marker.write_bytes(b"PRESERVE-MATERIALIZED-LATEST\n")
            result = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertNotEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertEqual(b"PRESERVE-MATERIALIZED-LATEST\n", marker.read_bytes())
            self.assertTrue((root / "ztp/.setup-switch-journal.json").is_file())

    def _assert_materialized_temp_cold_recovery(self, foreign_owner):
        with tempfile.TemporaryDirectory(prefix="r22-real-materialized-temp-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            output = q / "99-output-eth"
            latest, release = output / "latest", output / "20261007_010000_combine"
            for path in (latest, release):
                path.mkdir(parents=True)
                (path / ".published-complete").write_bytes(b"complete\n")
                (path / "leaf.yaml").write_bytes(b"hostname: independent-leaf\n")
            original = latest.stat()
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True,
                first_link_prefix=str(output / ".latest.tmp."),
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            temporary = Path(barrier.read_text().strip())
            self.assertEqual(output, temporary.parent)
            self.assertEqual(release.name, os.readlink(temporary))
            self.assertFalse(os.path.lexists(latest))
            journal = root / "ztp/.setup-switch-journal.json"
            evidence = journal.read_bytes()
            row = json.loads(evidence)["body"]["directories"][0]
            backup = Path(row["backup"])
            self.assertEqual((original.st_dev, original.st_ino),
                             (backup.stat().st_dev, backup.stat().st_ino))
            mixed = os.readlink(common)
            if foreign_owner:
                temporary.unlink()
                temporary.write_bytes(b"UNRELATED RECOVERY OBJECT\n")
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_materialized_temp_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        outcome = load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=descriptor, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        if sys.argv[2] != 'foreign' or 'disk state is unknown' not in str(exc):
            raise
        print('MATERIALIZED_TEMP_FOREIGN_OWNER_HELD')
    else:
        assert sys.argv[2] == 'owned', 'foreign temp must block all recovery writes'
        assert outcome == 'restored-p', outcome
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root),
                 "foreign" if foreign_owner else "owned"],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertFalse((root / "service-started").exists())
            if foreign_owner:
                self.assertIn("MATERIALIZED_TEMP_FOREIGN_OWNER_HELD", second.stdout)
                self.assertEqual(evidence, journal.read_bytes())
                self.assertEqual(mixed, os.readlink(common))
                self.assertFalse(os.path.lexists(latest))
                self.assertEqual(b"UNRELATED RECOVERY OBJECT\n", temporary.read_bytes())
                self.assertEqual((original.st_dev, original.st_ino),
                                 (backup.stat().st_dev, backup.stat().st_ino))
            else:
                self.assertFalse(os.path.lexists(temporary), "WAL must consume its temp link")
                self.assertFalse(os.path.lexists(backup))
                self.assertFalse(os.path.lexists(journal))
                self.assertEqual((original.st_dev, original.st_ino),
                                 (latest.stat().st_dev, latest.stat().st_ino))
                self.assertEqual(before["common"], os.readlink(common))
                self.assertEqual(before["p_only"], os.readlink(p_only))
                self.assertFalse(os.path.lexists(q_only))

    def test_real_setup_materialized_temp_death_cold_recovers_original_and_removes_temp(self):
        self._assert_materialized_temp_cold_recovery(False)

    def test_real_setup_materialized_temp_death_foreign_owner_blocks_all_restore(self):
        self._assert_materialized_temp_cold_recovery(True)

    def _assert_real_materialized_writer_registration_drift_holds(self, drift):
        with tempfile.TemporaryDirectory(prefix="r22-materialized-writer-drift-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, _common, _p_only, _q_only, _manifest, _before, hook, barrier, _home = fixture
            output = q / "99-output-eth"
            latest, release = output / "latest", output / "20261007_090000_combine"
            payload = b"hostname: independent migration owner\n"
            for path in (latest, release):
                path.mkdir(parents=True)
                (path / ".published-complete").write_bytes(b"complete\n")
                (path / "leaf.yaml").write_bytes(payload)
            original_source, original_target = latest.stat(), release.stat()
            site = hook / "sitecustomize.py"
            site.write_text(site.read_text(encoding="utf-8") + """
import json
previous_materialized_replace = os.replace
def drift_materialized_after_temp_registration(source, destination, *args, **kwargs):
    result = previous_materialized_replace(source, destination, *args, **kwargs)
    journal = Path(os.environ['HTTP_TEST_MATERIALIZED_WAL'])
    marker = journal.parent / 'materialized-writer-drifted'
    if Path(destination) == journal and not marker.exists():
        evidence = journal.read_bytes()
        body = json.loads(evidence)['body']
        latest = Path(os.environ['HTTP_TEST_MATERIALIZED_SOURCE'])
        if any(Path(row['path']).parent == latest.parent
               and Path(row['path']).name.startswith('.latest.tmp.') for row in body['links']):
            changed = latest if os.environ['HTTP_TEST_MATERIALIZED_DRIFT'] == 'source' else Path(os.environ['HTTP_TEST_MATERIALIZED_TARGET'])
            changed.rename(latest.parent / 'retained-original-migration-owner')
            changed.mkdir()
            (changed / '.published-complete').write_bytes(b'complete\\n')
            (changed / 'leaf.yaml').write_bytes(b'hostname: independent migration owner\\n')
            marker.write_bytes(evidence)
    return result
os.replace = drift_materialized_after_temp_registration
""", encoding="utf-8")
            journal = root / "ztp/.setup-switch-journal.json"
            with mock.patch.dict(os.environ, {
                "HTTP_TEST_MATERIALIZED_WAL": str(journal),
                "HTTP_TEST_MATERIALIZED_SOURCE": str(latest),
                "HTTP_TEST_MATERIALIZED_TARGET": str(release),
                "HTTP_TEST_MATERIALIZED_DRIFT": drift,
            }):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True, first_link_prefix=str(output / ".latest.tmp."),
                )
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertFalse(barrier.exists(), "registration drift reached latest publication")
            self.assertTrue(latest.is_dir(), "writer renamed a source before its HOLD")
            self.assertFalse(latest.is_symlink())
            evidence = (journal.parent / "materialized-writer-drifted").read_bytes()
            self.assertEqual(evidence, journal.read_bytes())
            row = next(row for row in json.loads(evidence)["body"]["directories"]
                       if row["path"] == str(latest))
            self.assertFalse(os.path.lexists(row["backup"]))
            retained = output / "retained-original-migration-owner"
            original = original_source if drift == "source" else original_target
            self.assertEqual(original.st_ino, retained.stat().st_ino)
            changed = latest if drift == "source" else release
            self.assertNotEqual(original.st_ino, changed.stat().st_ino)
            self.assertEqual(payload, (changed / "leaf.yaml").read_bytes())
            self.assertEqual(payload, (retained / "leaf.yaml").read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_materialized_writer_source_drift_blocks_rename(self):
        self._assert_real_materialized_writer_registration_drift_holds("source")

    def test_real_setup_materialized_writer_target_drift_blocks_rename(self):
        self._assert_real_materialized_writer_registration_drift_holds("target")

    def test_real_setup_materialized_latest_replace_cold_recovers_original_directory(self):
        """LOAD→SETUP death after replacing latest must restore its real Q bytes."""
        with tempfile.TemporaryDirectory(prefix="r22-real-latest-restore-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            output = q / "99-output-eth"
            latest = output / "latest"
            release = output / "20261005_120000_combine"
            for path in (latest, release):
                path.mkdir(parents=True)
                (path / ".published-complete").write_bytes(b"complete\n")
                (path / "leaf.yaml").write_bytes(b"hostname: leaf\n")
            original = latest.stat()
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, kill_after_replace=latest,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file(), first.stdout + first.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertTrue(latest.is_symlink())
            journal = root / "ztp/.setup-switch-journal.json"
            rows = json.loads(journal.read_bytes())["body"]["directories"]
            self.assertEqual(1, len(rows))
            backup = Path(rows[0]["backup"])
            self.assertEqual((original.st_dev, original.st_ino),
                             (backup.stat().st_dev, backup.stat().st_ino))

            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location(
    'fresh_materialized_switch_load', root / 'DAY0-Prepare/11-load.py',
)
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(os.path.lexists(backup))
            self.assertFalse(os.path.lexists(journal))
            self.assertFalse(latest.is_symlink())
            self.assertEqual((original.st_dev, original.st_ino),
                             (latest.stat().st_dev, latest.stat().st_ino))
            self.assertEqual(b"hostname: leaf\n", (latest / "leaf.yaml").read_bytes())

    def test_real_setup_materialized_cold_recovery_preserves_backup_replaced_after_unlink(self):
        with tempfile.TemporaryDirectory(prefix="r22-materialized-backup-owner-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, _before, _hook, _barrier, _home = fixture
            output = q / "99-output-eth"
            latest, release = output / "latest", output / "20261007_070000_combine"
            for path in (latest, release):
                path.mkdir(parents=True)
                (path / ".published-complete").write_bytes(b"complete\n")
                (path / "leaf.yaml").write_bytes(b"independent original leaf\n")
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True, kill_after_replace=latest)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            journal = root / "ztp/.setup-switch-journal.json"
            evidence, manifest_bytes, mixed = journal.read_bytes(), manifest.read_bytes(), os.readlink(common)
            row = json.loads(evidence)["body"]["directories"][0]
            backup = Path(row["backup"])
            original_inode = backup.lstat().st_ino
            recovery = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root, latest, backup = map(Path, sys.argv[1:4])
spec = importlib.util.spec_from_file_location('fresh_materialized_backup_owner_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
sync = load._fsync_parent_directory
injected = []
def replace_backup_after_unlink(path):
    sync(path)
    if Path(path) == latest.parent and not latest.is_symlink() and not latest.exists() and not injected:
        backup.rename(backup.with_name('operator-retained-original'))
        backup.mkdir()
        (backup / 'unrelated.txt').write_bytes(b'DO NOT MOVE\\n')
        injected.append(backup.stat().st_ino)
load._fsync_parent_directory = replace_backup_after_unlink
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
            deployment_lock_descriptor=descriptor, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as error:
        assert 'materialized' in str(error) and any(word in str(error) for word in ('owner', 'changed', 'identity')), str(error)
        assert len(injected) == 1, 'injection boundary was not reached'
        assert not latest.exists(), 'recovery moved an unknown backup into latest'
        assert backup.stat().st_ino == injected[0], 'unknown backup was consumed'
    else:
        raise AssertionError('cold recovery accepted a replaced materialized backup')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), str(latest), str(backup)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(latest))
            self.assertEqual(b"DO NOT MOVE\n", (backup / "unrelated.txt").read_bytes())
            retained = output / "operator-retained-original"
            self.assertEqual(original_inode, retained.lstat().st_ino)
            self.assertEqual(b"independent original leaf\n", (retained / "leaf.yaml").read_bytes())
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(manifest_bytes, manifest.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_materialized_cold_wal_drift_after_unlink_holds_backup(self):
        self._assert_real_materialized_cold_authority_holds("wal")

    def test_real_setup_materialized_cold_q_commit_after_unlink_holds_backup(self):
        self._assert_real_materialized_cold_authority_holds("q")

    def _assert_real_materialized_cold_authority_holds(self, kind):
        with tempfile.TemporaryDirectory(prefix="r22-materialized-total-authority-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, _before, _hook, barrier, _home = fixture
            output = q / "99-output-eth"
            latest, release = output / "latest", output / "20261007_070000_combine"
            for path in (latest, release):
                path.mkdir(parents=True)
                (path / ".published-complete").write_bytes(b"complete\n")
                (path / "leaf.yaml").write_bytes(b"independent original leaf\n")
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True, kill_after_replace=latest)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            journal = root / "ztp/.setup-switch-journal.json"
            evidence, manifest_bytes, mixed = journal.read_bytes(), manifest.read_bytes(), os.readlink(common)
            backup = Path(json.loads(evidence)["body"]["directories"][0]["backup"])
            original_owner = backup.stat().st_ino
            recovery = """
import importlib.util, os, sys
from pathlib import Path
from types import SimpleNamespace
root, latest, backup = map(Path, sys.argv[1:4])
kind = sys.argv[4]
spec = importlib.util.spec_from_file_location('fresh_materialized_authority_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
journal = root / 'ztp/.setup-switch-journal.json'
authority = journal if kind == 'wal' else latest.parent.parent / '99-output-ztp/current-release.json'
prior = journal.read_bytes()
sync = load._fsync_parent_directory
injected = []
def drift_after_unlink(parent):
    sync(parent)
    if parent == latest.parent and not os.path.lexists(latest) and not injected:
        if kind == 'wal':
            owner = journal.stat().st_ino
            journal.rename(journal.with_name('held-original-materialized-authority'))
            journal.write_bytes(prior)
            journal.chmod(0o600)
            assert owner != journal.stat().st_ino
        else:
            authority.parent.mkdir(parents=True, exist_ok=True)
            authority.write_bytes(b'independent Q commit marker\\n')
        injected.append(authority.stat().st_ino)
load._fsync_parent_directory = drift_after_unlink
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal,
            deployment_lock_descriptor=descriptor, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as error:
        assert ('journal changed' if kind == 'wal' else 'Q parent changed') in str(error), str(error)
        assert len(injected) == 1, 'authority injection boundary was not reached'
        assert not os.path.lexists(latest), 'authority drift licensed backup rename into latest'
        assert backup.is_dir(), 'original backup was consumed'
        assert authority.stat().st_ino == injected[0]
        raise SystemExit(7)
    raise AssertionError('cold recovery accepted authority drift')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), str(latest), str(backup), kind],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(7, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(latest))
            self.assertEqual(original_owner, backup.stat().st_ino)
            self.assertEqual(b"independent original leaf\n", (backup / "leaf.yaml").read_bytes())
            self.assertEqual(evidence, journal.read_bytes())
            if kind == "wal":
                self.assertEqual(evidence, journal.with_name("held-original-materialized-authority").read_bytes())
            else:
                self.assertEqual(b"independent Q commit marker\n",
                                 (q / "99-output-ztp/current-release.json").read_bytes())
            self.assertEqual(manifest_bytes, manifest.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_wal_refuses_unowned_infra_log_receipt_before_q_link(self):
        """An existing receipt outside this WAL cannot be inferred as Q-owned."""
        with tempfile.TemporaryDirectory(prefix="r22-real-infra-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            logs = root / "infra/logs"
            logs.mkdir(parents=True)
            marker = logs / "audit.log"
            marker.write_bytes(b"PRESERVE-INFRA-LOGS\n")
            receipt = root / "infra/.logs-migration.json"
            receipt.write_bytes(b'{"unowned":true}\n')
            receipt.chmod(0o600)
            result = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertNotEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertEqual(b"PRESERVE-INFRA-LOGS\n", marker.read_bytes())
            self.assertEqual(b'{"unowned":true}\n', receipt.read_bytes())
            self.assertTrue((root / "ztp/.setup-switch-journal.json").is_file())

    def test_real_setup_cold_recovery_preserves_unrecognized_wal_write_set(self):
        """Real SETUP death must not let a fresh LOAD ignore an unsupported write set."""
        with tempfile.TemporaryDirectory(prefix="r22-unknown-write-set-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, _before, _hook, _barrier, _home = fixture
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            journal = root / "ztp/.setup-switch-journal.json"
            envelope = json.loads(journal.read_bytes())
            envelope["body"]["unrecognized_shared_writes"] = [{"preimage": "do-not-retire"}]
            canonical = json.dumps(envelope["body"], sort_keys=True, separators=(",", ":"),
                                   ensure_ascii=True).encode("ascii")
            envelope["sha256"] = hashlib.sha256(canonical).hexdigest()
            journal.write_text(json.dumps(envelope, sort_keys=True) + "\n", encoding="ascii")
            evidence = journal.read_bytes()
            identity = journal.lstat()
            mixed = {path: os.readlink(path) if path.is_symlink() else None
                     for path in (common, p_only, q_only)}
            manifest_bytes = manifest.read_bytes()
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_unknown_write_set_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        assert 'fields' in str(exc) or 'unknown' in str(exc), str(exc)
    else:
        raise AssertionError('recovery ignored an unsupported WAL write set')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(manifest_bytes, manifest.read_bytes())
            self.assertEqual(mixed, {path: os.readlink(path) if path.is_symlink() else None
                                     for path in mixed})
            self.assertEqual(evidence, journal.read_bytes())
            now = journal.lstat()
            self.assertEqual((identity.st_ino, identity.st_size, identity.st_mtime_ns, identity.st_ctime_ns),
                             (now.st_ino, now.st_size, now.st_mtime_ns, now.st_ctime_ns))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_cold_recovery_preserves_replacement_wal_at_retirement(self):
        """Real SETUP death and a fresh LOAD cannot retire a later WAL owner."""
        with tempfile.TemporaryDirectory(prefix="r22-wal-retirement-owner-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, before, _hook, _barrier, _home = fixture
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            journal = root / "ztp/.setup-switch-journal.json"
            original = journal.read_bytes()
            retained = journal.with_name("operator-retained-wal.json")
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_wal_retirement_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
journal = root / 'ztp/.setup-switch-journal.json'
snapshot = load.snapshot_active_setup_state
def replace_after_postverify(*args, **kwargs):
    state = snapshot(*args, **kwargs)
    journal.rename(journal.with_name('operator-retained-wal.json'))
    journal.write_bytes(b'operator replacement evidence\\n')
    journal.chmod(0o600)
    return state
load.snapshot_active_setup_state = replace_after_postverify
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            journal, deployment_lock_descriptor=descriptor, services_quiesced=True,
            runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        assert 'journal' in str(exc) and 'changed' in str(exc), str(exc)
    else:
        raise AssertionError('recovery retired another owner WAL')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertEqual(original, retained.read_bytes())
            self.assertEqual(b"operator replacement evidence\n", journal.read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_cold_restore_preserves_foreign_link_stage(self):
        for kind in ("publish", "hold"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(prefix="r22-foreign-link-stage-") as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, _q, common, _p_only, _q_only, manifest, before, *_rest = fixture
                first = self._run_private_load_setup_switch(fixture, prepare_journal=True)
                self.assertEqual(23, first.returncode, first.stdout + first.stderr)
                target_before = (common.lstat().st_ino, os.readlink(common))
                manifest_before = manifest.read_bytes()
                journal = root / "ztp/.setup-switch-journal.json"
                wal = journal.read_bytes()
                recovery = r'''
import importlib.util, json, os, sys
from pathlib import Path
from types import SimpleNamespace
root, target, kind = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3]
spec = importlib.util.spec_from_file_location('foreign_link_stage_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
journal = root / 'ztp/.setup-switch-journal.json'
original, injected = Path.symlink_to, []
def stage_replaced(stage, raw, *args, **kwargs):
    result = original(stage, raw, *args, **kwargs)
    if stage.parent == target.parent and stage.name.startswith('.' + target.name + '.setup-rollback.') and not injected:
        held = stage.with_name('held-original-link-stage')
        stage.rename(held)
        if kind == 'publish':
            original(stage, 'independent-foreign-target')
        else:
            stage.write_bytes(b'independent foreign stage evidence\n')
            value = journal.read_bytes()
            journal.rename(root / 'held-original-link-stage-wal')
            journal.write_bytes(value)
            journal.chmod(0o600)
        injected.append(True)
        (root / 'link-stage-injection.json').write_text(json.dumps({'path': str(stage), 'ino': stage.lstat().st_ino,
            'held': str(held), 'wal_ino': journal.lstat().st_ino}))
    return result
Path.symlink_to = stage_replaced
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal, deployment_lock_descriptor=fd,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as exc:
        assert injected == [True], str(exc)
        print('EXPECTED FOREIGN LINK STAGE HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(fd)
'''
                second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root), str(common), kind],
                                        cwd=root, capture_output=True, text=True, check=False, timeout=30)
                marker = json.loads((root / "link-stage-injection.json").read_text())
                self.assertEqual(7, second.returncode, second.stdout + second.stderr)
                self.assertEqual(target_before, (common.lstat().st_ino, os.readlink(common)))
                self.assertEqual(manifest_before, manifest.read_bytes())
                self.assertEqual(wal, journal.read_bytes())
                self.assertEqual(marker["wal_ino"], journal.lstat().st_ino)
                stage, held = Path(marker["path"]), Path(marker["held"])
                self.assertTrue(os.path.lexists(stage), "cold recovery cleaned another owner's stage")
                self.assertEqual(marker["ino"], stage.lstat().st_ino)
                self.assertEqual(before["common"], os.readlink(held))
                if kind == "publish":
                    self.assertEqual("independent-foreign-target", os.readlink(stage))
                else:
                    self.assertEqual(b"independent foreign stage evidence\n", stage.read_bytes())
                    self.assertEqual(wal, (root / "held-original-link-stage-wal").read_bytes())
                self.assertFalse((root / "service-started").exists())

    def _assert_real_setup_restore_stage_authority_holds(self, surface, kind):
        with tempfile.TemporaryDirectory(prefix="r22-real-restore-stage-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, _before, *_rest = fixture
            for relative in (
                "tools/lldp-analyze-tool/99-output-p2p",
                "tools/lldp-analyze-tool/99-output-monitor",
                "tools/ibdiagnet-analyze-tool/99-output-p2p",
            ):
                leaf = root / relative
                if leaf.is_dir() and not leaf.is_symlink():
                    shutil.rmtree(leaf)
            latest = q / "99-output-eth/latest"
            if surface == "q-local":
                output = latest.parent
                output.mkdir()
                releases = [output / value for value in (
                    "20261005_110000_combine", "20261005_120000_combine",
                )]
                for release in releases:
                    release.mkdir()
                    (release / ".published-complete").write_bytes(b"complete\n")
                    (release / "leaf.yaml").write_bytes(b"hostname: independent-q\n")
                latest.symlink_to(releases[0].name)
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=False, auto_prepare_journal=True,
                kill_after_replace=manifest,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            target = manifest if surface == "manifest" else latest if surface == "q-local" else common
            target_owner = target.lstat().st_ino
            target_value = target.read_bytes() if surface == "manifest" else os.readlink(target)
            journal = root / "ztp/.setup-switch-journal.json"
            wal = journal.read_bytes()
            recovery = r'''
import importlib.util, json, os, stat, sys
from pathlib import Path
from types import SimpleNamespace
root, surface, kind, target = Path(sys.argv[1]), sys.argv[2], sys.argv[3], Path(sys.argv[4])
spec = importlib.util.spec_from_file_location('cold_setup_stage_guard', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
journal = root / 'ztp/.setup-switch-journal.json'
body = json.loads(journal.read_bytes())['body']
injected, symlink, sync = [], Path.symlink_to, os.fsync
def drift():
    if kind == 'wal':
        value, owner = journal.read_bytes(), journal.stat().st_ino
        journal.rename(root / 'held-original-restore-stage-wal')
        journal.write_bytes(value)
        journal.chmod(0o600)
        assert owner != journal.stat().st_ino
    else:
        parent = Path(body['q_parent_preimage']['path'])
        assert not os.path.lexists(parent)
        parent.parent.mkdir(parents=True, exist_ok=True)
        parent.write_bytes(b'independent Q parent during cold stage\n')
    injected.append(True)
def staged_link(path, raw, *args, **kwargs):
    result = symlink(path, raw, *args, **kwargs)
    if surface != 'manifest' and path.parent == target.parent and path.name.startswith('.' + target.name + '.setup-rollback.') and not injected:
        drift()
    return result
def staged_file(fd):
    result = sync(fd)
    if surface == 'manifest' and stat.S_ISREG(os.fstat(fd).st_mode) and not injected:
        drift()
    return result
Path.symlink_to, os.fsync = staged_link, staged_file
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal, deployment_lock_descriptor=fd,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as error:
        expected = 'journal changed' if kind == 'wal' else 'Q parent changed'
        assert injected == [True] and expected in str(error), str(error)
        print('R22_COLD_SETUP_RESTORE_STAGE_HELD')
        raise SystemExit(7)
    else:
        raise AssertionError('cold stage consumed changed authority')
finally:
    load.release_deployment_lock(fd)
'''
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root), surface, kind, str(target)],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(7, second.returncode, second.stdout + second.stderr)
            self.assertIn("R22_COLD_SETUP_RESTORE_STAGE_HELD", second.stdout)
            self.assertEqual(target_owner, target.lstat().st_ino,
                             "cold staged recovery replaced the target after authority drift")
            self.assertEqual(target_value, target.read_bytes() if surface == "manifest" else os.readlink(target))
            self.assertEqual(wal, journal.read_bytes())
            self.assertFalse(list(target.parent.glob(f".{target.name}.setup-rollback.*")))
            self.assertFalse(list(manifest.parent.glob(f".{manifest.name}.rollback.*")))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_shared_link_cold_restore_stage_holds_wal_and_q(self):
        for kind in ("wal", "q"):
            with self.subTest(kind=kind):
                self._assert_real_setup_restore_stage_authority_holds("shared", kind)

    def test_real_setup_q_local_cold_restore_stage_holds_wal_and_q(self):
        for kind in ("wal", "q"):
            with self.subTest(kind=kind):
                self._assert_real_setup_restore_stage_authority_holds("q-local", kind)

    def test_real_setup_manifest_cold_restore_stage_holds_wal_and_q(self):
        for kind in ("wal", "q"):
            with self.subTest(kind=kind):
                self._assert_real_setup_restore_stage_authority_holds("manifest", kind)

    def test_real_setup_q_manifest_publication_stage_holds_without_eager_rollback(self):
        for kind in ("wal", "q"):
            with self.subTest(kind=kind):
                with tempfile.TemporaryDirectory(prefix="r22-q-manifest-publish-stage-") as name:
                    fixture = self._private_setup_switch_fixture(Path(name))
                    root, _p, _q, common, _p_only, _q_only, manifest, before, hook, _barrier, _home = fixture
                    for relative in (
                        "tools/lldp-analyze-tool/99-output-p2p",
                        "tools/lldp-analyze-tool/99-output-monitor",
                        "tools/ibdiagnet-analyze-tool/99-output-p2p",
                    ):
                        leaf = root / relative
                        if leaf.is_dir() and not leaf.is_symlink():
                            shutil.rmtree(leaf)
                    hook_file = hook / "sitecustomize.py"
                    hook_file.write_text(hook_file.read_text() + f'''
import json
import stat
root = Path({str(root)!r})
manifest = root / 'ztp/.setup_manifest'
kind = {kind!r}
observed = root / 'q-manifest-stage-evidence.json'
original_sync = os.fsync
def drift_manifest_publication(fd):
    result = original_sync(fd)
    info = os.fstat(fd)
    stages = (*manifest.parent.glob('.setup_manifest.*'),
              *manifest.parent.glob('.' + manifest.name + '.rollback.*'))
    if (not observed.exists() and stat.S_ISREG(info.st_mode)
            and any(path.lstat().st_ino == info.st_ino for path in stages)):
        journal = root / 'ztp/.setup-switch-journal.json'
        body = json.loads(journal.read_bytes())['body']
        common = root / 'ztp/config/cumulus/template/01-global.yaml'
        evidence = {{'wal': journal.read_text(), 'manifest_ino': manifest.lstat().st_ino,
                    'manifest': manifest.read_text(), 'common_ino': common.lstat().st_ino,
                    'common': os.readlink(common)}}
        if kind == 'wal':
            journal.rename(root / 'held-original-publish-wal')
            journal.write_text(evidence['wal'])
            journal.chmod(0o600)
            evidence['foreign_ino'] = journal.lstat().st_ino
        else:
            marker = Path(body['new_project']['path']) / '99-output-ztp/current-release.json'
            marker.parent.mkdir(parents=True, exist_ok=True)
            marker.write_bytes(b'independent Q parent appeared during publication\\n')
            evidence['foreign_ino'] = marker.lstat().st_ino
        observed.write_text(json.dumps(evidence))
    return result
os.fsync = drift_manifest_publication
''', encoding="utf-8")
                    result = self._run_private_load_setup_switch(
                        fixture, prepare_journal=False, auto_prepare_journal=True,
                        kill_after_replace=root / "never-replaced",
                    )
                    observed = root / "q-manifest-stage-evidence.json"
                    self.assertTrue(observed.is_file(), result.stdout + result.stderr)
                    evidence = json.loads(observed.read_text())
                    self.assertEqual(24, result.returncode, result.stdout + result.stderr)
                    self.assertEqual(before["manifest"], manifest.read_bytes())
                    self.assertEqual(evidence["manifest_ino"], manifest.lstat().st_ino,
                                     "publication or eager rollback replaced P manifest after HOLD")
                    self.assertEqual((evidence["common_ino"], evidence["common"]),
                                     (common.lstat().st_ino, os.readlink(common)),
                                     "eager rollback consumed the changed total authority")
                    journal = root / "ztp/.setup-switch-journal.json"
                    self.assertEqual(evidence["wal"], journal.read_text())
                    authority = (journal if kind == "wal" else fixture[2] / "99-output-ztp/current-release.json")
                    self.assertEqual(evidence["foreign_ino"], authority.lstat().st_ino)
                    self.assertFalse(list(manifest.parent.glob(".setup_manifest.*")))
                    self.assertFalse(list(manifest.parent.glob(f".{manifest.name}.rollback.*")))
                    self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_manifest_publication_preserves_foreign_staged_owner(self):
        for kind in ("inode", "bytes"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(prefix="r22-manifest-stage-owner-") as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, _q, common, _p_only, _q_only, manifest, before, hook, _barrier, _home = fixture
                for relative in (
                    "tools/lldp-analyze-tool/99-output-p2p",
                    "tools/lldp-analyze-tool/99-output-monitor",
                    "tools/ibdiagnet-analyze-tool/99-output-p2p",
                ):
                    leaf = root / relative
                    if leaf.is_dir() and not leaf.is_symlink():
                        shutil.rmtree(leaf)
                hook_file = hook / "sitecustomize.py"
                hook_file.write_text(hook_file.read_text() + f'''
import json, stat
root, kind = Path({str(root)!r}), {kind!r}
manifest = root / 'ztp/.setup_manifest'
observed = root / 'manifest-stage-owner-evidence.json'
original_sync = os.fsync
def change_stage_owner(fd):
    result = original_sync(fd)
    info = os.fstat(fd)
    if stat.S_ISREG(info.st_mode) and not observed.exists():
        for path in manifest.parent.glob('.' + manifest.name + '.rollback.*'):
            if path.lstat().st_ino != info.st_ino:
                continue
            journal = root / 'ztp/.setup-switch-journal.json'
            common = root / 'ztp/config/cumulus/template/01-global.yaml'
            evidence = {{'manifest_ino': manifest.lstat().st_ino,
                        'common_ino': common.lstat().st_ino, 'common': os.readlink(common),
                        'wal': journal.read_text(), 'stage': str(path)}}
            if kind == 'inode':
                path.rename(root / 'held-original-manifest-stage')
            path.write_bytes(b'independent foreign Q manifest stage\\n')
            path.chmod(0o644)
            evidence['stage_ino'] = path.lstat().st_ino
            observed.write_text(json.dumps(evidence))
            break
    return result
os.fsync = change_stage_owner
''', encoding="utf-8")
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=False, auto_prepare_journal=True,
                    kill_after_replace=root / "never-replaced",
                )
                observed = root / "manifest-stage-owner-evidence.json"
                self.assertTrue(observed.is_file(), result.stdout + result.stderr)
                evidence = json.loads(observed.read_text())
                self.assertEqual((evidence["manifest_ino"], before["manifest"]),
                                 (manifest.lstat().st_ino, manifest.read_bytes()),
                                 "real publisher replaced P using a foreign stage")
                self.assertEqual(24, result.returncode, result.stdout + result.stderr)
                self.assertEqual((evidence["common_ino"], evidence["common"]),
                                 (common.lstat().st_ino, os.readlink(common)))
                self.assertEqual(evidence["wal"], (root / "ztp/.setup-switch-journal.json").read_text())
                stage = Path(evidence["stage"])
                self.assertEqual(evidence["stage_ino"], stage.lstat().st_ino,
                                 "HOLD cleanup deleted or consumed the foreign stage")
                self.assertEqual(b"independent foreign Q manifest stage\n", stage.read_bytes())
                self.assertFalse((root / "service-started").exists())

    def test_real_setup_parent_commit_cannot_bypass_wal_via_legacy_restore(self):
        with tempfile.TemporaryDirectory(prefix="r22-legacy-parent-fence-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, q_only, manifest, before, _hook, _barrier, _home = fixture
            for relative in ("tools/lldp-analyze-tool/99-output-p2p",
                             "tools/lldp-analyze-tool/99-output-monitor",
                             "tools/ibdiagnet-analyze-tool/99-output-p2p"):
                leaf = root / relative
                if leaf.is_dir() and not leaf.is_symlink():
                    shutil.rmtree(leaf)
            first = self._run_private_load_setup_switch(fixture, prepare_journal=False,
                auto_prepare_journal=True, kill_after_replace=root / "never-replaced")
            self.assertEqual(25, first.returncode, first.stdout + first.stderr)
            # SETUP points at LOAD's not-yet-generated P2P artifacts. Supply
            # private generator dependencies, not rollback expected values.
            for value in manifest.read_text().splitlines()[1:]:
                link = Path(value)
                if link.is_symlink() and not link.exists():
                    dependency = link.resolve(strict=False)
                    dependency.relative_to(q / "99-output-p2p")
                    dependency.parent.mkdir(parents=True, exist_ok=True)
                    dependency.write_bytes(b"{}\n")
            journal = root / "ztp/.setup-switch-journal.json"
            mixed = (common.lstat().st_ino, os.readlink(common), q_only.lstat().st_ino,
                     os.readlink(q_only), manifest.lstat().st_ino, manifest.read_bytes(), journal.read_bytes())
            driver = f'''
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root, p, q = map(Path, sys.argv[1:4])
spec = importlib.util.spec_from_file_location('legacy_parent_fence_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
state = load.ActiveSetupState(project=p, manifest={before['manifest']!r}, manifest_mode=0o644,
    links={{root / 'ztp/config/cumulus/template/01-global.yaml': {before['common']!r},
            root / 'ztp/config/publickey/p-only.pub': {before['p_only']!r}}})
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    candidate = load.prepare_current_release(q, {{'schema_version': 1, 'release_id': 'independent-Q-parent'}})
    load.commit_prepared_release(candidate)
    try:
        load.restore_active_setup_state(state, current_project=q, deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as exc:
        print('EXPECTED LEGACY WAL HOLD:', exc)
        raise SystemExit(7)
finally:
    load.release_deployment_lock(descriptor)
'''
            second = subprocess.run([sys.executable, "-B", "-c", driver, str(root), str(p), str(q)],
                                    cwd=root, capture_output=True, text=True, check=False)
            parent = q / "99-output-ztp/current-release.json"
            self.assertEqual("independent-Q-parent", json.loads(parent.read_bytes())["release_id"])
            self.assertEqual(mixed, (common.lstat().st_ino, os.readlink(common), q_only.lstat().st_ino if q_only.is_symlink() else None,
                                    os.readlink(q_only) if q_only.is_symlink() else None, manifest.lstat().st_ino,
                                    manifest.read_bytes(), journal.read_bytes()),
                             "legacy API must not roll P back over a pending WAL/visible Q parent")
            self.assertEqual(7, second.returncode, second.stdout + second.stderr)
            self.assertIn("pending setup switch WAL:", second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(p_only))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_manifest_commit_cold_recovers_p_and_original_logs(self):
        """Q manifest publication and successful SETUP must retain P recovery bytes."""
        for death_at_manifest in (True, False):
            with self.subTest(death_at_manifest=death_at_manifest):
                with tempfile.TemporaryDirectory(prefix="r22-q-manifest-cold-") as name:
                    fixture = self._private_setup_switch_fixture(Path(name))
                    root, _p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
                    # The source tools tree contains real output placeholders.
                    # This initialized switch fixture needs absent link leaves,
                    # not unrelated copies blocking SETUP before manifest commit.
                    for relative in (
                        "tools/lldp-analyze-tool/99-output-p2p",
                        "tools/lldp-analyze-tool/99-output-monitor",
                        "tools/ibdiagnet-analyze-tool/99-output-p2p",
                    ):
                        leaf = root / relative
                        if leaf.is_dir() and not leaf.is_symlink():
                            shutil.rmtree(leaf)
                    logs = root / "infra/logs"
                    logs.mkdir()
                    latest = q / "99-output-eth/latest"
                    release = q / "99-output-eth/20261005_120000_combine"
                    for path in (latest, release):
                        path.mkdir(parents=True)
                        (path / ".published-complete").write_bytes(b"complete\n")
                        (path / "leaf.yaml").write_bytes(b"hostname: leaf\n")
                    latest_identity = (latest.stat().st_dev, latest.stat().st_ino)
                    payload = b"P-era logs must survive Q setup success\n"
                    (logs / "audit.log").write_bytes(payload)
                    identity = (logs.stat().st_dev, logs.stat().st_ino)
                    first = self._run_private_load_setup_switch(
                        fixture, prepare_journal=False, auto_prepare_journal=True,
                        kill_after_replace=manifest if death_at_manifest else root / "never-replaced",
                    )
                    self.assertEqual(23 if death_at_manifest else 25, first.returncode,
                                     first.stdout + first.stderr)
                    self.assertNotEqual(before["manifest"], manifest.read_bytes())
                    self.assertTrue(logs.is_symlink())
                    journal = root / "ztp/.setup-switch-journal.json"
                    body = json.loads(journal.read_bytes())["body"]
                    backup = Path(body["infra_logs"][0]["backup"])
                    self.assertEqual(identity, (backup.stat().st_dev, backup.stat().st_ino))
                    latest_backup = Path(body["directories"][0]["backup"])
                    self.assertEqual(latest_identity, (
                        latest_backup.stat().st_dev, latest_backup.stat().st_ino,
                    ))
                    recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_q_manifest_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
                    second = subprocess.run(
                        [sys.executable, "-B", "-c", recovery, str(root)],
                        cwd=root, capture_output=True, text=True, check=False,
                    )
                    self.assertEqual(0, second.returncode, second.stdout + second.stderr)
                    self.assertEqual(before["manifest"], manifest.read_bytes())
                    self.assertEqual(before["common"], os.readlink(common))
                    self.assertEqual(before["p_only"], os.readlink(p_only))
                    self.assertFalse(os.path.lexists(q_only))
                    self.assertFalse(logs.is_symlink())
                    self.assertFalse(latest.is_symlink())
                    self.assertEqual(latest_identity, (latest.stat().st_dev, latest.stat().st_ino))
                    self.assertEqual(b"hostname: leaf\n", (latest / "leaf.yaml").read_bytes())
                    self.assertEqual(identity, (logs.stat().st_dev, logs.stat().st_ino))
                    self.assertEqual(payload, (logs / "audit.log").read_bytes())
                    self.assertFalse(os.path.lexists(q / "99-output-infra/audit.log"))
                    self.assertFalse(os.path.lexists(root / "infra/.logs-migration.json"))
                    self.assertFalse(journal.exists())
                    self.assertFalse((root / "service-started").exists())

    def test_real_load_setup_rejects_q_parent_owner_changed_during_wal_freeze(self):
        """LOAD must refuse before launching the real SETUP with an unstable Q preimage."""
        with tempfile.TemporaryDirectory(prefix="r22-q-parent-freeze-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, q_only, manifest, before, _hook, _barrier, _home = fixture
            parent = q / "99-output-ztp/current-release.json"
            parent.parent.mkdir()
            payload = b'{"release_id":"independent-Q-parent"}\n'
            parent.write_bytes(payload)
            operation = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root, p, q = map(Path, sys.argv[1:])
spec = importlib.util.spec_from_file_location('q_parent_freeze_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
parent = q / '99-output-ztp/current-release.json'
reader = load._read_rollback_regular
mutated = False
setup_invoked = False
run = load.run
def replace_after_read(path, **kwargs):
    global mutated
    result = reader(path, **kwargs)
    if Path(path) == parent and not mutated:
        mutated = True
        parent.rename(parent.with_name('operator-retained-parent.json'))
        parent.write_bytes(b'{"release_id":"independent-Q-parent"}\\n')
    return result
def observe_setup(command, **kwargs):
    global setup_invoked
    setup_invoked = True
    return run(command, **kwargs)
load._read_rollback_regular = replace_after_read
load.run = observe_setup
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.activate_project(
            q, q / 'p2p.xlsx', strict=False, no_upgrade=True,
            deployment_lock_descriptor=descriptor, host_role='workstation',
            prepare_switch_wal=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as error:
        rejection = str(error)
    else:
        raise AssertionError('LOAD adopted the replacement Q parent')
    assert mutated
    assert not setup_invoked, 'unstable preimage must stop before SETUP'
    assert 'Q parent' in rejection, rejection
    assert not (root / 'ztp/.setup-switch-journal.json').exists(), 'unstable preimage must not publish a WAL'
    print('R22_Q_PARENT_FREEZE_HELD')
finally:
    load.release_deployment_lock(descriptor)
"""
            result = subprocess.run(
                [sys.executable, "-B", "-c", operation, str(root), str(p), str(q)],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, result.returncode, result.stdout + result.stderr)
            self.assertIn("R22_Q_PARENT_FREEZE_HELD", result.stdout)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertEqual(payload, parent.read_bytes())
            self.assertEqual(payload, parent.with_name("operator-retained-parent.json").read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_then_uncertain_parent_retains_dhcp_a_evidence(self):
        """SETUP success must not license deleting A backups behind uncertain B."""
        self._real_setup_uncertain_parent_dhcp_evidence(reenter=False)

    def test_real_setup_then_uncertain_parent_reentry_refuses_new_dhcp_writes(self):
        """An already-visible uncertain B cannot authorize a second installation."""
        self._real_setup_uncertain_parent_dhcp_evidence(reenter=True)

    def test_real_setup_then_uncertain_dhcp_transaction_fences_fresh_load_process(self):
        """A new LOAD process must honor retained disk evidence, not candidate flags."""
        self._real_setup_uncertain_parent_dhcp_evidence(reenter=False, cold=True)

    def test_real_setup_then_fresh_load_inspects_disk_dhcp_preimage_without_activation(self):
        self._real_setup_uncertain_parent_dhcp_evidence(reenter=False, cold=True, inspect_preimage=True)

    def test_real_setup_then_fresh_load_inspection_holds_on_parent_directory_rebind(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, cold=True, inspect_parent_rebind=True, outcomes=("fsync",),
        )

    def test_real_setup_then_parent_success_retains_dhcp_a_until_whole_b_acceptance(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, cold=True, inspect_preimage=True, outcomes=("success-setup-wal",),
        )

    def test_real_setup_then_failed_dhcp_a_restore_preserves_transaction_evidence(self):
        """A partial precommit restore cannot erase the only intact A copies."""
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("restore-return", "restore-raise"),
        )

    def test_real_setup_then_silent_wrong_dhcp_a_restore_holds_evidence(self):
        """Command exit0 cannot certify restored bytes, metadata or leaf type."""
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("restore-bytes", "restore-mode", "restore-link"),
        )

    def test_real_setup_then_verified_dhcp_a_restore_preserves_old_leaf_shapes(self):
        """An intact A link and originally absent leaf still restore and clean up."""
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("precommit-shapes",),
        )

    def test_real_setup_then_dhcp_a_restore_sync_failure_retains_evidence(self):
        """Verified A bytes must not license cleanup after a sync failure."""
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("restore-file-sync", "restore-dir-open", "restore-dir-sync"),
        )

    def test_real_setup_then_unverified_dhcp_a_backup_blocks_first_b_install(self):
        """An exit0 copy with a wrong backup cannot authorize overwriting A."""
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("backup-bytes", "backup-mode", "backup-link", "backup-read"),
        )

    def test_real_setup_then_dhcp_a_backup_sync_failure_blocks_first_b_install(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("backup-file-sync", "backup-dir-sync", "backup-transaction-sync", "backup-runtime-sync"),
        )

    def test_real_setup_then_dhcp_a_backup_syncs_before_first_b_install(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("backup-sync-success",),
        )

    def test_real_setup_then_dhcp_persistent_a_preimage_precedes_b_write(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("preimage-before-write",),
        )

    def test_real_setup_publication_wal_survives_parent_writer_death(self):
        self._real_setup_publication_writer_death("parent")

    def test_real_setup_publication_cold_resume_after_report_writer_death(self):
        self._real_setup_publication_writer_death("report")

    def test_real_setup_publication_cold_resume_after_selector_writer_death(self):
        self._real_setup_publication_writer_death("selector")

    def test_real_setup_publication_cold_cli_resumes_without_regeneration(self):
        self._real_setup_publication_writer_death("parent", resume_cli=True)

    def test_real_setup_publication_cold_reader_holds_foreign_q_manifest_and_flat_runtime(self):
        for fault in ("manifest-owner", "flat-runtime"):
            with self.subTest(fault=fault):
                self._real_setup_publication_writer_death("parent", late_fault=fault)

    def test_real_setup_publication_cold_stopped_b_preflight_retains_original_wal(self):
        self._real_setup_publication_writer_death("selector", b_preflight="pass")

    def test_real_setup_publication_cold_stopped_b_preflight_late_input_drift_holds(self):
        self._real_setup_publication_writer_death("selector", b_preflight="late-input")

    def test_real_setup_publication_cold_b_activation_retains_original_wal(self):
        self._real_setup_publication_writer_death("selector", b_preflight="activate-pass")

    def test_real_setup_publication_cold_b_health_failure_stops_without_a_restore(self):
        self._real_setup_publication_writer_death("selector", b_preflight="activate-fail")

    def _real_setup_publication_writer_death(self, death_phase, *, resume_cli=False, late_fault=None,
                                            b_preflight=None):
        with tempfile.TemporaryDirectory(prefix="r22-publication-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, _common, _p_only, _q_only, manifest, _before, *_rest = fixture
            if b_preflight and b_preflight.startswith("activate-"):
                from test_cases.test_ssh_key_preparation_contract import public_line
                (q / "mgmt-server.pub").write_bytes(public_line("ssh-ed25519", 47, "management"))
                document = yaml.safe_load((q / "01-global.yaml").read_text())
                document["common"]["mgmt"]["http"].update(
                    status="enabled", address="192.0.2.40", http_root=str(root))
                document["switches"][0]["eth"]["version"] = "5.12.0"
                (q / "01-global.yaml").write_text(yaml.safe_dump(document, sort_keys=False))
                image = root / "ztp/image/cumulus/cumulus-linux-5.12.0-mlx-amd64.bin"
                image.parent.mkdir(parents=True, exist_ok=True)
                image.write_bytes(b"#!/bin/sh\n" + b"x" * 1048576)
            for relative in ("tools/lldp-analyze-tool/99-output-p2p",
                             "tools/lldp-analyze-tool/99-output-monitor",
                             "tools/ibdiagnet-analyze-tool/99-output-p2p"):
                path = root / relative
                if path.is_dir() and not path.is_symlink():
                    shutil.rmtree(path)
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=False, auto_prepare_journal=True,
                first_link=root / "never-created-link",
            )
            self.assertEqual(25, first.returncode, first.stdout + first.stderr)
            if b_preflight and b_preflight.startswith("activate-"):
                # The report fixture's two independent literal fingerprints
                # correspond to these public wire bytes, not producer output.
                (q / "laptop.pub").write_bytes(public_line("ssh-ed25519", 31, "laptop"))
            api = load_script("r22_publication_writer_probe", root / "DAY0-Prepare/11-load.py")
            self.assertTrue(callable(getattr(api, "register_setup_switch_publication_bundle", None)),
                            "real LOAD cannot register its publication group in the original SETUP WAL")
            q_manifest = manifest.read_bytes()
            operation = r'''
import os, signal, sys
from pathlib import Path
from types import SimpleNamespace
root, q, test_root = map(Path, sys.argv[1:4])
sys.path.insert(0, str(test_root))
from test_cases.module_loader import load_script
from test_cases.test_load_release_transaction import private_publication_candidates
api = load_script('r22_bound_publication_child', root / 'DAY0-Prepare/11-load.py')
api.active_managed_services = lambda *args, **kwargs: ()
descriptor = api.acquire_deployment_lock(exclusive=True)
guard = dict(deployment_lock_descriptor=descriptor, services_quiesced=True,
             runtime_backend=SimpleNamespace(name='supervisor'))
candidates = private_publication_candidates(api, q)
image_options = {}
if sys.argv[5].startswith('activate-'):
    image_options['images'] = {'eth': root / 'ztp/image/cumulus/cumulus-linux-5.12.0-mlx-amd64.bin'}
token = api.register_setup_switch_publication_bundle(*candidates, guard=guard, **image_options)
api.commit_prepared_release(candidates[0], require_authority=lambda:
    api.require_setup_switch_publication_bundle(token, guard=guard))
if sys.argv[4] in ('report', 'selector'):
    api.generation_report.commit_prepared_generation_report(candidates[1], require_authority=lambda:
        api.require_setup_switch_publication_bundle(token, guard=guard, phase='report'))
if sys.argv[4] == 'selector':
    import json
    staged = json.loads(token['payload'])['body']['publication_bundle']['selector']['staged']
    api.commit_pinned_generation(q, candidates[2], expected_current=None, prepared_selector=staged,
        require_authority=lambda: api.require_setup_switch_publication_bundle(token, guard=guard, phase='selector'))
os.kill(os.getpid(), signal.SIGKILL)
'''
            child = subprocess.run([sys.executable, "-B", "-c", operation,
                                    str(root), str(q), str(ROOT), death_phase, b_preflight or "none"], cwd=root,
                                   capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(-signal.SIGKILL, child.returncode, child.stdout + child.stderr)
            journal = root / "ztp/.setup-switch-journal.json"
            wal = journal.read_bytes()
            binding = json.loads(wal)["body"]["publication_bundle"]
            parent_path = Path(binding["parent"]["destination"])
            parent_bytes = parent_path.read_bytes()
            self.assertEqual(base64.b64decode(binding["parent"]["candidate"]["payload"]), parent_bytes)
            self.assertEqual(death_phase == "parent", Path(binding["report"]["temporary"]).exists())
            self.assertEqual(death_phase != "parent", Path(binding["report"]["destination"]).exists())
            descriptor = api.acquire_deployment_lock(exclusive=True)
            try:
                with mock.patch.object(api, "active_managed_services", return_value=()):
                    with self.assertRaisesRegex(api.LoadError, "publication.*HOLD"):
                        api.recover_setup_switch_journal(
                            journal, restore_runtime_from_wal=True, requested_new_project=q,
                            deployment_lock_descriptor=descriptor, services_quiesced=True,
                            runtime_backend=SimpleNamespace(name="supervisor"))
            finally:
                api.release_deployment_lock(descriptor)
            self.assertEqual(wal, journal.read_bytes())
            self.assertEqual(parent_bytes, parent_path.read_bytes())
            self.assertEqual(q_manifest, manifest.read_bytes())
            self.assertFalse((root / "service-started").exists())
            if late_fault == "manifest-owner":
                held = manifest.with_name("independent-held-shared-manifest")
                manifest.rename(held)
                manifest.write_bytes(q_manifest)
                manifest.chmod(0o644)
                self.assertNotEqual(held.lstat().st_ino, manifest.lstat().st_ino)
            elif late_fault == "flat-runtime":
                runtime = root / "ztp/ztp-bootstrap_oob.sh"
                runtime.write_bytes(b"independent unregistered bootstrap\n")
                runtime.chmod(0o755)
            resume = r'''
import json, sys
from pathlib import Path
from types import SimpleNamespace
root, q, test_root = map(Path, sys.argv[1:4])
sys.path.insert(0, str(test_root))
from test_cases.module_loader import load_script
api = load_script('r22_cold_publication_reader', root / 'DAY0-Prepare/11-load.py')
api.active_managed_services = lambda *args, **kwargs: ()
def forbidden(*args, **kwargs):
    raise AssertionError('cold publication attempted generation or a service mutation')
api.start_services = api.quiesce_services = api.generate_configs = forbidden
if sys.argv[4] == 'cli':
    api.runtime_os = lambda: 'Linux'
    api.service_runtime_backend = lambda: SimpleNamespace(name='supervisor')
    api.parse_args = lambda argv: api._build_parser().parse_args(argv)
    status = api.main([str(q), '--host-role', 'management-server', '--resume-setup-publication'])
    assert status == 0, 'explicit cold resume CLI did not complete'
    print('R22_RESUME='+json.dumps(dict(before='report', after='published',
                                     result='published-stopped-wal-retained')))
    sys.exit(0)
descriptor = api.acquire_deployment_lock(exclusive=True)
token = None
try:
    guard = dict(deployment_lock_descriptor=descriptor, services_quiesced=True,
                 runtime_backend=SimpleNamespace(name='supervisor'))
    token = api.open_setup_switch_publication_bundle(q, guard=guard)
    before = token['phase']
    result = api.resume_setup_switch_publication_bundle(token, guard=guard)
    if sys.argv[5].startswith('activate-'):
        parent = json.loads((q / '99-output-ztp/current-release.json').read_bytes())
        paths = parent['input_paths']
        inputs = api.ProjectInputs(global_file=q / paths['global'], devices_file=q / paths['devices'],
            subnet_file=q / paths['subnet'], p2p_file=q / paths['p2p'], device_types=frozenset(),
            pubkeys=(), settings=api.load_global(q / paths['global']))
        active, actions = set(), []
        api.active_managed_services = lambda *a, **kw: tuple(sorted(active))
        api.plan_local_dhcp_runtime = lambda *a, **kw: SimpleNamespace(endpoint_ips=(), listener_names=())
        api.validate_management_host = lambda *a, **kw: True
        # External auth/Apache/live HTTP are hermetic; real report inspection,
        # normal preflight, Q owner gates and real SETUP stay on this path.
        api.verify_control_auth = lambda **kw: None
        api.verify_apache_publication_boundary = lambda **kw: None
        def command(argv, **kw):
            assert argv[-2:] == ['apache2ctl', 'configtest'], argv
        api.run = command
        def service(backend, action, name, **kw):
            assert action == 'start' and name == 'apache2', (action, name)
            active.add(name)
            actions.append((action, name))
        api._runtime_service_action = service
        def health(*a, **kw):
            if sys.argv[5] == 'activate-fail':
                raise api.LoadError('independent cold B health rejection')
        api.verify_http_publication = health
        def stop(*a, **kw):
            active.clear()
            actions.append(('stop', 'all'))
        api.quiesce_services = stop
        api.restore_cross_project_rollback_context = forbidden
        original_wal = token['payload']
        try:
            image = root / 'ztp/image/cumulus/cumulus-linux-5.12.0-mlx-amd64.bin'
            status = api.activate_setup_switch_b_publication(token, guard=guard, inputs=inputs,
                images={'eth': image}, confirmed=True)
        except api.LoadError as exc:
            assert sys.argv[5] == 'activate-fail' and 'B health rejection' in str(exc), str(exc)
            assert active == set() and actions == [('start', 'apache2'), ('stop', 'all')], actions
            status = 'b-health-failed-stopped-wal-retained'
        else:
            assert sys.argv[5] == 'activate-pass', 'B failure was accepted'
            assert active == {'apache2'} and actions == [('start', 'apache2')], actions
            assert status == 'b-health-passed-wal-retained', status
        assert original_wal == token['journal'].read_bytes()
        print('R22_PREFLIGHT=' + status)
    elif sys.argv[5] != 'none':
        parent = json.loads((q / '99-output-ztp/current-release.json').read_bytes())
        paths = parent['input_paths']
        inputs = api.ProjectInputs(global_file=q / paths['global'], devices_file=q / paths['devices'],
            subnet_file=q / paths['subnet'], p2p_file=q / paths['p2p'], device_types=frozenset(),
            pubkeys=(), settings=api.load_global(q / paths['global']))
        trace = []
        def plan(*args, **kwargs):
            trace.append('plan')
            if sys.argv[5] == 'late-input':
                inputs.devices_file.write_bytes(b'independent drift after first helper\n')
            return SimpleNamespace(listener_names=())
        api.plan_local_dhcp_runtime = plan
        api.supervisor_service_availability = lambda *a, **kw: trace.append('availability') or (False, False)
        api.run = forbidden
        original_wal = token['payload']
        try:
            status = api.preflight_setup_switch_b_publication(token, guard=guard, inputs=inputs, images={})
        except api.LoadError as exc:
            assert sys.argv[5] == 'late-input' and 'HOLD' in str(exc), str(exc)
            assert trace == ['plan'], trace
            status = 'held-still-stopped-wal-retained'
        else:
            assert sys.argv[5] == 'pass', 'late-input drift was accepted'
            assert trace == ['plan', 'availability'], trace
            assert status == 'b-preflight-passed-still-stopped-wal-retained', status
        assert original_wal == token['journal'].read_bytes()
        print('R22_PREFLIGHT=' + status)
    print('R22_RESUME='+json.dumps(dict(before=before, after=token['phase'], result=result)))
finally:
    api.close_setup_switch_publication_bundle(token)
    api.release_deployment_lock(descriptor)
'''
            recovered = subprocess.run([sys.executable, "-B", "-c", resume,
                                        str(root), str(q), str(ROOT), "cli" if resume_cli else "api",
                                        b_preflight or "none"], cwd=root,
                                       capture_output=True, text=True, check=False, timeout=30)
            if late_fault is not None:
                self.assertEqual(1, recovered.returncode, recovered.stdout + recovered.stderr)
                self.assertIn("HOLD", recovered.stderr)
                self.assertEqual(wal, journal.read_bytes())
                self.assertEqual(parent_bytes, parent_path.read_bytes())
                self.assertEqual(q_manifest, manifest.read_bytes())
                self.assertTrue(Path(binding["report"]["temporary"]).is_file())
                self.assertFalse(Path(binding["report"]["destination"]).exists())
                self.assertFalse((q / "99-output-ztp/current").exists())
                self.assertFalse((root / "service-started").exists())
                if late_fault == "flat-runtime":
                    self.assertEqual(b"independent unregistered bootstrap\n", runtime.read_bytes())
                else:
                    self.assertEqual(q_manifest, held.read_bytes())
                return
            self.assertEqual(0, recovered.returncode, recovered.stdout + recovered.stderr)
            if b_preflight:
                expected = {"pass": "b-preflight-passed-still-stopped-wal-retained",
                    "late-input": "held-still-stopped-wal-retained",
                    "activate-pass": "b-health-passed-wal-retained",
                    "activate-fail": "b-health-failed-stopped-wal-retained"}[b_preflight]
                self.assertIn("R22_PREFLIGHT=" + expected, recovered.stdout)
            payload = json.loads(next(line.split("=", 1)[1] for line in recovered.stdout.splitlines()
                                      if line.startswith("R22_RESUME=")))
            self.assertEqual({"before": {"parent": "report", "report": "selector", "selector": "published"}[death_phase],
                              "after": "published", "result": "published-stopped-wal-retained"}, payload)
            self.assertEqual(wal, journal.read_bytes())
            self.assertEqual(parent_bytes, parent_path.read_bytes())
            self.assertEqual(q_manifest, manifest.read_bytes())
            self.assertEqual(binding["selector"]["staged"]["identity"][1],
                             (q / "99-output-ztp/current").lstat().st_ino)
            self.assertEqual(binding["report"]["candidate"]["sha256"],
                             hashlib.sha256(Path(binding["report"]["destination"]).read_bytes()).hexdigest())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_then_load_binds_dhcp_preimage_before_b_install(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, bind_setup=True, outcomes=("bound-setup-before-write", "bound-setup-precommit"),
        )

    def test_real_setup_then_load_holds_replaced_grouped_dhcp_authorities(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, bind_setup=True,
            outcomes=("bound-setup-register-wal", "bound-setup-register-receipt", "bound-setup-late-wal", "bound-setup-parent-wal"),
        )

    def test_real_setup_then_dhcp_writer_death_cold_group_recovery_restores_stopped_a(self):
        self._real_dhcp_writer_death_group_recovery()

    def test_real_setup_then_dhcp_writer_death_explicit_cli_recovers_group_twice(self):
        self._real_dhcp_writer_death_group_recovery(cli=True)

    def test_real_setup_then_dhcp_writer_death_cli_restores_link_and_missing_a(self):
        self._real_dhcp_writer_death_group_recovery(cli=True, shapes=True)

    def test_real_setup_then_dhcp_writer_death_cli_holds_malformed_group_binding(self):
        self._real_dhcp_writer_death_group_recovery(cli=True, fault="malformed-binding")

    def test_real_setup_dhcp_restore_death_before_directory_sync_cold_reenters(self):
        self._real_dhcp_writer_death_group_recovery(cli=True, recovery_death=True)

    def test_real_setup_then_dhcp_writer_death_group_recovery_holds_foreign_evidence(self):
        for fault in ("receipt-owner", "second-transaction", "restore-stage", "parent-b-visible"):
            with self.subTest(fault=fault):
                self._real_dhcp_writer_death_group_recovery(fault=fault)

    def _real_dhcp_writer_death_group_recovery(self, *, fault=None, cli=False, shapes=False,
                                            recovery_death=False):
        fixture = RELEASE_FIXTURE.ReleaseTransactionTests()
        fixture.setUp()
        try:
            state, _context = fixture._prepare_persisted_a_runtime_switch()
            root = fixture.root.resolve()
            p, q, common = state[:3]
            journal, descriptor, _backend = state[-3:]
            shutil.copytree(ROOT / "tools", root / "tools")
            for name in ("11-load.py", "01-a-setup.py"):
                shutil.copy2(ROOT / "DAY0-Prepare" / name, root / "DAY0-Prepare" / name)
            RELEASE_FIXTURE.LOAD.release_deployment_lock(descriptor)
            fixture._cleanups = [item for item in fixture._cleanups
                                 if item[0] != RELEASE_FIXTURE.LOAD.release_deployment_lock]
            operation = r'''
import sys, os, importlib.util
from pathlib import Path
from types import SimpleNamespace
root, q, common, test_root = map(Path, sys.argv[1:5])
sys.path.insert(0, str(test_root))
from test_cases.module_loader import load_script
from test_cases.test_load_release_transaction import assert_private_dhcp_transaction_retention
load = load_script('dhcp_group_cold_load', root / 'DAY0-Prepare/11-load.py')
load.active_managed_services = lambda *a, **kw: ()
backend = SimpleNamespace(name='supervisor')
fd = load.acquire_deployment_lock(exclusive=True)
guard = dict(deployment_lock_descriptor=fd, services_quiesced=True, runtime_backend=backend)
try:
    if sys.argv[5] == 'writer':
        setup = load_script('dhcp_group_real_setup', root / 'DAY0-Prepare/01-a-setup.py')
        load.service_runtime_backend = lambda: backend
        setup._SETUP_SWITCH_WAL_CONTEXT = (vars(load), fd)
        setup._register_setup_switch_link_before_write(str(common), str(q / '01-global.yaml'))
        common.unlink()
        common.symlink_to(os.path.relpath(q / '01-global.yaml', common.parent))
        (root / 'ztp/ztp-bootstrap_oob.sh').write_bytes(b'independent partial B runtime\n')
        shapes = sys.argv[6] == 'shapes'
        assert_private_dhcp_transaction_retention(load, root, q,
            'bound-setup-kill-last' if shapes or sys.argv[7] == 'recovery-death' else 'bound-setup-kill-first',
            setup_guard=guard, a_shapes=shapes)
        raise AssertionError('writer did not die at real first B install')
    names = ('dhcpd.conf', 'dhcpd_eth.hosts', 'dhcpd_ib.hosts', 'dhcpd_nvl.hosts')
    runtime = root / 'private-etc-dhcp'
    load.dhcp_file_mappings = lambda *a: {q / 'independent-dhcp-B' / n: runtime / n for n in names}
    if sys.argv[5] == 'kill-recovery':
        original_sync = load._fsync_parent_directory
        def kill_after_a_replace_before_sync(path):
            if Path(path) == runtime:
                assert (runtime / 'dhcpd.conf').read_bytes() == b'independent-A:dhcpd.conf\n'
                for name in names[1:]:
                    assert (runtime / name).read_bytes() == (q / 'independent-dhcp-B' / name).read_bytes()
                os.kill(os.getpid(), 9)
            return original_sync(path)
        load._fsync_parent_directory = kill_after_a_replace_before_sync
    if sys.argv[5] == 'cli':
        load.release_deployment_lock(fd)
        fd = None
        load.runtime_os = lambda: 'Linux'
        load.service_runtime_backend = lambda: backend
        def forbidden(*a, **kw):
            raise AssertionError('cold recovery invoked generation or service action')
        for name in ('initialize_from_template', 'activate_project', 'render_ztp_runtime',
                     'start_services', 'quiesce_services', 'validate_inputs', 'prepare_pubkeys'):
            setattr(load, name, forbidden)
        result = load.main([str(q), '--host-role', 'management-server', '--recover-setup-switch'])
        assert result == 0, result
    else:
        result = load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
            restore_runtime_from_wal=True, recover_dhcp_group=True, requested_new_project=q, **guard)
        assert result == 'restored-p-a-dhcp-retained', result
    assert (root / 'ztp/ztp-bootstrap_oob.sh').read_bytes() == b'A oob bootstrap\n'
    for n in names:
        if sys.argv[6] == 'shapes' and n == 'dhcpd_eth.hosts':
            assert os.readlink(runtime / n) == '../independent-original-A-source'
        elif sys.argv[6] == 'shapes' and n == 'dhcpd_nvl.hosts':
            assert not os.path.lexists(runtime / n)
        else:
            assert (runtime / n).read_bytes() == ('independent-A:' + n + '\n').encode()
    print('R22_BOUND_DHCP_GROUP_RESTORED_STOPPED_A')
finally:
    if fd is not None:
        load.release_deployment_lock(fd)
'''
            command = [sys.executable, "-B", "-c", operation, str(root), str(q), str(common), str(ROOT)]
            shape_arg = "shapes" if shapes else "files"
            death_arg = "recovery-death" if recovery_death else "normal"
            writer = subprocess.run(command + ["writer", shape_arg, death_arg], capture_output=True, text=True, timeout=30)
            self.assertEqual(-9, writer.returncode, writer.stdout + writer.stderr)
            before = journal.read_bytes()
            self.assertNotEqual(state[9][common], os.readlink(common))
            self.assertTrue((root / "private-etc-dhcp/dhcpd.conf").read_bytes().startswith(b"# independent-B"))
            runtime = root / "private-etc-dhcp"
            transaction, = runtime.glob(".load-dhcp-transaction-*")
            if fault == "receipt-owner":
                receipt = transaction / "a-preimage.json"
                original = receipt.read_bytes()
                replacement = transaction / "independent-receipt"
                replacement.write_bytes(original)
                replacement.chmod(0o600)
                os.replace(replacement, receipt)
            elif fault == "second-transaction":
                (runtime / ".load-dhcp-transaction-independent").mkdir(mode=0o700)
            elif fault == "restore-stage":
                (runtime / ".dhcpd.conf.rollback-independent").write_bytes(b"independent stage\n")
            elif fault == "parent-b-visible":
                candidate = root / "DAY0-Prepare" / q.name / "99-output-ztp/current-release.json"
                candidate.write_bytes(b"independent committed B\n")
            elif fault == "malformed-binding":
                envelope = json.loads(journal.read_bytes())
                del envelope["body"]["dhcp_transaction"]["body"]["runtime"]
                envelope["sha256"] = hashlib.sha256(json.dumps(
                    envelope["body"], sort_keys=True, separators=(",", ":"), ensure_ascii=False,
                ).encode()).hexdigest()
                journal.write_text(json.dumps(envelope), encoding="utf-8")
                before = journal.read_bytes()
            if fault is not None:
                p_before = {path: os.readlink(path) for path in state[9]}
                runtime_before = {path: path.read_bytes() for path in runtime.iterdir() if path.is_file()}
                bootstrap_before = (root / "ztp/ztp-bootstrap_oob.sh").read_bytes()
                recovery = subprocess.run(command + ["cli" if cli else "recover", shape_arg, death_arg], capture_output=True, text=True, timeout=30)
                self.assertNotEqual(0, recovery.returncode, recovery.stdout + recovery.stderr)
                self.assertIn("HOLD", recovery.stderr)
                self.assertEqual(p_before, {path: os.readlink(path) for path in state[9]})
                self.assertEqual(runtime_before, {path: path.read_bytes() for path in runtime.iterdir() if path.is_file()})
                self.assertEqual(bootstrap_before, (root / "ztp/ztp-bootstrap_oob.sh").read_bytes())
                self.assertEqual(before, journal.read_bytes())
                self.assertTrue((transaction / "a-preimage.json").is_file())
                self.assertFalse((root / "service-started").exists())
                return
            if recovery_death:
                killed = subprocess.run(command + ["kill-recovery", shape_arg, death_arg],
                                        capture_output=True, text=True, timeout=30)
                self.assertEqual(-9, killed.returncode, killed.stdout + killed.stderr)
                self.assertEqual(state[9], {path: os.readlink(path) for path in state[9]})
                self.assertEqual(before, journal.read_bytes())
                self.assertEqual(b"independent-A:dhcpd.conf\n", (runtime / "dhcpd.conf").read_bytes())
                for name in ("dhcpd_eth.hosts", "dhcpd_ib.hosts", "dhcpd_nvl.hosts"):
                    self.assertEqual((q / "independent-dhcp-B" / name).read_bytes(), (runtime / name).read_bytes())
                self.assertFalse((root / "service-started").exists())
            for _ in range(2):
                recovery = subprocess.run(command + ["cli" if cli else "recover", shape_arg, death_arg], capture_output=True, text=True, timeout=30)
                self.assertEqual(0, recovery.returncode, recovery.stdout + recovery.stderr)
                self.assertIn("R22_BOUND_DHCP_GROUP_RESTORED_STOPPED_A", recovery.stdout)
                self.assertEqual(state[9], {path: os.readlink(path) for path in state[9]})
                self.assertEqual(before, journal.read_bytes())
            self.assertFalse((root / "service-started").exists())
        finally:
            fixture.doCleanups()
            fixture.tearDown()

    def test_real_setup_then_dhcp_a_preimage_sync_failure_holds_original_a(self):
        self._real_setup_uncertain_parent_dhcp_evidence(
            reenter=False, outcomes=("preimage-file-sync",),
        )

    def _real_setup_uncertain_parent_dhcp_evidence(
        self, *, reenter, cold=False, inspect_preimage=False, inspect_parent_rebind=False,
        outcomes=("open", "fsync"), bind_setup=False,
    ):
        for boundary in outcomes:
            with self.subTest(boundary=boundary), tempfile.TemporaryDirectory() as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, q, common, _p_only, _q_only, manifest, before, _hook, _barrier, _home = fixture
                for relative in ("tools/lldp-analyze-tool/99-output-p2p",
                                 "tools/lldp-analyze-tool/99-output-monitor",
                                 "tools/ibdiagnet-analyze-tool/99-output-p2p"):
                    path = root / relative
                    if path.is_dir() and not path.is_symlink():
                        shutil.rmtree(path)
                first = self._run_private_load_setup_switch(
                    fixture, prepare_journal=False, auto_prepare_journal=True,
                    first_link=root / "never-created-link",
                )
                self.assertEqual(25, first.returncode, first.stdout + first.stderr)
                q_manifest, q_link = manifest.read_bytes(), os.readlink(common)
                self.assertNotEqual(before["manifest"], q_manifest)
                journal = root / "ztp/.setup-switch-journal.json"
                wal = journal.read_bytes()
                api = load_script("r22_private_dhcp_retention_" + boundary,
                                  root / "DAY0-Prepare/11-load.py")
                descriptor = api.acquire_deployment_lock(exclusive=True)
                try:
                    guard = (dict(deployment_lock_descriptor=descriptor, services_quiesced=True,
                                  runtime_backend=SimpleNamespace(name="supervisor")) if bind_setup else None)
                    with mock.patch.object(api, "active_managed_services", return_value=()):
                        RELEASE_FIXTURE.assert_private_dhcp_transaction_retention(
                            api, root, q, boundary, reenter=reenter, setup_guard=guard,
                        )
                finally:
                    api.release_deployment_lock(descriptor)
                if cold:
                    operation = """
import sys
from pathlib import Path
root, q, test_root = map(Path, sys.argv[1:4])
sys.path.insert(0, str(test_root))
from test_cases.module_loader import load_script
from test_cases.test_load_release_transaction import assert_private_dhcp_cold_pending_hold
load = load_script('r22_cold_pending_dhcp_load', root / 'DAY0-Prepare/11-load.py')
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert_private_dhcp_cold_pending_hold(load, root, q, inspect_preimage=(sys.argv[4] == 'inspect'),
                                         inspect_parent_rebind=(sys.argv[4] == 'parent-rebind'))
    print('R22_COLD_PENDING_DHCP_HELD')
finally:
    load.release_deployment_lock(descriptor)
"""
                    second = subprocess.run(
                        [sys.executable, "-B", "-c", operation, str(root), str(q), str(ROOT),
                         'parent-rebind' if inspect_parent_rebind else ('inspect' if inspect_preimage else 'hold-only')],
                        cwd=root, capture_output=True, text=True, check=False, timeout=30,
                    )
                    self.assertEqual(0, second.returncode, second.stdout + second.stderr)
                    self.assertIn("R22_COLD_PENDING_DHCP_HELD", second.stdout)
                self.assertEqual(q_manifest, manifest.read_bytes())
                self.assertEqual(q_link, os.readlink(common))
                if bind_setup and boundary not in {"bound-setup-register-wal", "bound-setup-register-receipt"}:
                    self.assertNotEqual(wal, journal.read_bytes())
                    prior, current = json.loads(wal)["body"], json.loads(journal.read_bytes())["body"]
                    self.assertEqual(prior, {k: v for k, v in current.items() if k != "dhcp_transaction"})
                else:
                    self.assertEqual(wal, journal.read_bytes())
                self.assertFalse((root / "service-started").exists())

    def test_real_setup_parent_visible_sync_failure_stops_success_and_cold_p_rollback(self):
        """A visible B with unconfirmed directory durability is HOLD, not A rollback."""
        for boundary in ("open", "fsync"):
            with self.subTest(boundary=boundary), \
                    tempfile.TemporaryDirectory(prefix="r22-parent-visible-sync-") as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, q, common, _p_only, _q_only, manifest, before, _hook, _barrier, _home = fixture
                for relative in ("tools/lldp-analyze-tool/99-output-p2p",
                                 "tools/lldp-analyze-tool/99-output-monitor",
                                 "tools/ibdiagnet-analyze-tool/99-output-p2p"):
                    path = root / relative
                    if path.is_dir() and not path.is_symlink():
                        shutil.rmtree(path)
                first = self._run_private_load_setup_switch(
                    fixture, prepare_journal=False, auto_prepare_journal=True,
                    first_link=root / "never-created-link",
                )
                self.assertEqual(25, first.returncode, first.stdout + first.stderr)
                q_manifest, q_link = manifest.read_bytes(), os.readlink(common)
                self.assertNotEqual(before["manifest"], q_manifest)
                journal = root / "ztp/.setup-switch-journal.json"
                wal = journal.read_bytes()
                operation = """
import importlib.util, os, stat, sys
from pathlib import Path
root, q = map(Path, sys.argv[1:3])
boundary = sys.argv[3]
spec = importlib.util.spec_from_file_location('visible_parent_sync_writer', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
parent = q / '99-output-ztp/current-release.json'
parent.parent.mkdir(exist_ok=True)
temporary = parent.with_name('.independent-parent-B.json')
payload = b'{"release_id":"independent-Q-B"}\\n'
with temporary.open('wb') as stream:
    stream.write(payload)
    stream.flush()
    os.fsync(stream.fileno())
candidate = load.PreparedParentRelease(parent, temporary, 'independent-Q-B')
descriptor = load.acquire_deployment_lock(exclusive=True)
real_open, real_sync = os.open, os.fsync
injected = []
def fail_open(path, flags, *args, **kwargs):
    if boundary == 'open' and Path(path) == parent.parent:
        injected.append('open')
        raise OSError('independent directory open failure')
    return real_open(path, flags, *args, **kwargs)
def fail_sync(fd):
    metadata = os.fstat(fd)
    if boundary == 'fsync' and stat.S_ISDIR(metadata.st_mode) and metadata.st_ino == parent.parent.stat().st_ino:
        injected.append('fsync')
        raise OSError('independent directory sync failure')
    return real_sync(fd)
os.open, os.fsync = fail_open, fail_sync
try:
    try:
        load.commit_prepared_release(candidate)
        (root / 'service-started').write_bytes(b'UNEXPECTED SUCCESS CONTINUATION')
    except load.LoadError as exc:
        assert '持久化' in str(exc) and '未确认' in str(exc), str(exc)
    else:
        raise AssertionError('unconfirmed parent durability was reported as success')
    assert candidate.committed and parent.read_bytes() == payload
    assert injected == [boundary], injected
finally:
    os.open, os.fsync = real_open, real_sync
    load.release_deployment_lock(descriptor)
"""
                second = subprocess.run(
                    [sys.executable, "-B", "-c", operation, str(root), str(q), boundary],
                    cwd=root, capture_output=True, text=True, check=False, timeout=30,
                )
                self.assertEqual(0, second.returncode, second.stdout + second.stderr)
                self.assertFalse((root / "service-started").exists())
                recovery = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('visible_parent_sync_cold_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        assert 'Q parent' in str(exc), str(exc)
    else:
        raise AssertionError('visible B licensed P rollback after sync failure')
finally:
    load.release_deployment_lock(descriptor)
"""
                third = subprocess.run(
                    [sys.executable, "-B", "-c", recovery, str(root)],
                    cwd=root, capture_output=True, text=True, check=False, timeout=30,
                )
                self.assertEqual(0, third.returncode, third.stdout + third.stderr)
                self.assertEqual(q_manifest, manifest.read_bytes())
                self.assertEqual(q_link, os.readlink(common))
                self.assertEqual(wal, journal.read_bytes())
                self.assertEqual(b'{"release_id":"independent-Q-B"}\n',
                                 (q / "99-output-ztp/current-release.json").read_bytes())
                self.assertFalse((root / "service-started").exists())

    def test_real_setup_then_parent_commit_cannot_cold_restore_p(self):
        """Actual Q parent replace must revoke the old P rollback authority."""
        with tempfile.TemporaryDirectory(prefix="r22-q-parent-commit-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, before, _hook, _barrier, _home = fixture
            for relative in ("tools/lldp-analyze-tool/99-output-p2p",
                             "tools/lldp-analyze-tool/99-output-monitor",
                             "tools/ibdiagnet-analyze-tool/99-output-p2p"):
                leaf = root / relative
                if leaf.is_dir() and not leaf.is_symlink():
                    shutil.rmtree(leaf)
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=False, auto_prepare_journal=True,
                first_link=root / "never-created-link",
            )
            self.assertEqual(25, first.returncode, first.stdout + first.stderr)
            committed_setup_manifest = manifest.read_bytes()
            self.assertNotEqual(before["manifest"], committed_setup_manifest)
            mixed = os.readlink(common)
            evidence = (root / "ztp/.setup-switch-journal.json").read_bytes()
            operation = """
import importlib.util
import os
import sys
from pathlib import Path
from types import SimpleNamespace
root, q = map(Path, sys.argv[1:])
spec = importlib.util.spec_from_file_location('fresh_parent_commit_guard', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
parent = q / '99-output-ztp/current-release.json'
parent.parent.mkdir(exist_ok=True)
temporary = parent.with_name('.candidate-parent.json')
with temporary.open('wb') as stream:
    stream.write(b'{"release_id":"Q-committed-B"}\\n')
    stream.flush()
    os.fsync(stream.fileno())
candidate = load.PreparedParentRelease(parent, temporary, 'Q-committed-B')
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    load.commit_prepared_release(candidate)
    assert candidate.committed
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as exc:
        assert 'Q parent' in str(exc), str(exc)
    else:
        raise AssertionError('committed Q parent incorrectly licensed P restoration')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", operation, str(root), str(q)],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(mixed, os.readlink(common))
            self.assertEqual(committed_setup_manifest, manifest.read_bytes())
            self.assertEqual(evidence, (root / "ztp/.setup-switch-journal.json").read_bytes())
            self.assertEqual(b'{"release_id":"Q-committed-B"}\n',
                             (q / "99-output-ztp/current-release.json").read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_q_local_latest_preimage_survives_temp_and_pointer_death(self):
        """Q's old latest is not a P manifest link or an absent Q-only name."""
        for boundary in ("temporary", "pointer", "q-manifest-committed"):
            with self.subTest(boundary=boundary), \
                    tempfile.TemporaryDirectory(prefix="r22-q-local-latest-") as name:
                fixture = self._private_setup_switch_fixture(Path(name))
                root, _p, q, common, p_only, q_only, manifest, before, _hook, _barrier, _home = fixture
                for relative in ("tools/lldp-analyze-tool/99-output-p2p",
                                 "tools/lldp-analyze-tool/99-output-monitor",
                                 "tools/ibdiagnet-analyze-tool/99-output-p2p"):
                    leaf = root / relative
                    if leaf.is_dir() and not leaf.is_symlink():
                        shutil.rmtree(leaf)
                output = q / "99-output-eth"
                output.mkdir()
                old, new = (output / value for value in (
                    "20261005_110000_combine", "20261005_120000_combine",
                ))
                for release in (old, new):
                    release.mkdir()
                    (release / ".published-complete").write_bytes(b"complete\n")
                    (release / "leaf.yaml").write_bytes(release.name.encode() + b"\n")
                latest = output / "latest"
                latest.symlink_to(old.name)
                nvos_output = q / "99-output-ib_nvl"
                nvos_output.mkdir()
                nvos_old, nvos_new = (nvos_output / value for value in (
                    "20261005_110000-ib", "20261005_120000-combine",
                ))
                for release in (nvos_old, nvos_new):
                    release.mkdir()
                    (release / ".published-complete").write_bytes(b"complete\n")
                    (release / "leaf.yaml").write_bytes(release.name.encode() + b"\n")
                nvos_latest = nvos_output / "latest"
                nvos_latest.symlink_to(nvos_old.name)
                first = self._run_private_load_setup_switch(
                    fixture, prepare_journal=False, auto_prepare_journal=True,
                    first_link=root / "never-created-link",
                    first_link_prefix=str(nvos_output / ".latest.tmp.") if boundary == "temporary" else None,
                    kill_after_replace=nvos_latest if boundary == "pointer" else root / "never-replaced",
                )
                self.assertEqual(25 if boundary == "q-manifest-committed" else 23,
                                 first.returncode, first.stdout + first.stderr)
                self.assertEqual(new.name, os.readlink(latest))
                self.assertEqual(nvos_old.name if boundary == "temporary" else nvos_new.name,
                                 os.readlink(nvos_latest))
                if boundary == "q-manifest-committed":
                    self.assertNotEqual(before["manifest"], manifest.read_bytes())
                    self.assertEqual("template/99-output/latest", os.readlink(root / "ztp/config/cumulus/latest_yaml"))
                journal = root / "ztp/.setup-switch-journal.json"
                body = json.loads(journal.read_bytes())["body"]
                self.assertEqual(old.name, body["q_local_links"][0]["before"])
                self.assertEqual(new.name, body["q_local_links"][0]["after"])
                self.assertEqual(2, len(body["q_local_links"]))
                recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_q_local_latest', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
        services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
    ) == 'restored-p'
finally:
    load.release_deployment_lock(descriptor)
"""
                second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                        cwd=root, capture_output=True, text=True, check=False)
                self.assertEqual(0, second.returncode, second.stdout + second.stderr)
                self.assertEqual(old.name, os.readlink(latest))
                self.assertEqual(nvos_old.name, os.readlink(nvos_latest))
                self.assertEqual(before["manifest"], manifest.read_bytes())
                self.assertEqual(before["common"], os.readlink(common))
                self.assertEqual(before["p_only"], os.readlink(p_only))
                self.assertFalse(os.path.lexists(q_only))
                self.assertFalse(list(output.glob(".latest.tmp.*")))
                self.assertFalse(list(nvos_output.glob(".latest.tmp.*")))
                for release in (old, new, nvos_old, nvos_new):
                    self.assertEqual(release.name.encode() + b"\n", (release / "leaf.yaml").read_bytes())
                self.assertFalse(journal.exists())
                self.assertFalse((root / "service-started").exists())

    def test_real_setup_infra_log_death_cold_recovery_preserves_replacement_receipt(self):
        """A later same-byte receipt inode must not authorize or be deleted by recovery."""
        with tempfile.TemporaryDirectory(prefix="r22-infra-receipt-owner-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, _p_only, _q_only, manifest, _before, _hook, barrier, _home = fixture
            source = root / "infra/logs"
            source.mkdir()
            payload = b"independent P-era incident evidence\n"
            (source / "audit.log").write_bytes(payload)
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link=source,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            self.assertTrue(source.is_symlink())
            receipt = root / "infra/.logs-migration.json"
            original = receipt.read_bytes()
            journal = root / "ztp/.setup-switch-journal.json"
            evidence, mixed, q_manifest = journal.read_bytes(), os.readlink(common), manifest.read_bytes()
            backup = source.resolve() / ".logs-migration-source"
            copied = source.resolve() / "audit.log"
            recovery = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('cold_infra_receipt_owner_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
receipt = root / 'infra/.logs-migration.json'
restore = load._setup_switch_restore_infra_log_tree
def replace_before_first_write(snapshot, **guard):
    original, before = receipt.read_bytes(), receipt.stat()
    receipt.rename(receipt.with_name('operator-retained-receipt.json'))
    receipt.write_bytes(original)
    receipt.chmod(0o600)
    assert receipt.stat().st_ino != before.st_ino
    restore(snapshot, **guard)
load._setup_switch_restore_infra_log_tree = replace_before_first_write
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as error:
        assert 'infra log receipt changed' in str(error), str(error)
        print('R22_REPLACEMENT_INFRA_RECEIPT_HELD')
    else:
        raise AssertionError('cold recovery consumed replacement receipt ownership')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertIn("R22_REPLACEMENT_INFRA_RECEIPT_HELD", second.stdout)
            self.assertEqual(original, receipt.read_bytes())
            self.assertEqual(original, receipt.with_name("operator-retained-receipt.json").read_bytes())
            self.assertEqual(evidence, journal.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            self.assertEqual(q_manifest, manifest.read_bytes())
            self.assertTrue(source.is_symlink())
            self.assertEqual(payload, (backup / "audit.log").read_bytes())
            self.assertEqual(payload, copied.read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_infra_log_death_cold_recovery_preserves_replaced_q_copy(self):
        """Only the receipt-bound original hardlink has Q cleanup authority."""
        with tempfile.TemporaryDirectory(prefix="r22-infra-copy-owner-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, _p_only, _q_only, manifest, _before, _hook, barrier, _home = fixture
            source = root / "infra/logs"
            source.mkdir()
            payload = b"independent original infra audit evidence\n"
            (source / "audit.log").write_bytes(payload)
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True, first_link=source)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            self.assertTrue(source.is_symlink())
            destination = source.resolve()
            copy, backup = destination / "audit.log", destination / ".logs-migration-source"
            receipt = root / "infra/.logs-migration.json"
            journal = root / "ztp/.setup-switch-journal.json"
            receipt_bytes, wal = receipt.read_bytes(), journal.read_bytes()
            mixed, manifest_bytes = os.readlink(common), manifest.read_bytes()
            recovery = """
import importlib.util, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('cold_infra_copy_owner_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
copy = (root / 'infra/logs').resolve() / 'audit.log'
restore = load._setup_switch_restore_infra_log_tree
def replace_before_first_write(snapshot, **guard):
    original, before = copy.read_bytes(), copy.stat()
    copy.rename(copy.parent.parent / 'operator-retained-audit.log')
    copy.write_bytes(original)
    assert copy.stat().st_ino != before.st_ino
    restore(snapshot, **guard)
load._setup_switch_restore_infra_log_tree = replace_before_first_write
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as error:
        assert 'infra log' in str(error) and 'changed' in str(error), str(error)
        print('R22_REPLACEMENT_INFRA_COPY_HELD')
    else:
        raise AssertionError('cold recovery deleted a replacement Q log owner')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertIn("R22_REPLACEMENT_INFRA_COPY_HELD", second.stdout)
            self.assertEqual(payload, copy.read_bytes())
            self.assertEqual(payload, (destination.parent / "operator-retained-audit.log").read_bytes())
            self.assertEqual(payload, (backup / "audit.log").read_bytes())
            self.assertEqual(receipt_bytes, receipt.read_bytes())
            self.assertEqual(wal, journal.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            self.assertEqual(manifest_bytes, manifest.read_bytes())
            self.assertTrue(source.is_symlink())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_infra_log_death_cold_recovery_does_not_restore_replaced_backup(self):
        """A backup replaced after the Q pointer unlink must stay outside P."""
        with tempfile.TemporaryDirectory(prefix="r22-infra-backup-owner-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, _p_only, _q_only, manifest, _before, _hook, barrier, _home = fixture
            source = root / "infra/logs"
            source.mkdir()
            payload = b"independent original P log directory\n"
            (source / "audit.log").write_bytes(payload)
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True, first_link=source)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            destination = source.resolve()
            backup = destination / ".logs-migration-source"
            retained = destination.parent / "operator-retained-original-infra"
            receipt, journal = root / "infra/.logs-migration.json", root / "ztp/.setup-switch-journal.json"
            receipt_bytes, wal = receipt.read_bytes(), journal.read_bytes()
            mixed, manifest_bytes = os.readlink(common), manifest.read_bytes()
            recovery = """
import importlib.util, os, sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('cold_infra_backup_owner_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
source = root / 'infra/logs'
destination = source.resolve()
backup = destination / '.logs-migration-source'
sync = load._fsync_parent_directory
mutated = False
def replace_after_unlink(parent):
    global mutated
    sync(parent)
    if parent == source.parent and not os.path.lexists(source) and not mutated:
        mutated = True
        backup.rename(destination.parent / 'operator-retained-original-infra')
        backup.mkdir()
        (backup / 'audit.log').write_bytes(b'foreign replacement backup evidence\\n')
load._fsync_parent_directory = replace_after_unlink
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(
            root / 'ztp/.setup-switch-journal.json', deployment_lock_descriptor=descriptor,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'),
        )
    except load.LoadError as error:
        assert 'infra log' in str(error) and 'changed' in str(error), str(error)
        assert mutated
        assert not os.path.lexists(source), 'replacement backup was moved into P'
        print('R22_REPLACEMENT_INFRA_BACKUP_HELD')
    else:
        raise AssertionError('cold recovery accepted a replacement backup')
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run([sys.executable, "-B", "-c", recovery, str(root)],
                                    cwd=root, capture_output=True, text=True, check=False, timeout=30)
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertIn("R22_REPLACEMENT_INFRA_BACKUP_HELD", second.stdout)
            self.assertFalse(os.path.lexists(source))
            self.assertEqual(b"foreign replacement backup evidence\n", (backup / "audit.log").read_bytes())
            self.assertEqual(payload, (retained / "audit.log").read_bytes())
            self.assertEqual(payload, (destination / "audit.log").read_bytes())
            self.assertEqual(receipt_bytes, receipt.read_bytes())
            self.assertEqual(wal, journal.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            self.assertEqual(manifest_bytes, manifest.read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_infra_log_cold_wal_drift_after_unlink_holds_remaining_writes(self):
        self._assert_real_infra_recovery_late_authority_holds("wal", "pointer-unlink")

    def test_real_setup_infra_log_cold_q_commit_after_rename_holds_remaining_writes(self):
        self._assert_real_infra_recovery_late_authority_holds("q", "directory-restore")

    def _assert_real_infra_recovery_late_authority_holds(self, kind, boundary):
        with tempfile.TemporaryDirectory(prefix="r22-infra-total-authority-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, _p_only, _q_only, manifest, _before, _hook, barrier, _home = fixture
            source = root / "infra/logs"
            source.mkdir()
            payload = b"independent P audit evidence for cold recovery\n"
            (source / "audit.log").write_bytes(payload)
            owner = source.stat().st_ino
            first = self._run_private_load_setup_switch(fixture, prepare_journal=True, first_link=source)
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file())
            self.assertTrue(source.is_symlink())
            journal, receipt = root / "ztp/.setup-switch-journal.json", root / "infra/.logs-migration.json"
            row = json.loads(journal.read_bytes())["body"]["infra_logs"][0]
            backup, destination = Path(row["backup"]), Path(row["destination"])
            wal, incident, p_manifest, mixed = journal.read_bytes(), receipt.read_bytes(), manifest.read_bytes(), os.readlink(common)
            copy_owner = (destination / "audit.log").stat().st_ino
            recovery = r'''
import importlib.util, os, sys
from pathlib import Path
from types import SimpleNamespace
root, kind, boundary = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
spec = importlib.util.spec_from_file_location('cold_infra_total_authority', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
journal, source = root / 'ztp/.setup-switch-journal.json', root / 'infra/logs'
import json
body = json.loads(journal.read_bytes())['body']
destination = Path(body['infra_logs'][0]['destination'])
sync, injected = load._fsync_parent_directory, []
def change_after_owned_step(parent):
    result = sync(parent)
    reached = (not os.path.lexists(source) if boundary == 'pointer-unlink'
               else source.is_dir() and not source.is_symlink())
    if parent == source.parent and reached and not injected:
        if kind == 'wal':
            original, before = journal.read_bytes(), journal.stat().st_ino
            journal.rename(journal.with_name('held-original-infra-wal'))
            journal.write_bytes(original)
            journal.chmod(0o600)
            assert before != journal.stat().st_ino
        else:
            marker = Path(body['q_parent_preimage']['path'])
            assert not os.path.lexists(marker)
            marker.parent.mkdir(parents=True, exist_ok=True)
            marker.write_bytes(b'independent Q commit marker during cold recovery\n')
        injected.append(True)
    return result
load._fsync_parent_directory = change_after_owned_step
fd = load.acquire_deployment_lock(exclusive=True)
try:
    try:
        load.recover_setup_switch_journal(journal, deployment_lock_descriptor=fd,
            services_quiesced=True, runtime_backend=SimpleNamespace(name='supervisor'))
    except load.LoadError as exc:
        expected = 'journal changed' if kind == 'wal' else 'Q parent changed'
        assert injected == [True] and expected in str(exc), str(exc)
        assert (destination / 'audit.log').is_file(), 'drift licensed Q copy deletion'
        assert (root / 'infra/.logs-migration.json').is_file(), 'drift licensed incident retirement'
        if boundary == 'pointer-unlink':
            assert not os.path.lexists(source), 'drift licensed backup restoration'
        print('R22_INFRA_TOTAL_AUTHORITY_HELD_BETWEEN_WRITES')
        raise SystemExit(7)
    else:
        raise AssertionError('cold recovery consumed an unknown authority')
finally:
    load.release_deployment_lock(fd)
'''
            result = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root), kind, boundary],
                cwd=root, capture_output=True, text=True, check=False, timeout=30,
            )
            self.assertEqual(7, result.returncode, result.stdout + result.stderr)
            self.assertIn("R22_INFRA_TOTAL_AUTHORITY_HELD_BETWEEN_WRITES", result.stdout)
            self.assertEqual(wal, journal.read_bytes())
            self.assertEqual(incident, receipt.read_bytes())
            self.assertEqual(p_manifest, manifest.read_bytes())
            self.assertEqual(mixed, os.readlink(common))
            self.assertEqual(copy_owner, (destination / "audit.log").stat().st_ino)
            self.assertEqual(payload, (destination / "audit.log").read_bytes())
            if boundary == "pointer-unlink":
                self.assertFalse(os.path.lexists(source))
                self.assertEqual(owner, backup.stat().st_ino)
                self.assertEqual(wal, journal.with_name("held-original-infra-wal").read_bytes())
            else:
                self.assertEqual(owner, source.stat().st_ino)
                self.assertFalse(source.is_symlink())
                self.assertFalse(os.path.lexists(backup))
                self.assertEqual(b"independent Q commit marker during cold recovery\n",
                                 (q / "99-output-ztp/current-release.json").read_bytes())
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_infra_log_pointer_death_cold_restores_real_directory(self):
        """LOAD→SETUP death after real log migration restores exact old bytes."""
        with tempfile.TemporaryDirectory(prefix="r22-real-infra-restore-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            logs = root / "infra/logs"
            logs.mkdir(parents=True)
            payload = b"P-era audit bytes\n"
            (logs / "audit.log").write_bytes(payload)
            original = logs.stat()
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link=logs,
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertTrue(barrier.is_file(), first.stdout + first.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertTrue(logs.is_symlink())
            destination = q / "99-output-infra"
            self.assertEqual(payload, (destination / "audit.log").read_bytes())
            journal = root / "ztp/.setup-switch-journal.json"
            rows = json.loads(journal.read_bytes())["body"]["infra_logs"]
            self.assertEqual(1, len(rows))
            backup = Path(rows[0]["backup"])
            self.assertEqual((original.st_dev, original.st_ino),
                             (backup.stat().st_dev, backup.stat().st_ino))
            receipt = root / "infra/.logs-migration.json"
            self.assertTrue(receipt.is_file())

            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location(
    'fresh_infra_switch_load', root / 'DAY0-Prepare/11-load.py',
)
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(logs.is_symlink())
            self.assertEqual((original.st_dev, original.st_ino),
                             (logs.stat().st_dev, logs.stat().st_ino))
            self.assertEqual(payload, (logs / "audit.log").read_bytes())
            self.assertFalse(os.path.lexists(destination / "audit.log"))
            self.assertFalse(os.path.lexists(backup))
            self.assertFalse(os.path.lexists(receipt))
            self.assertFalse(os.path.lexists(journal))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_infra_log_partial_copy_death_cold_removes_owned_hardlink(self):
        """Receipt plus an interrupted first Q hardlink can restore exact P."""
        with tempfile.TemporaryDirectory(prefix="r22-real-infra-copy-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            logs = root / "infra/logs"
            logs.mkdir(parents=True)
            payload = b"P-era audit bytes\n"
            (logs / "audit.log").write_bytes(payload)
            original = logs.stat()
            destination = q / "99-output-infra"
            first = self._run_private_load_setup_switch(
                fixture, prepare_journal=True,
                kill_after_hardlink=destination / "audit.log",
            )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertEqual(b"infra logs first hardlink written\n", barrier.read_bytes())
            self.assertFalse(logs.is_symlink())
            self.assertEqual((original.st_dev, original.st_ino),
                             (logs.stat().st_dev, logs.stat().st_ino))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            receipt = root / "infra/.logs-migration.json"
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertTrue(receipt.is_file())
            self.assertTrue(journal.is_file())
            self.assertEqual(payload, (destination / "audit.log").read_bytes())

            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location(
    'fresh_infra_copy_load', root / 'DAY0-Prepare/11-load.py',
)
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertEqual((original.st_dev, original.st_ino),
                             (logs.stat().st_dev, logs.stat().st_ino))
            self.assertEqual(payload, (logs / "audit.log").read_bytes())
            self.assertFalse(os.path.lexists(destination / "audit.log"))
            self.assertFalse(os.path.lexists(receipt))
            self.assertFalse(os.path.lexists(journal))
            self.assertFalse((root / "service-started").exists())

    def test_real_setup_wal_refuses_unjournaled_p_only_cleanup_before_q_link(self):
        """A P-owned legacy link absent from the WAL cannot be deleted silently."""
        with tempfile.TemporaryDirectory(prefix="r22-real-p-only-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            legacy = root / "ztp/config/publickey/legacy.pub"
            legacy.symlink_to(os.path.relpath(p / "operator.pub", legacy.parent))
            legacy_raw = os.readlink(legacy)
            result = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertNotEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertEqual(legacy_raw, os.readlink(legacy))
            self.assertFalse(os.path.lexists(q_only))
            self.assertTrue((root / "ztp/.setup-switch-journal.json").is_file())

    def test_real_setup_wal_refuses_unlisted_legacy_p2p_before_any_shared_write(self):
        """A legacy P2P link outside P's manifest cannot be silently removed."""
        with tempfile.TemporaryDirectory(prefix="r22-real-p2p-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            legacy = root / "ztp/config/cumulus/template/P2P/legacy.xlsx"
            legacy.parent.mkdir(parents=True, exist_ok=True)
            legacy.symlink_to(os.path.relpath(p / "p2p.xlsx", legacy.parent))
            original = os.readlink(legacy)
            result = self._run_private_load_setup_switch(fixture, prepare_journal=True)
            self.assertNotEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertIn("WAL 未覆盖旧链接清理", result.stdout + result.stderr)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertEqual(original, os.readlink(legacy))
            self.assertFalse(os.path.lexists(q_only))
            self.assertTrue((root / "ztp/.setup-switch-journal.json").is_file())

    def test_real_setup_wal_retains_legacy_p2p_inserted_after_preflight(self):
        """A new same-UID legacy link at cleanup time must survive the switch."""
        with tempfile.TemporaryDirectory(prefix="r22-real-p2p-race-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, _q, common, p_only, _q_only, manifest, before, hook, _barrier, _home = fixture
            legacy = root / "ztp/config/cumulus/template/P2P/raced.xlsx"
            legacy.parent.mkdir(parents=True, exist_ok=True)
            raw_target = os.path.relpath(p / "p2p.xlsx", legacy.parent)
            hook_file = hook / "sitecustomize.py"
            hook_file.write_text(hook_file.read_text(encoding="utf-8") + """
import glob
import inspect
original_glob = glob.glob
def insert_at_legacy_cleanup(pattern, *args, **kwargs):
    legacy = Path(os.environ['HTTP_TEST_RACE_LEGACY_LINK'])
    if (pattern == str(legacy.parent / '*.xlsx')
            and any(frame.function == '_remove_legacy_p2p_links'
                    for frame in inspect.stack())
            and not os.path.lexists(legacy)):
        os.symlink(os.environ['HTTP_TEST_RACE_LEGACY_TARGET'], legacy)
    return original_glob(pattern, *args, **kwargs)
glob.glob = insert_at_legacy_cleanup
""", encoding="utf-8")
            with mock.patch.dict(os.environ, {
                "HTTP_TEST_RACE_LEGACY_LINK": str(legacy),
                "HTTP_TEST_RACE_LEGACY_TARGET": raw_target,
            }):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True,
                    first_link=root / "never-created-link",
                )
            self.assertNotEqual(0, result.returncode, result.stdout + result.stderr)
            self.assertIn("旧 P2P 链接缺少 WAL", result.stdout + result.stderr)
            self.assertEqual(raw_target, os.readlink(legacy))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))

    def test_real_setup_wal_retains_generic_cleanup_link_inserted_after_preflight(self):
        """The actual LOAD→SETUP cleanup cannot delete an unjournaled late P link."""
        with tempfile.TemporaryDirectory(prefix="r22-real-cleanup-race-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, _q, common, p_only, q_only, manifest, before, hook, barrier, _home = fixture
            late = root / "ztp/config/publickey/late.pub"
            raw_target = os.path.relpath(p / "operator.pub", late.parent)
            hook_file = hook / "sitecustomize.py"
            hook_file.write_text(hook_file.read_text(encoding="utf-8") + """
import glob
import inspect
original_glob_for_cleanup = glob.glob
def insert_at_generic_cleanup(pattern, *args, **kwargs):
    late = Path(os.environ['HTTP_TEST_RACE_GENERIC_LINK'])
    if (pattern == str(late.parent / '*.pub')
            and any(frame.function == '_unsetup_previous' for frame in inspect.stack())
            and not os.path.lexists(late)):
        os.symlink(os.environ['HTTP_TEST_RACE_GENERIC_TARGET'], late)
    return original_glob_for_cleanup(pattern, *args, **kwargs)
glob.glob = insert_at_generic_cleanup
""", encoding="utf-8")
            with mock.patch.dict(os.environ, {
                "HTTP_TEST_RACE_GENERIC_LINK": str(late),
                "HTTP_TEST_RACE_GENERIC_TARGET": raw_target,
            }):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True,
                    first_link=root / "never-created-link",
                )
            self.assertNotEqual(0, result.returncode, result.stdout + result.stderr)
            self.assertIn("WAL", result.stdout + result.stderr)
            self.assertEqual(raw_target, os.readlink(late))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(barrier.exists())

    def test_real_setup_sample_unlink_is_durable_before_next_link_and_cold_recovers_p(self):
        """SETUP must fsync the known P unlink before publishing another sample name."""
        with tempfile.TemporaryDirectory(prefix="r22-real-sample-unlink-fsync-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, q_only, manifest, before, hook, barrier, _home = fixture
            sample = root / f"ztp/optimize/{q.name}-sample"
            sample.mkdir()
            stale = sample / "old-managed-link"
            raw = os.path.relpath(p / "01-global.yaml", sample)
            stale.symlink_to(raw)
            manifest.write_bytes(before["manifest"] + (str(stale) + "\n").encode())
            before_manifest = manifest.read_bytes()
            hook_file = hook / "sitecustomize.py"
            hook_file.write_text(hook_file.read_text(encoding="utf-8") + """
sample_unlink_pending = False
original_sample_remove = os.remove
original_sample_unlink = os.unlink
original_sample_fsync = os.fsync
def record_sample_remove(path, *args, **kwargs):
    global sample_unlink_pending
    result = original_sample_remove(path, *args, **kwargs)
    if os.path.abspath(path) == os.environ['HTTP_TEST_SAMPLE_OLD_LINK']:
        sample_unlink_pending = True
    return result
def record_sample_unlink(path, *args, **kwargs):
    global sample_unlink_pending
    result = original_sample_unlink(path, *args, **kwargs)
    if os.path.abspath(path) == os.environ['HTTP_TEST_SAMPLE_OLD_LINK']:
        sample_unlink_pending = True
    return result
def kill_after_sample_parent_sync(fd):
    result = original_sample_fsync(fd)
    if sample_unlink_pending:
        held = os.fstat(fd)
        named = Path(os.environ['HTTP_TEST_SAMPLE_OLD_LINK']).parent.stat()
        if (held.st_dev, held.st_ino) == (named.st_dev, named.st_ino):
            Path(os.environ['HTTP_TEST_Q_LINK_BARRIER']).write_bytes(b'sample unlink parent fsynced\\n')
            os.kill(os.getpid(), signal.SIGKILL)
    return result
def refuse_next_sample_symlink(source, destination, *args, **kwargs):
    if sample_unlink_pending:
        raise ValueError('sample unlink not durable before next publication')
    return kill_after_real_q_link(source, destination, *args, **kwargs)
os.remove = record_sample_remove
os.unlink = record_sample_unlink
os.fsync = kill_after_sample_parent_sync
os.symlink = refuse_next_sample_symlink
""", encoding="utf-8")
            with mock.patch.dict(os.environ, {"HTTP_TEST_SAMPLE_OLD_LINK": str(stale)}):
                first = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True, first_link=root / "never-created-link",
                )
            self.assertEqual(23, first.returncode, first.stdout + first.stderr)
            self.assertEqual(b"sample unlink parent fsynced\n", barrier.read_bytes())
            self.assertFalse(os.path.lexists(stale))
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('sample_unlink_fresh_load', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *a, **k: ()
fd = load.acquire_deployment_lock(exclusive=True)
try:
    assert load.recover_setup_switch_journal(root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=fd, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor')) == 'restored-p'
finally:
    load.release_deployment_lock(fd)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)], cwd=root,
                capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(raw, os.readlink(stale))
            self.assertEqual(before_manifest, manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse((root / "ztp/.setup-switch-journal.json").exists())

    def test_real_setup_pubkey_cleanup_retains_late_unowned_link(self):
        """A real secondary writer must not delete a link inserted after generic cleanup."""
        with tempfile.TemporaryDirectory(prefix="r22-real-pub-cleanup-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, before, hook, _barrier, _home = fixture
            late = root / "ztp/config/publickey/late-foreign.pub"
            foreign = root / "operator-owned.pub"
            foreign.write_bytes(b"operator-owned sentinel\n")
            raw = os.path.relpath(foreign, late.parent)
            hook_file = hook / "sitecustomize.py"
            hook_file.write_text(hook_file.read_text(encoding="utf-8") + """
import glob
import inspect
original_glob_for_secondary_cleanup = glob.glob
def insert_before_public_key_cleanup(pattern, *args, **kwargs):
    late = Path(os.environ['HTTP_TEST_SECONDARY_LINK'])
    if (pattern == str(late.parent / '*.pub')
            and any(frame.function == '_process_pubkeys' for frame in inspect.stack())
            and not os.path.lexists(late)):
        original_symlink(os.environ['HTTP_TEST_SECONDARY_TARGET'], late)
    return original_glob_for_secondary_cleanup(pattern, *args, **kwargs)
glob.glob = insert_before_public_key_cleanup
""", encoding="utf-8")
            with mock.patch.dict(os.environ, {
                "HTTP_TEST_SECONDARY_LINK": str(late), "HTTP_TEST_SECONDARY_TARGET": raw,
            }):
                result = self._run_private_load_setup_switch(
                    fixture, prepare_journal=True, first_link=root / "never-created-link",
                )
            self.assertEqual(24, result.returncode, result.stdout + result.stderr)
            self.assertIn("WAL", result.stdout + result.stderr)
            self.assertEqual(raw, os.readlink(late))
            self.assertEqual(b"operator-owned sentinel\n", foreign.read_bytes())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertTrue((root / "ztp/.setup-switch-journal.json").is_file())

    def test_real_setup_wal_refuses_unlisted_fixed_q_repoint_before_write(self):
        """P's fixed link cannot be repointed when its WAL row has no Q value."""
        with tempfile.TemporaryDirectory(prefix="r22-real-fixed-q-hold-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, omit_common_from_wal=True,
            )
            self.assertNotEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertFalse(barrier.exists(), result.stdout + result.stderr)
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertTrue((root / "ztp/.setup-switch-journal.json").is_file())

    def test_real_setup_wal_registers_dynamic_q_key_before_child_write(self):
        """A late-discovered Q key must be durable before a killed setup publishes it."""
        with tempfile.TemporaryDirectory(prefix="r22-real-q-key-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, q, _common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            (q / "late.pub").write_bytes((q / "operator.pub").read_bytes())
            late = root / "ztp/config/publickey/late.pub"
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link=late,
            )
            self.assertEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertEqual(b"first Q managed link written\n", barrier.read_bytes())
            journal = root / "ztp/.setup-switch-journal.json"
            body = json.loads(journal.read_bytes())["body"]
            self.assertIn(str(late), {row["path"] for row in body["links"]})
            self.assertEqual(before["manifest"], manifest.read_bytes())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_dynamic_q_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(late))
            self.assertFalse(journal.exists())
            self.assertEqual(before["manifest"], manifest.read_bytes())

    def test_real_setup_wal_registers_missing_infra_logs_before_child_write(self):
        """The first infra/logs link needs its own durable Q-only WAL row."""
        with tempfile.TemporaryDirectory(prefix="r22-real-infra-link-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, _common, _p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            logs = root / "infra/logs"
            self.assertFalse(os.path.lexists(logs))
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link=logs,
            )
            self.assertEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertEqual(b"first Q managed link written\n", barrier.read_bytes())
            journal = root / "ztp/.setup-switch-journal.json"
            body = json.loads(journal.read_bytes())["body"]
            self.assertIn(str(logs), {row["path"] for row in body["links"]})
            self.assertEqual(before["manifest"], manifest.read_bytes())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_infra_link_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(logs))
            self.assertFalse(journal.exists())
            self.assertEqual(before["manifest"], manifest.read_bytes())

    def test_real_setup_exact_monitor_temp_sigkill_cold_recovers_p(self):
        """An atomic exact-link temp must be WAL-known before SIGKILL."""
        with tempfile.TemporaryDirectory(prefix="r22-real-exact-temp-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            prefix = str(root / "ethernet/monitor/.eth.csv.tmp.")
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, first_link_prefix=prefix,
            )
            self.assertEqual(23, result.returncode, result.stdout + result.stderr)
            temporary = Path(barrier.read_text().strip())
            self.assertTrue(str(temporary).startswith(prefix))
            self.assertTrue(temporary.is_symlink())
            journal = root / "ztp/.setup-switch-journal.json"
            body = json.loads(journal.read_bytes())["body"]
            final = root / "ethernet/monitor/eth.csv"
            self.assertIn(str(temporary), {row["path"] for row in body["links"]})
            self.assertIn(str(final), {row["path"] for row in body["links"]})
            self.assertEqual(before["manifest"], manifest.read_bytes())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_exact_temp_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(temporary))
            self.assertFalse(journal.exists())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))

    def test_real_setup_exact_monitor_replace_sigkill_cold_recovers_p(self):
        """A completed exact-link replace is still Q-in-flight before commit."""
        with tempfile.TemporaryDirectory(prefix="r22-real-exact-replace-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, _p, _q, common, p_only, _q_only, manifest, before, _hook, barrier, _home = fixture
            final = root / "ethernet/monitor/eth.csv"
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, kill_after_replace=final,
            )
            self.assertEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertEqual(b"infra logs pointer replaced\n", barrier.read_bytes())
            self.assertTrue(final.is_symlink())
            journal = root / "ztp/.setup-switch-journal.json"
            rows = json.loads(journal.read_bytes())["body"]["links"]
            self.assertIn(str(final), {row["path"] for row in rows})
            self.assertEqual(before["manifest"], manifest.read_bytes())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_exact_replace_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(final))
            self.assertFalse(journal.exists())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))

    def test_real_setup_wal_registers_infra_repoint_temp_before_child_write(self):
        """A killed infra/logs re-point leaves only journalled P and temp links."""
        with tempfile.TemporaryDirectory(prefix="r22-real-infra-repoint-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            p_logs = p / "99-output-infra"
            q_logs = q / "99-output-infra"
            p_logs.mkdir(exist_ok=True)
            q_logs.mkdir(exist_ok=True)
            logs = root / "infra/logs"
            logs.symlink_to(os.path.relpath(p_logs, logs.parent))
            old_logs = os.readlink(logs)
            manifest.write_bytes(before["manifest"] + (str(logs) + "\n").encode())
            before["manifest"] = manifest.read_bytes()
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True,
                first_link_prefix=str(root / "infra/.logs.setup."),
            )
            self.assertEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertTrue(barrier.read_text().startswith(
                str(root / "infra/.logs.setup.")), barrier.read_text())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(old_logs, os.readlink(logs))
            journal = root / "ztp/.setup-switch-journal.json"
            body = json.loads(journal.read_bytes())["body"]
            temp_rows = [row for row in body["links"]
                         if row["path"].startswith(str(root / "infra/.logs.setup."))]
            self.assertEqual(1, len(temp_rows), (body["links"], list((root / "infra").iterdir()), result.stdout, result.stderr))
            temporary = Path(temp_rows[0]["path"])
            self.assertTrue(temporary.is_symlink())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_infra_repoint_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertFalse(os.path.lexists(temporary))
            self.assertEqual(old_logs, os.readlink(logs))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(journal.exists())

    def test_real_setup_wal_recovers_after_infra_logs_pointer_replace(self):
        """A killed re-point after replace restores P using the durable WAL."""
        with tempfile.TemporaryDirectory(prefix="r22-real-infra-replaced-wal-") as name:
            fixture = self._private_setup_switch_fixture(Path(name))
            root, p, q, common, p_only, q_only, manifest, before, _hook, barrier, _home = fixture
            p_logs, q_logs = p / "99-output-infra", q / "99-output-infra"
            p_logs.mkdir(exist_ok=True)
            q_logs.mkdir(exist_ok=True)
            logs = root / "infra/logs"
            logs.symlink_to(os.path.relpath(p_logs, logs.parent))
            old_logs = os.readlink(logs)
            manifest.write_bytes(before["manifest"] + (str(logs) + "\n").encode())
            before["manifest"] = manifest.read_bytes()
            result = self._run_private_load_setup_switch(
                fixture, prepare_journal=True, kill_after_replace=logs,
            )
            self.assertEqual(23, result.returncode, result.stdout + result.stderr)
            self.assertEqual(b"infra logs pointer replaced\n", barrier.read_bytes())
            self.assertEqual(q_logs, logs.resolve())
            self.assertEqual(before["manifest"], manifest.read_bytes())
            journal = root / "ztp/.setup-switch-journal.json"
            self.assertTrue(journal.is_file())
            recovery = """
import importlib.util
import sys
from pathlib import Path
from types import SimpleNamespace
root = Path(sys.argv[1])
spec = importlib.util.spec_from_file_location('fresh_infra_replaced_recovery', root / 'DAY0-Prepare/11-load.py')
load = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = load
spec.loader.exec_module(load)
load.active_managed_services = lambda *args, **kwargs: ()
descriptor = load.acquire_deployment_lock(exclusive=True)
try:
    outcome = load.recover_setup_switch_journal(
        root / 'ztp/.setup-switch-journal.json',
        deployment_lock_descriptor=descriptor, services_quiesced=True,
        runtime_backend=SimpleNamespace(name='supervisor'),
    )
    if outcome != 'restored-p':
        raise SystemExit('unexpected outcome: ' + repr(outcome))
finally:
    load.release_deployment_lock(descriptor)
"""
            second = subprocess.run(
                [sys.executable, "-B", "-c", recovery, str(root)],
                cwd=root, capture_output=True, text=True, check=False,
            )
            self.assertEqual(0, second.returncode, second.stdout + second.stderr)
            self.assertEqual(old_logs, os.readlink(logs))
            self.assertEqual(before["manifest"], manifest.read_bytes())
            self.assertEqual(before["common"], os.readlink(common))
            self.assertEqual(before["p_only"], os.readlink(p_only))
            self.assertFalse(os.path.lexists(q_only))
            self.assertFalse(journal.exists())

    def test_c2_process_death_between_version_and_pointer_commits(self):
        """Fresh readers reject an orphan but accept only a committed pointer."""
        child = """
import os
from pathlib import Path
import signal
import sys
from unittest import mock
from test_cases.module_loader import load_script

root, base_name, crash_after = sys.argv[1], sys.argv[2], int(sys.argv[3])
c2 = load_script("r22_c2_sigkill_child", Path(root) / "ztp/config/cumulus/template/90-c2-generate_configs.py")
base = Path(base_name)
real_replace = os.replace
replaces = 0

def patch_candidate(_dot, staging, _display, _authority):
    (Path(staging) / "Border01.yaml").write_bytes(b"patched B\\n")

def die_after_replace(src, dst):
    global replaces
    result = real_replace(src, dst)
    replaces += 1
    if replaces == crash_after:
        os.kill(os.getpid(), signal.SIGKILL)
    return result

with mock.patch.object(c2, "_require_csv_input_authority"), mock.patch.object(
    c2, "_patch_descriptions_in_directory", side_effect=patch_candidate,
), mock.patch.object(c2.os, "replace", side_effect=die_after_replace):
    c2._run_patch_descriptions("private-topology.dot", str(base), expected_csv_authority=object())
"""
        for crash_after in (1, 2):
            with self.subTest(crash_after=crash_after), tempfile.TemporaryDirectory(
                prefix="r22-c2-sigkill-",
            ) as name:
                output = Path(name) / "ztp/config/cumulus/template/99-output"
                output.mkdir(parents=True)
                timestamp = "20261002_120000"
                base = output / timestamp
                base.mkdir()
                (base / "Border01.yaml").write_bytes(b"plain base\n")
                legacy = output / f"{timestamp}_with_desc"
                legacy.mkdir()
                (legacy / "Border01.yaml").write_bytes(b"accepted legacy A\n")
                before_legacy = _artifact_tree(legacy)
                result = subprocess.run(
                    [sys.executable, "-B", "-c", child, str(ROOT), str(base),
                     str(crash_after)],
                    cwd=ROOT, capture_output=True, text=True, check=False,
                )
                self.assertEqual(-signal.SIGKILL, result.returncode, result.stderr)
                self.assertEqual(before_legacy, _artifact_tree(legacy))
                pointer = output / f"{timestamp}_with_desc_current"
                context = PUBLISHER._cumulus_dir_context(str(base))
                self.assertIsNotNone(context)
                if crash_after == 1:
                    self.assertFalse(os.path.lexists(pointer))
                    self.assertEqual(
                        "version-residue-no-current",
                        C2._inspect_description_output_recovery(str(legacy))["state"],
                    )
                    with self.assertRaisesRegex(ValueError, "version without a current pointer"):
                        PUBLISHER._preferred_cumulus_contexts([context])
                else:
                    self.assertTrue(pointer.is_symlink())
                    self.assertEqual(pointer, LOAD._select_cumulus_description_publication(base))
                    selected = PUBLISHER._preferred_cumulus_contexts([context])
                    self.assertEqual(str(pointer), selected[0]["path"])
                    self.assertEqual(
                        b"patched B\n", (Path(selected[0]["read_path"]) / "Border01.yaml").read_bytes(),
                    )

    def test_c2_single_pointer_is_the_real_d_input_and_tamper_never_falls_back(self):
        """C2's immutable publish pointer is what the real d resolver consumes."""
        with tempfile.TemporaryDirectory(prefix="r22-description-pointer-") as name:
            output = Path(name) / "ztp/config/cumulus/template/99-output"
            output.mkdir(parents=True)
            timestamp = "20261002_120000"
            base = output / timestamp
            base.mkdir()
            (base / "Border01.yaml").write_bytes(b"plain base\n")
            legacy = output / f"{timestamp}_with_desc"
            legacy.mkdir()
            (legacy / "Border01.yaml").write_bytes(b"complete legacy A\n")
            pointer = output / f"{timestamp}_with_desc_current"

            def patch_staged(_dot, staging, _display, _authority):
                (Path(staging) / "Border01.yaml").write_bytes(b"patched B\n")

            with mock.patch.object(
                C2, "_require_csv_input_authority",
            ), mock.patch.object(
                C2, "_patch_descriptions_in_directory",
                side_effect=patch_staged,
            ):
                C2._run_patch_descriptions(
                    "private-topology.dot", str(base),
                    expected_csv_authority=object(),
                )
            self.assertTrue(pointer.is_symlink())
            self.assertEqual(pointer, LOAD._select_cumulus_description_publication(base))
            self.assertEqual(b"complete legacy A\n",
                             (legacy / "Border01.yaml").read_bytes())
            context = PUBLISHER._cumulus_dir_context(str(base))
            self.assertIsNotNone(context)
            selected = PUBLISHER._preferred_cumulus_contexts([context])
            self.assertEqual(str(pointer), selected[0]["path"])
            version = PUBLISHER._description_current_version(str(pointer), timestamp)
            self.assertEqual(version, selected[0]["read_path"])
            self.assertEqual(b"patched B\n",
                             (Path(version) / "Border01.yaml").read_bytes())
            (Path(version) / "Border01.yaml").write_bytes(b"tampered\n")
            with self.assertRaisesRegex(ValueError, "hash|manifest"):
                PUBLISHER._preferred_cumulus_contexts([context])
            self.assertEqual(b"complete legacy A\n",
                             (legacy / "Border01.yaml").read_bytes())

    def test_real_setup_manifest_switch_then_load_restores_previous_project(self):
        """A failed Q generation restores P links and its setup authority."""
        with tempfile.TemporaryDirectory(prefix="r22-project-rollback-") as name:
            root = Path(name)
            p = root / "DAY0-Prepare/p"
            q = root / "DAY0-Prepare/q"
            p.mkdir(parents=True)
            q.mkdir(parents=True)
            for project in (p, q):
                (project / "01-global.yaml").write_text(project.name + "\n")
            ztp = root / "ztp"
            common = ztp / "config/cumulus/template/01-global.yaml"
            p_only = ztp / "p-only"
            q_only = ztp / "q-only"
            for path in (common, p_only, q_only):
                path.parent.mkdir(parents=True, exist_ok=True)
            manifest = ztp / ".setup_manifest"
            lock_path = root / ".deployment.lock"
            with mock.patch.multiple(
                LOAD, HTTP_ROOT=root, ZTP_DIR=ztp, MANIFEST=manifest,
                IMAGE_DIR=root / "image", DEPLOYMENT_LOCK=lock_path,
            ), mock.patch.object(SETUP, "MANIFEST_FILE", str(manifest)):
                SETUP._SetupLinkTransaction._restore_link(
                    str(common), os.path.relpath(p / "01-global.yaml", common.parent),
                )
                SETUP._SetupLinkTransaction._restore_link(
                    str(p_only), os.path.relpath(p, p_only.parent),
                )
                with mock.patch.object(
                    SETUP, "_collect_expected_links",
                    return_value=[str(common), str(p_only)],
                ):
                    SETUP._write_manifest(str(p.resolve()))
                old = LOAD.snapshot_active_setup_state(expected_project=p)
                SETUP._SetupLinkTransaction._restore_link(
                    str(common), os.path.relpath(q / "01-global.yaml", common.parent),
                )
                SETUP._SetupLinkTransaction._restore_link(
                    str(q_only), os.path.relpath(q, q_only.parent),
                )
                fsynced_directories = []
                real_fsync = os.fsync

                def record_fsync(descriptor):
                    metadata = os.fstat(descriptor)
                    if stat.S_ISDIR(metadata.st_mode):
                        fsynced_directories.append(metadata.st_ino)
                    return real_fsync(descriptor)

                with mock.patch.object(
                    SETUP, "_collect_expected_links",
                    return_value=[str(common), str(q_only)],
                ), mock.patch.object(SETUP.os, "fsync", side_effect=record_fsync):
                    SETUP._write_manifest(str(q.resolve()))
                self.assertIn(
                    manifest.parent.stat().st_ino, fsynced_directories,
                    "Q manifest replacement must fsync its parent directory",
                )
                self.assertEqual(q.resolve(), LOAD.active_project())
                descriptor = LOAD.acquire_deployment_lock(exclusive=True)
                try:
                    with mock.patch.object(
                        LOAD, "active_managed_services", return_value=(),
                    ):
                        LOAD.restore_active_setup_state(
                            old, current_project=q,
                            deployment_lock_descriptor=descriptor,
                            services_quiesced=True,
                            runtime_backend=SimpleNamespace(name="supervisor"),
                        )
                finally:
                    LOAD.release_deployment_lock(descriptor)
                self.assertEqual(p.resolve(), LOAD.active_project())
                self.assertEqual(old.manifest, manifest.read_bytes())
                self.assertEqual(old.links[common], os.readlink(common))
                self.assertEqual(old.links[p_only], os.readlink(p_only))
                self.assertFalse(os.path.lexists(q_only))

    def test_c2_crash_residue_blocks_real_publisher_base_fallback(self):
        """C2's cold gap cannot make d publish the unpatched base tree."""
        with tempfile.TemporaryDirectory(prefix="r22-c2-publisher-gap-") as name:
            output = Path(name) / "ztp/config/cumulus/template/99-output"
            output.mkdir(parents=True)
            timestamp = "20261002_120000"
            base = output / timestamp
            base.mkdir()
            (base / "Border01.yaml").write_bytes(b"base without descriptions\n")
            final = output / f"{timestamp}_with_desc"
            old = output / f".{final.name}.tmp.crash.old"
            stage = output / f".{final.name}.tmp.crash"
            old.mkdir()
            stage.mkdir()
            (old / "Border01.yaml").write_bytes(b"complete A descriptions\n")
            (stage / "Border01.yaml").write_bytes(b"incomplete B descriptions\n")
            before = _artifact_tree(output)

            inventory = C2._inspect_description_output_recovery(str(final))
            self.assertEqual("gap-after-old-rename", inventory["state"])
            self.assertEqual(str(old), inventory["old_candidate"])
            context = PUBLISHER._cumulus_dir_context(str(base))
            self.assertIsNotNone(context)
            with self.assertRaisesRegex(ValueError, "unresolved description"):
                PUBLISHER._preferred_cumulus_contexts([context])
            self.assertEqual(before, _artifact_tree(output))

    def test_real_nvos_scoped_generation_parent_remains_activatable(self):
        """Real IB/NVL generation must retain accepted empty unselected hosts."""
        private_parent = os.environ.get("R22_TEST_PRIVATE_ROOT")
        for selected in ("ib", "nvl"):
            with self.subTest(selected=selected), tempfile.TemporaryDirectory(
                prefix=f"r22-nvos-{selected}-activation-", dir=private_parent,
            ) as name:
                private = Path(name)

                class ScopedNVOSFlow(PLATFORM_FLOW.PlatformReleaseFlowTests):
                    selected_scope = selected

                    @classmethod
                    def _run_dhcp_generator(cls):
                        document = yaml.safe_load(cls.global_file.read_text(encoding="utf-8"))
                        document["common"]["mgmt"]["dhcp-server"] = {
                            "status": "enabled", "package": "isc-dhcp-server",
                        }
                        cls.global_file.write_text(
                            yaml.safe_dump(document, sort_keys=False), encoding="utf-8",
                        )
                        for filename in ("dhcpd_eth.hosts", "dhcpd_ib.hosts",
                                         "dhcpd_nvl.hosts", "dhcpd.conf"):
                            (cls.dhcp / filename).write_bytes(b"")
                        executed = []

                        def real_child(command, **_kwargs):
                            script = command[1]
                            executed.append((script, list(command[2:])))
                            if (command[command.index("--switch") + 1]
                                    != cls.selected_scope):
                                raise AssertionError((cls.selected_scope, command))
                            if (command[command.index("--deployment-scope") + 1]
                                    != "prod"):
                                raise AssertionError(command)
                            if script == "c1-generate_dhcp.py":
                                bindings = {
                                    "HTTP_ROOT": str(cls.root),
                                    "SCRIPT_DIR": str(cls.dhcp),
                                    "OUTPUT_ETH": str(cls.dhcp / "dhcpd_eth.hosts"),
                                    "OUTPUT_IB": str(cls.dhcp / "dhcpd_ib.hosts"),
                                    "OUTPUT_NVL": str(cls.dhcp / "dhcpd_nvl.hosts"),
                                    "OUTPUT_CONF": str(cls.dhcp / "dhcpd.conf"),
                                    "OUTPUT_MANIFEST": str(cls.dhcp / "dhcp-release-manifest.json"),
                                    "SUBNET_CSV": str(cls.dhcp / "02-subnet_config.csv"),
                                    "GLOBAL_YAML": str(cls.dhcp / "01-global.yaml"),
                                    "P2P_AIR_JSON": str(cls.dhcp / "p2p-air.json"),
                                    "DEVICES_CSV": str(cls.dhcp / "02-devices_config.csv"),
                                    "_AUTO_YES": False,
                                }
                                with mock.patch.multiple(PLATFORM_FLOW.DHCP, **bindings), \
                                        mock.patch.object(sys, "argv", list(command[1:])):
                                    PLATFORM_FLOW.DHCP.main()
                            elif script == "90-c2-generate_configs.py":
                                template = cls.nvos / "template"
                                ib_dir = template / f"99-output-ib_nvl/{PLATFORM_FLOW.TIMESTAMP}-ib"
                                nvl_dir = template / f"99-output-ib_nvl/{PLATFORM_FLOW.TIMESTAMP}-nvl"
                                with mock.patch.multiple(
                                    PLATFORM_FLOW.NVOS_GENERATOR,
                                    SCRIPT_DIR=str(template),
                                    _CSV_FILE=str(template / "02-devices_config.csv"),
                                    _GLOBAL_FILE=str(template / "01-global.yaml"),
                                    OUTPUT_IB_NVL_ROOT=str(template / "99-output-ib_nvl"),
                                    OUTPUT_IB_DIR=str(ib_dir),
                                    OUTPUT_NVL_DIR=str(nvl_dir),
                                ):
                                    generated = PLATFORM_FLOW.NVOS_GENERATOR._generate_all_ib(
                                        switch_scope=cls.selected_scope,
                                    )
                                if generated != 1:
                                    raise AssertionError((cls.selected_scope, generated))
                            elif script == "d-hostname2mac.py":
                                devices = PLATFORM_FLOW.PUBLISHER.load_csv(
                                    cls.devices_file, switch_scope=cls.selected_scope,
                                )
                                contexts = [
                                    PLATFORM_FLOW.PUBLISHER._nvos_dir_context(path)
                                    for path in command[2:]
                                    if str(path).endswith(f"-{cls.selected_scope}")
                                ]
                                if len(contexts) != 1:
                                    raise AssertionError((cls.selected_scope, command, contexts))
                                with mock.patch.object(PLATFORM_FLOW.PUBLISHER, "_AUTO_YES", True):
                                    published = PLATFORM_FLOW.PUBLISHER._publish_single_nvos(
                                        contexts[0], devices,
                                        deployment_scope="prod",
                                        switch_scope=cls.selected_scope,
                                    )
                                if not published:
                                    raise AssertionError("real scoped NVOS publisher rejected artifacts")
                            else:
                                raise AssertionError(f"unexpected real child: {command!r}")

                        with mock.patch.object(LOAD, "ZTP_DIR", cls.ztp), \
                                mock.patch.object(LOAD, "run", side_effect=real_child), \
                                redirect_stdout(io.StringIO()):
                            LOAD.generate_configs(
                                frozenset({cls.selected_scope}),
                                install_dhcp=False, dry_run=False, dhcp_enabled=True,
                                schema_version=1, deployment_scope="prod",
                                switch_scope=cls.selected_scope,
                            )
                        cls.executed = executed
                        _publish_private_dhcp_outputs(cls.project, cls.ztp)

                    @classmethod
                    def _run_cumulus_generator_and_publisher(cls, _p2p_output):
                        # A selected NVOS run does not execute Cumulus generation.
                        return None

                    @classmethod
                    def _run_nvos_generator_and_publisher(cls):
                        # Already executed by the real LOAD child dispatch above.
                        cls.nvos_release = (cls.nvos / "latest_yaml").resolve(strict=True)

                    @classmethod
                    def _run_parent_consumer(cls):
                        settings = LOAD.GlobalSettings(
                            dhcp_enabled=True, dhcp_package="isc-dhcp-server",
                            http_enabled=True, http_package="apache2",
                            http_root=cls.root, ztp_enabled=True,
                            ztp_prefix="/ztp", ztp_ips={"prod_oob": ("192.0.2.2",)},
                            versions={"cumulus": "5.16.4", "nvos": "25.02"},
                        )
                        inputs = LOAD.ProjectInputs(
                            global_file=cls.global_file, devices_file=cls.devices_file,
                            subnet_file=cls.subnet_file, p2p_file=cls.p2p_file,
                            device_types=frozenset({cls.selected_scope}),
                            pubkeys=(), settings=settings,
                            deployment_scope="prod",
                            switch_scope=cls.selected_scope,
                        )
                        with mock.patch.object(LOAD, "ZTP_DIR", cls.ztp):
                            LOAD.validate_dhcp_output_sources(LOAD.dhcp_file_mappings())
                            cls.parent = LOAD.validate_and_publish_release(
                                cls.project, inputs, publish=True,
                            )

                with redirect_stdout(io.StringIO()):
                    ScopedNVOSFlow._build_flow(private)
                self.assertEqual(
                    ["c1-generate_dhcp.py", "90-c2-generate_configs.py",
                     "d-hostname2mac.py"],
                    [script for script, _ in ScopedNVOSFlow.executed],
                )
                dhcp_dir = ScopedNVOSFlow.dhcp
                selected_hosts = dhcp_dir / f"dhcpd_{selected}.hosts"
                self.assertGreater(selected_hosts.stat().st_size, 0)
                for other in {"eth", "ib", "nvl"} - {selected}:
                    self.assertEqual(0, (dhcp_dir / f"dhcpd_{other}.hosts").stat().st_size)
                parent_path = ScopedNVOSFlow.project / "99-output-ztp/current-release.json"
                self.assertEqual(
                    ScopedNVOSFlow.parent["release_id"],
                    json.loads(parent_path.read_text())["release_id"],
                )
                settings = ACTIVATE.Settings(
                    http_root=private, project_name=ScopedNVOSFlow.project.name,
                    scope="prod", switch_scope=selected,
                )
                before = (_artifact_tree(ScopedNVOSFlow.project),
                          _artifact_tree(ScopedNVOSFlow.ztp))
                identity = ACTIVATE.validate_parent_release(settings, parent_path)
                self.assertEqual(before, (_artifact_tree(ScopedNVOSFlow.project),
                                          _artifact_tree(ScopedNVOSFlow.ztp)))
                self.assertEqual(
                    hashlib.sha256(parent_path.read_bytes()).hexdigest(),
                    identity["parent_release_sha256"],
                )

    def test_scoped_empty_dhcp_hosts_accepted_by_load_remain_activatable(self):
        """A LOAD-accepted zero-host family cannot invalidate its parent."""
        private_parent = os.environ.get("R22_TEST_PRIVATE_ROOT")
        with tempfile.TemporaryDirectory(
            prefix="r22-empty-host-activation-", dir=private_parent,
        ) as name:
            private = Path(name)
            project, ztp = _private_real_chain(private, "192.0.2.53")
            for relative in REAL_CHILD_SCRIPTS:
                self.assertEqual((ROOT / relative).read_bytes(),
                                 (private / relative).read_bytes())
            errors, warnings = SETUP._validate_v2_mlag_project(
                str(project / "01-global.yaml"),
                str(project / "02-devices_config.csv"),
            )
            self.assertEqual(([], []), (errors, warnings))
            args = SimpleNamespace(
                switch_scope="eth", deployment_scope="prod", p2p_file=None,
                no_upgrade=True, mini=None, ssh_dir=None, dry_run=True,
            )
            validated, _ = LOAD.validate_inputs(
                project, args, allow_management_key_generation=False,
                host_role="workstation",
            )
            self.assertEqual(frozenset({"eth"}), validated.device_types)

            original_subprocess = LOAD._run_subprocess
            executed = []

            def real_private_child(command, **kwargs):
                script = command[1] if len(command) > 1 else command[0]
                cwd = kwargs.get("cwd")
                self.assertIsNotNone(cwd)
                child_dir = Path(cwd).resolve(strict=True)
                self.assertTrue(child_dir.is_relative_to(private.resolve()), child_dir)
                self.assertIn(script, {Path(path).name for path in REAL_CHILD_SCRIPTS})
                self.assertTrue((child_dir / script).is_file())
                if not kwargs.get("capture_output"):
                    kwargs = {
                        **kwargs, "stdout": subprocess.PIPE,
                        "stderr": subprocess.STDOUT, "text": True,
                    }
                result = original_subprocess(command, timeout=120, **kwargs)
                executed.append((script, result.returncode))
                self.assertEqual(0, result.returncode, result.stdout)
                return result

            with (
                mock.patch.object(LOAD, "ZTP_DIR", ztp),
                mock.patch.object(LOAD, "_run_subprocess", side_effect=real_private_child),
                redirect_stdout(io.StringIO()),
            ):
                LOAD.generate_configs(
                    validated.device_types, install_dhcp=False, dry_run=False,
                    dhcp_enabled=True, schema_version=validated.settings.schema_version,
                    deployment_scope="prod", switch_scope="eth",
                    source_identities=validated.source_identities,
                )
                self.assertEqual(
                    [Path(path).name for path in REAL_CHILD_SCRIPTS],
                    [script for script, _ in executed],
                )
                _publish_private_dhcp_outputs(project, ztp)
                dhcp_dir = ztp / "config/isc-dhcp-server"
                self.assertEqual(
                    0, (dhcp_dir / "dhcpd_ib.hosts").stat().st_size,
                )
                self.assertEqual(
                    0, (dhcp_dir / "dhcpd_nvl.hosts").stat().st_size,
                )
                LOAD.validate_dhcp_output_sources(LOAD.dhcp_file_mappings())
                parent = LOAD.validate_and_publish_release(project, validated)

            parent_path = project / "99-output-ztp/current-release.json"
            self.assertEqual(parent["release_id"],
                             json.loads(parent_path.read_text())["release_id"])
            settings = ACTIVATE.Settings(
                http_root=private, project_name=project.name,
                scope="prod", switch_scope="eth",
            )
            before = (_artifact_tree(project), _artifact_tree(ztp))
            activation_error = None
            try:
                identity = ACTIVATE.validate_parent_release(settings, parent_path)
            except ACTIVATE.ActivationError as error:
                activation_error = str(error)
            self.assertEqual(before, (_artifact_tree(project), _artifact_tree(ztp)))
            self.assertIsNone(activation_error, activation_error)
            self.assertEqual(
                hashlib.sha256(parent_path.read_bytes()).hexdigest(),
                identity["parent_release_sha256"],
            )

    def test_real_load_child_chain_rejects_same_inode_global_aba_before_publisher(self):
        """A c2-only B read cannot publish B after LOAD validated A."""
        private_parent = os.environ.get("R22_TEST_PRIVATE_ROOT")
        with tempfile.TemporaryDirectory(
            prefix="r22-real-generation-", dir=private_parent,
        ) as name:
            # macOS may report /var for a /private/var subprocess cwd; the
            # candidate API deliberately requires one canonical project path.
            private = Path(name).resolve()
            project, ztp = _private_real_chain(
                private / "A", "192.0.2.53", http_enabled=True,
            )
            project_b, _ = _private_real_chain(
                private / "B", "192.0.2.54", http_enabled=True,
            )
            # Activation currently requires nonempty private includes even
            # when this scoped ETH run does not own IB/NVL.  Valid DHCP
            # comments model pre-existing, unchanged family files; c1 still
            # produces and hashes the complete five-output child receipt.
            for family in ("ib", "nvl"):
                (ztp / "config/isc-dhcp-server" / f"dhcpd_{family}.hosts").write_text(
                    "# no selected hosts for this family\n", encoding="utf-8",
                )
            for relative in REAL_CHILD_SCRIPTS:
                self.assertEqual(
                    (ROOT / relative).read_bytes(),
                    (private / "A" / relative).read_bytes(),
                )
            source = project / "01-global.yaml"
            bytes_a = source.read_bytes()
            bytes_b = (project_b / "01-global.yaml").read_bytes()
            self.assertNotEqual(bytes_a, bytes_b)
            self.assertEqual(len(bytes_a), len(bytes_b))
            self.assertIn(b"192.0.2.53", bytes_a)
            self.assertIn(b"192.0.2.54", bytes_b)

            args = SimpleNamespace(
                switch_scope="eth", deployment_scope="prod", p2p_file=None,
                no_upgrade=True, mini=None, ssh_dir=None, dry_run=True,
            )
            for candidate in (project, project_b):
                errors, warnings = SETUP._validate_v2_mlag_project(
                    str(candidate / "01-global.yaml"),
                    str(candidate / "02-devices_config.csv"),
                )
                self.assertEqual(([], []), (errors, warnings))
                validated, _ = LOAD.validate_inputs(
                    candidate, args, allow_management_key_generation=False,
                    host_role="workstation",
                )
                self.assertEqual(frozenset({"eth"}), validated.device_types)

            original_subprocess = LOAD._run_subprocess
            state = {
                "inject": False, "injected": False, "children": [],
                "post_c1_activation_error": None,
            }
            original = None
            transcript = io.StringIO()

            def real_private_child(command, **kwargs):
                script = command[1] if len(command) > 1 else command[0]
                cwd = kwargs.get("cwd")
                if cwd is None:
                    self.assertEqual("/usr/bin/ssh-keygen", command[0])
                    self.assertEqual("-l", command[1])
                else:
                    child_dir = Path(cwd).resolve()
                    self.assertTrue(child_dir.is_relative_to(private.resolve()), child_dir)
                    self.assertIn(script, {Path(path).name for path in REAL_CHILD_SCRIPTS})
                    self.assertTrue((child_dir / script).is_file())
                    if not kwargs.get("capture_output"):
                        kwargs = {
                            **kwargs, "stdout": subprocess.PIPE,
                            "stderr": subprocess.STDOUT, "text": True,
                        }
                is_c2 = script == "90-c2-generate_configs.py" and state["inject"]
                if is_c2:
                    self.assertFalse(state["injected"])
                    state["injected"] = True
                    _write_same_inode(source, bytes_b, original)
                started = time.monotonic()
                try:
                    result = original_subprocess(command, timeout=120, **kwargs)
                finally:
                    if is_c2:
                        _write_same_inode(source, bytes_a, original)
                state["children"].append((script, result.returncode))
                if cwd is not None and result.stdout:
                    transcript.write(result.stdout)
                if result.returncode and (
                    not state["inject"] or script == "c1-generate_dhcp.py"
                ):
                    self.fail(
                        f"real private child {script} exited {result.returncode}: "
                        f"{result.stdout!r}"
                    )
                if state["inject"] and script == "c1-generate_dhcp.py":
                    try:
                        ACTIVATE.validate_parent_release(settings, parent_path)
                    except ACTIVATE.ActivationError as error:
                        state["post_c1_activation_error"] = str(error)
                self.assertLess(time.monotonic() - started, 120)
                return result

            def run_chain(validated, *, dhcp_candidate_dir=None):
                candidate_options = (
                    {} if dhcp_candidate_dir is None
                    else {"dhcp_candidate_dir": dhcp_candidate_dir}
                )
                return LOAD.generate_configs(
                    validated.device_types, install_dhcp=False, dry_run=False,
                    dhcp_enabled=validated.settings.dhcp_enabled,
                    schema_version=validated.settings.schema_version,
                    deployment_scope="prod", switch_scope="eth",
                    source_identities=validated.source_identities,
                    **candidate_options,
                )

            with (
                mock.patch.object(LOAD, "ZTP_DIR", ztp),
                mock.patch.object(LOAD, "_run_subprocess", side_effect=real_private_child),
                redirect_stdout(transcript),
            ):
                validated_a, _ = LOAD.validate_inputs(
                    project, args, allow_management_key_generation=False,
                    host_role="workstation",
                )
                self.assertEqual(
                    hashlib.sha256(bytes_a).hexdigest(),
                    validated_a.source_identities["global"],
                )
                run_chain(validated_a)

                # Publish the real c1 child manifest through the same
                # project-owned path and canonical public link that Docker
                # activation requires.  No release JSON or hash is forged.
                dhcp_dir = ztp / "config/isc-dhcp-server"
                _publish_private_dhcp_outputs(project, ztp)
                parent = LOAD.validate_and_publish_release(project, validated_a)
                parent_path = project / "99-output-ztp/current-release.json"
                self.assertEqual(parent["release_id"],
                                 json.loads(parent_path.read_text())["release_id"])
                settings = ACTIVATE.Settings(
                    http_root=private / "A", project_name=project.name,
                    scope="prod", switch_scope="eth",
                )
                accepted_parent = ACTIVATE.validate_parent_release(
                    settings, parent_path,
                )
                self.assertEqual(
                    hashlib.sha256(parent_path.read_bytes()).hexdigest(),
                    accepted_parent["parent_release_sha256"],
                )
                # The same preflight LOAD now uses before any B generation
                # must migrate the already committed A while its source bytes
                # and published DHCP child are still independently verifiable.
                self.assertEqual(
                    f"generations/{parent['release_id']}",
                    LOAD.ensure_previous_generation_frozen(project),
                )
                committed_a = (
                    project / "99-output-ztp/generations" / parent["release_id"]
                )
                selector = project / "99-output-ztp/current"
                self.assertEqual(
                    parent["release_id"],
                    ACTIVATE.validate_pinned_generation(
                        settings, selector, expected_release_id=parent["release_id"],
                    )["release_id"],
                )
                # LOAD has selected A, so ACTIVATE's legacy flat-input writer
                # must not publish a runtime plan without binding that selector.
                with mock.patch.object(
                    ACTIVATE, "ensure_runtime_directories",
                ) as runtime_dirs, mock.patch.object(
                    ACTIVATE, "_atomic_write",
                ) as runtime_write:
                    with self.assertRaisesRegex(
                        ACTIVATE.ActivationError, "selector|generation",
                    ):
                        ACTIVATE.prepare_runtime(settings)
                runtime_dirs.assert_not_called()
                runtime_write.assert_not_called()
                frozen_global = committed_a / "inputs/global"
                original_global = frozen_global.read_bytes()
                frozen_global.write_bytes(b"other generation input\n")
                with self.assertRaisesRegex(ACTIVATE.ActivationError, "pinned input"):
                    ACTIVATE.validate_pinned_generation(
                        settings, selector, expected_release_id=parent["release_id"],
                    )
                frozen_global.write_bytes(original_global)

                # The prior release is a real successful publication, not a
                # same-timestamp c2 staging directory or a synthetic marker.
                latest = project / "99-output-eth/latest"
                prior_name = os.readlink(latest)
                prior_release = latest.parent / prior_name
                prior_state = _artifact_tree(prior_release)
                self.assertIn("Border01.yaml", prior_state)
                dhcp_dir = ztp / "config/isc-dhcp-server"
                dhcp_before = {
                    item: (dhcp_dir / item).read_bytes() for item in DHCP_OUTPUTS
                }
                boundary = int(time.time())
                while int(time.time()) == boundary:
                    time.sleep(0.025)

                original = source.stat()
                state["inject"] = True
                private_candidate = (
                    project / "99-output-dhcp" /
                    ".load-dhcp-candidate-0123456789abcdef"
                )
                private_candidate.mkdir(mode=0o700)
                validated_a, _ = LOAD.validate_inputs(
                    project, args, allow_management_key_generation=False,
                    host_role="workstation",
                )
                self.assertEqual(
                    hashlib.sha256(bytes_a).hexdigest(),
                    validated_a.source_identities["global"],
                )
                failure = None
                try:
                    run_chain(validated_a, dhcp_candidate_dir=private_candidate)
                except LOAD.LoadError as error:
                    failure = error

            self.assertTrue(state["injected"], state)
            self.assertEqual(
                parent["release_id"],
                ACTIVATE.validate_pinned_generation(
                    settings, selector, expected_release_id=parent["release_id"],
                )["release_id"],
            )
            self.assertEqual(
                (original.st_ino, original.st_size, original.st_mtime_ns),
                (source.stat().st_ino, source.stat().st_size, source.stat().st_mtime_ns),
            )
            self.assertEqual(bytes_a, source.read_bytes())
            final_name = os.readlink(latest)
            final_yaml = (latest / "Border01.yaml").read_bytes()
            dhcp_after = {item: (dhcp_dir / item).read_bytes() for item in DHCP_OUTPUTS}
            dhcp_changed = [
                key for key in DHCP_OUTPUTS if dhcp_before[key] != dhcp_after[key]
            ]
            detail = (
                f"children={state['children']!r}; prior={prior_name}; final={final_name}; "
                f"new YAML sha256={hashlib.sha256(final_yaml).hexdigest()}; "
                f"DHCP changed={dhcp_changed}; "
                f"post-c1 activation={state['post_c1_activation_error']!r}; "
                f"child output tail={transcript.getvalue()[-1000:]!r}"
            )
            with self.subTest("accepted parent survives completed c1 before c2"):
                self.assertIsNone(state["post_c1_activation_error"], detail)
            with self.subTest("refuse drift from validated A"):
                self.assertIsNotNone(failure, detail)
                self.assertRegex(
                    (str(failure) + transcript.getvalue()).lower(),
                    r"input|source|authority|drift|chang",
                )
            with self.subTest("latest remains the previously published release"):
                self.assertEqual(prior_name, final_name, detail)
            with self.subTest("prior published release bytes are untouched"):
                self.assertEqual(prior_state, _artifact_tree(prior_release), detail)
            with self.subTest("validated A never publishes B"):
                self.assertFalse(b"192.0.2.54" in final_yaml, detail)
            with self.subTest("refusal leaves five prior DHCP outputs untouched"):
                self.assertEqual([], dhcp_changed, detail)
            with self.subTest("c1 wrote private complete candidate before c2 refused B"):
                self.assertEqual(
                    set(DHCP_OUTPUTS),
                    {item.name for item in private_candidate.iterdir()}, detail,
                )
            # The actual LOAD child chain has left the previously committed A
            # selected.  A subsequent, independently legal B edit to the
            # mutable source must not be used when ACTIVATE observes A's plan.
            # This is read-only planning, not a service activation receipt.
            network_links = [{
                "ifindex": 4, "ifname": "eth0",
                "flags": ["BROADCAST", "MULTICAST", "UP", "LOWER_UP"],
                "address": "02:00:00:00:00:04", "operstate": "UP",
                "link_type": "ether",
            }]
            network_addresses = [{
                "ifindex": 4, "ifname": "eth0",
                "addr_info": [{
                    "family": "inet", "local": "192.0.2.20", "prefixlen": 24,
                    "scope": "global", "flags": [],
                    "valid_life_time": "forever",
                    "preferred_life_time": "forever",
                }],
            }]
            source.write_bytes(bytes_b)
            try:
                with mock.patch.object(
                    ACTIVATE, "capture_stable_network_snapshot",
                    return_value=(network_links, network_addresses),
                ):
                    _plan, _runtime, payload, _apache = (
                        ACTIVATE.observe_pinned_runtime(
                            settings, selector,
                            expected_release_id=parent["release_id"],
                        )
                    )
                self.assertEqual(parent["release_id"],
                                 payload["selected_generation_release_id"])
                self.assertEqual(
                    hashlib.sha256((committed_a / "inputs/subnet").read_bytes()).hexdigest(),
                    payload["subnet_sha256"],
                )
                # Exercise the real ACTIVATE publication boundary after the
                # real LOAD/c1/c2/d chain rejects B.  Keep every destination
                # private and leave service-start authority absent.
                pinned_state = private / "pinned-state"
                pinned_state.mkdir()
                pinned_settings = ACTIVATE.Settings(
                    http_root=private / "A", project_name=project.name,
                    scope="prod", switch_scope="eth",
                    state_root=pinned_state, log_root=private / "pinned-log",
                    apache_listeners=private / "pinned-apache-listeners.conf",
                    dhcp_config=private / "pinned-dhcpd.conf",
                    dhcp_leases=private / "pinned-leases",
                )
                pinned_settings.activation_marker.write_text("old marker\n")
                pinned_settings.precommit_marker.write_text("old precommit\n")
                with mock.patch.object(
                    ACTIVATE, "capture_stable_network_snapshot",
                    return_value=(network_links, network_addresses),
                ), mock.patch.object(ACTIVATE, "ensure_runtime_directories"):
                    _selected, _runtime, published = ACTIVATE.prepare_pinned_runtime(
                        pinned_settings, selector_path=selector,
                        expected_release_id=parent["release_id"],
                    )
                self.assertEqual(parent["release_id"],
                                 published["selected_generation_release_id"])
                self.assertEqual(
                    published,
                    json.loads(pinned_settings.runtime_plan.read_text()),
                )
                self.assertTrue(pinned_settings.apache_listeners.is_file())
                self.assertFalse(pinned_settings.activation_marker.exists())
                self.assertFalse(pinned_settings.precommit_marker.exists())
            finally:
                source.write_bytes(bytes_a)

    def test_load_setup_c2_rejects_same_inode_second_read_and_preserves_existing_same_name_output(self):
        rows_a = V2MlagSetupPreflightTests().pair()
        rows_b = copy.deepcopy(rows_a)
        rows_b[0][0] = "Border03"
        rows_b[1][0] = "Border04"
        with tempfile.TemporaryDirectory(prefix="r22-generation-workflow-") as name:
            private = Path(name)
            project_a = private / "project-a"
            project_b = private / "project-b"
            _write_legal_project(project_a, rows_a)
            _write_legal_project(project_b, rows_b)
            source = project_a / "02-devices_config.csv"
            bytes_a = source.read_bytes()
            bytes_b = (project_b / "02-devices_config.csv").read_bytes()
            self.assertNotEqual(bytes_a, bytes_b)
            self.assertEqual(len(bytes_a), len(bytes_b))

            # Both generations independently pass the actual setup validator.
            for project in (project_a, project_b):
                errors, warnings = SETUP._validate_v2_mlag_project(
                    str(project / "01-global.yaml"),
                    str(project / "02-devices_config.csv"),
                )
                self.assertEqual(([], []), (errors, warnings))

            args = SimpleNamespace(
                switch_scope="all", deployment_scope="prod", p2p_file=None,
                no_upgrade=True, mini=None, ssh_dir=None, dry_run=True,
            )
            inputs_a, images_a = LOAD.validate_inputs(
                project_a, args, allow_management_key_generation=False,
                host_role="workstation",
            )
            inputs_b, images_b = LOAD.validate_inputs(
                project_b, args, allow_management_key_generation=False,
                host_role="workstation",
            )
            self.assertEqual({}, images_a)
            self.assertEqual({}, images_b)
            self.assertEqual(frozenset({"eth"}), inputs_a.device_types)
            self.assertEqual(frozenset({"eth"}), inputs_b.device_types)
            self.assertEqual(
                hashlib.sha256(bytes_a).hexdigest(),
                inputs_a.source_identities["devices"],
            )
            self.assertEqual(
                hashlib.sha256(bytes_b).hexdigest(),
                inputs_b.source_identities["devices"],
            )

            dot = private / "fixture-lldpq.dot"
            dot.write_text("", encoding="utf-8")
            generation = private / "generation"
            generation.mkdir()
            (generation / "new-marker.txt").write_bytes(b"candidate-from-A")
            existing = private / "generation_with_desc"
            existing.mkdir()
            old_artifact = existing / "old-marker.txt"
            old_artifact.write_bytes(b"existing-same-name")
            initial = source.stat()
            real_copytree = C2.shutil.copytree
            observed = {"injected": False}

            def mutate_after_copy(src, dst, *copy_args, **copy_kwargs):
                result = real_copytree(src, dst, *copy_args, **copy_kwargs)
                source.write_bytes(bytes_b)
                os.utime(source, ns=(initial.st_atime_ns, initial.st_mtime_ns))
                changed = source.stat()
                observed.update({
                    "injected": True,
                    "same_inode": changed.st_ino == initial.st_ino,
                    "same_size": changed.st_size == initial.st_size,
                    "same_mtime": changed.st_mtime_ns == initial.st_mtime_ns,
                })
                return result

            transcript = io.StringIO()
            failure = None
            with (
                mock.patch.object(C2, "_CSV_FILE", str(source)),
                mock.patch.object(C2.shutil, "copytree", side_effect=mutate_after_copy),
                redirect_stdout(transcript),
            ):
                try:
                    C2._run_patch_descriptions(str(dot), str(generation))
                except Exception as error:
                    failure = error

            self.assertTrue(observed["injected"], observed)
            self.assertTrue(observed["same_inode"], observed)
            self.assertTrue(observed["same_size"], observed)
            self.assertTrue(observed["same_mtime"], observed)
            with self.subTest("refuse cross-generation input"):
                self.assertIsNotNone(
                    failure,
                    "c2 accepted project B after setup/load validated and bound A; "
                    f"actual c2 output was: {transcript.getvalue()!r}",
                )
                self.assertRegex(
                    str(failure).lower(), r"input|source|authority|drift|chang",
                )
            with self.subTest("preserve same-name existing output"):
                self.assertTrue(old_artifact.is_file())
                self.assertEqual(
                    b"existing-same-name", old_artifact.read_bytes(),
                )
                self.assertFalse((existing / "new-marker.txt").exists())


if __name__ == "__main__":
    unittest.main()
