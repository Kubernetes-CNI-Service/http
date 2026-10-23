"""The CVT lives in the project's 99-output-p2p (owner 2026-10-21: "放项目的 99-output-p2p，和历史用法以及
ETH 一致"), beside Ethernet's <p2p>-lldpq.dot, and the converter runs on its own from its directory.

Only that one place opens: every other DAY0 project path stays refused, so the two REQ7 refusal tests
(test_req7_ib_provenance*.py) keep their meaning (infiniband/monitor/UFM-MONITOR.md section 13)."""

import hashlib
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import xlsxwriter

ROOT = Path(__file__).resolve().parents[1]
CONVERTER = ROOT / "ztp/config/nvos/template/P2P/p2p-to-validation.py"


def load_converter():
    spec = importlib.util.spec_from_file_location("ib_cvt_project_output_converter", CONVERTER)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    try:
        spec.loader.exec_module(module)
    finally:
        sys.modules.pop(spec.name, None)
    return module


def write_p2p(path: Path) -> None:
    book = xlsxwriter.Workbook(str(path))
    sheet = book.add_worksheet("CL links")
    for column, value in enumerate(("Name", "Port", "Name", "Port")):
        sheet.write(0, column, value)
    for column, value in enumerate(("leaf01", "sw1p1", "server01", "mlx5_0")):
        sheet.write(1, column, value)
    book.close()


class CvtProjectOutputDirectTest(unittest.TestCase):
    def setUp(self):
        self.converter = load_converter()
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name).resolve()
        self.config = self.root / "ztp/config"
        self.config.mkdir(parents=True)
        self.project = self.root / "DAY0-Prepare/demo"
        (self.project / "99-output-p2p").mkdir(parents=True)
        patcher = mock.patch.object(self.converter, "CONFIG_ROOT", self.config)
        patcher.start()
        self.addCleanup(patcher.stop)

    def tearDown(self):
        self.tmp.cleanup()

    def test_a_cvt_may_go_into_the_project_99_output_p2p(self):
        target = self.project / "99-output-p2p/fabric-blue-cvt.xlsx"
        self.converter.publish_cvt_workbook(target, [], [], [], [])
        self.assertTrue(target.is_file())

    def test_every_other_day0_path_stays_refused(self):
        (self.project / "01-global.yaml").write_bytes(b"PRESERVE\n")
        for relative in ("01-global.yaml", "99-output-p2p/report.xlsx", "99-output-p2p/sub/x-cvt.xlsx",
                         "99-output-ztp/x-cvt.xlsx", "x-cvt.xlsx"):
            with self.subTest(relative=relative):
                with self.assertRaises(self.converter.ConversionError):
                    self.converter.publish_cvt_workbook(self.project / relative, [], [], [], [])
        self.assertEqual(b"PRESERVE\n", (self.project / "01-global.yaml").read_bytes())
        self.assertFalse((self.project / "99-output-p2p/report.xlsx").exists())


class ShippedInventoryRulesTest(unittest.TestCase):
    """Owner 2026-10-22: the shipped 01-inventory.log also names IB switches called ew-spineNNN /
    ew-leafNNN (the B300 fabric of 2026-10, short or FQDN); a project with other names adds its own."""

    def test_ew_spine_and_ew_leaf_are_ib_switches_and_other_names_keep_their_type(self):
        converter = load_converter()
        p2p = ROOT / "ztp/config/cumulus/template/P2P"
        rules = converter.RuleSet(p2p / "01-inventory.log", p2p / "02-port-mapping.log", p2p / "03-splitter.log")
        expected = {
            "ew-spine001.site.example.test": "IB-SW", "ew-leaf108": "IB-SW", "EW-LEAF001": "IB-SW",
            "IB-Spine01": "IB-SW", "IBSW-SU1-Leaf01": "IB-SW", "TAN-SU1-Leaf01": "Eth-SW",
            "OOB-SU1-Leaf01": "Eth-SW", "hgx0019": "SMC-B300", "ufm01": "ufm",
            # the B300 GPU servers by their P2P Hostname (owner 2026-10-22: GPUs use that column too)
            "su01-gpu-b300l-019.site.example.test": "SMC-B300", "gpu-gb300-001": "GB300-GPU",
        }
        self.assertEqual({name: rules.device_type(name) for name in expected}, expected)


DEVICES_CONFIG = ("hostname,type,eth0_ip\n"
                  "ew-leaf001.site.example.test,ib,192.0.2.11\n"
                  "ew-spine001.site.example.test,ib,192.0.2.12\n"
                  "ufm01.site.example.test,server_ufm,192.0.2.30\n")
INVENTORY_RULES = "[IB-SW]\n*ew-leaf*\n*ew-spine*\n\n[SMC-B300]\n*gpu*\n*hgx*\n\n[ufm]\n*ufm*\n"
PORT_RULES = "[IB-SW]\n1to2#/1,sw#p1\n1to2#/2,sw#p2\n"


def write_endpoint_sheets(path: Path, gpu_rows, switch_rows, *, name_first=True) -> None:
    """Two CL/CS sheets laid out like the 2026-10 B300 P2P: each side has two name columns before its
    port column. name_first=True: `name` (a rack label) then `Hostname`; False: `Hostname` then `name`."""
    book = xlsxwriter.Workbook(str(path))
    labels = ("name", "Hostname") if name_first else ("Hostname", "name")
    for sheet_name, port_header, rows in (("IB GPU-CL", "HCA/port", gpu_rows), ("IB CL-CS", "port", switch_rows)):
        sheet = book.add_worksheet(sheet_name)
        header = ("Room", "rack", "U", *labels, port_header, "Transceiver", "Room", "rack", "U", *labels, "port")
        for column, value in enumerate(header):
            sheet.write(0, column, value)
        for number, row in enumerate(rows, 1):
            first_a, second_a, port_a, first_b, second_b, port_b = row
            values = ("HALL", "R1", "1", first_a, second_a, port_a, "X", "HALL", "R2", "2", first_b, second_b, port_b)
            for column, value in enumerate(values):
                sheet.write(number, column, value)
    book.close()


class CvtDeviceNameColumnTest(unittest.TestCase):
    """Owner 2026-10-22: the converter finds which P2P column holds the device names by matching it
    against the hostnames in 02-devices_config.csv, and every device (GPUs too) takes its name from
    that column. Without 02-devices_config.csv the column next to the port is used, as before."""

    def run_converter(self, gpu_rows, switch_rows, *, devices=True, name_first=True):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        root = Path(self.tmp.name).resolve()
        config = root / "ztp/config"
        p2p_dir = config / "nvos/template/P2P"
        p2p_dir.mkdir(parents=True)
        shutil.copy2(CONVERTER, p2p_dir / "p2p-to-validation.py")
        for name in ("topology_rules.py", "ib_topology_provenance.py"):
            shutil.copy2(ROOT / "ztp/config" / name, config / name)
        project = root / "DAY0-Prepare/fabric-a"
        output = project / "99-output-p2p"
        output.mkdir(parents=True)
        write_endpoint_sheets(project / "fabric-blue.xlsx", gpu_rows, switch_rows, name_first=name_first)
        (project / "p2p.xlsx").symlink_to("fabric-blue.xlsx")
        (project / "02-devices_config.csv").write_text(DEVICES_CONFIG, encoding="utf-8")
        # the links setup creates in the converter's own directory
        (p2p_dir / "p2p.xlsx").symlink_to("../../../../../DAY0-Prepare/fabric-a/p2p.xlsx")
        (p2p_dir / "output-p2p").symlink_to("../../../../../DAY0-Prepare/fabric-a/99-output-p2p")
        if devices:
            (p2p_dir / "02-devices_config.csv").symlink_to("../../../../../DAY0-Prepare/fabric-a/02-devices_config.csv")
        for name, text in (("01-inventory.log", INVENTORY_RULES), ("02-port-mapping.log", PORT_RULES),
                           ("03-splitter.log", "")):
            (p2p_dir / name).write_text(text, encoding="utf-8")
        completed = subprocess.run(
            [sys.executable, "-B", str(p2p_dir / "p2p-to-validation.py")],
            cwd=p2p_dir, text=True, capture_output=True, timeout=120, check=False,
            env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
        return completed, output / "fabric-blue-cvt.xlsx"

    @staticmethod
    def links(cvt: Path) -> set:
        import openpyxl
        book = openpyxl.load_workbook(cvt, read_only=True, data_only=True)
        rows = list(book["Links"].iter_rows(min_row=2, values_only=True))
        return {(row[1], row[2], row[6], row[7]) for row in rows if row and row[1]}

    GPU = [("hgx0001", "gpu-001.site.example.test", "1", "IBSW-Leaf01", "ew-leaf001", "1/1")]
    SWITCHES = [("IBSW-Leaf01", "ew-leaf001", "37/1", "IBSW-Spine01", "ew-spine001", "1/1")]

    def test_the_column_holding_devices_config_names_names_every_device(self):
        completed, cvt = self.run_converter(self.GPU, self.SWITCHES)
        self.assertEqual(0, completed.returncode, completed.stdout + completed.stderr)
        self.assertEqual(self.links(cvt), {
            ("gpu-001.site.example.test", "1", "ew-leaf001", "sw1p1"),
            ("ew-leaf001", "sw37p1", "ew-spine001", "sw1p1")})
        self.assertIn("'Hostname'", completed.stdout)
        provenance = json.loads(Path(f"{cvt}.provenance.json").read_text(encoding="utf-8"))
        self.assertTrue(provenance["sources"]["devices"]["resolved_path"].endswith("02-devices_config.csv"))
        sys.path.insert(0, str(ROOT / "ztp/config"))
        try:
            from ib_topology_provenance import ProvenanceError, read_cvt_provenance
            self.assertEqual(read_cvt_provenance(cvt)["sources"]["devices"], provenance["sources"]["devices"])
            # only `devices` may join the six roles: any other extra role is refused
            sidecar = Path(f"{cvt}.provenance.json")
            provenance["sources"]["extra"] = provenance["sources"]["devices"]
            sidecar.write_text(json.dumps(provenance), encoding="utf-8")
            with self.assertRaises(ProvenanceError):
                read_cvt_provenance(cvt)
        finally:
            sys.path.remove(str(ROOT / "ztp/config"))

    def test_the_matching_column_wins_even_when_it_is_not_next_to_the_port(self):
        # here `name` holds the switch names and sits before `Hostname`, which holds rack labels
        gpu = [("gpu-001.site.example.test", "R3-U7", "1", "ew-leaf001", "IBSW-Leaf01", "1/1")]
        switches = [("ew-leaf001", "IBSW-Leaf01", "37/1", "ew-spine001", "IBSW-Spine01", "1/1")]
        completed, cvt = self.run_converter(gpu, switches, name_first=True)
        self.assertEqual(0, completed.returncode, completed.stdout + completed.stderr)
        self.assertEqual(self.links(cvt), {
            ("gpu-001.site.example.test", "1", "ew-leaf001", "sw1p1"),
            ("ew-leaf001", "sw37p1", "ew-spine001", "sw1p1")})

    def test_without_devices_config_the_column_next_to_the_port_is_used(self):
        gpu = [("gpu-001.site.example.test", "R3-U7", "1", "ew-leaf001", "IBSW-Leaf01", "1/1")]
        switches = [("ew-leaf001", "IBSW-Leaf01", "37/1", "ew-spine001", "IBSW-Spine01", "1/1")]
        completed, cvt = self.run_converter(gpu, switches, devices=False)
        self.assertEqual(0, completed.returncode, completed.stdout + completed.stderr)
        self.assertEqual({(a, c) for a, _b, c, _d in self.links(cvt)},
                         {("R3-U7", "IBSW-Leaf01"), ("IBSW-Leaf01", "IBSW-Spine01")})
        self.assertIn("02-devices_config.csv", completed.stdout)
        provenance = json.loads(Path(f"{cvt}.provenance.json").read_text(encoding="utf-8"))
        self.assertNotIn("devices", provenance["sources"], "the six source roles as before")

    def test_switch_and_ufm_names_missing_from_devices_config_are_reported(self):
        gpu = self.GPU + [("UFM-01", "ufm-01", "IB0", "IBSW-Leaf01", "ew-leaf001", "FNM")]
        completed, _cvt = self.run_converter(gpu, self.SWITCHES)
        self.assertEqual(0, completed.returncode, completed.stdout + completed.stderr)
        self.assertIn("ufm-01", completed.stderr)
        self.assertNotIn("gpu-001", completed.stderr, "GPUs are not in 02-devices_config.csv")

    def test_no_link_record_publishes_nothing(self):
        completed, cvt = self.run_converter([], [])
        self.assertNotEqual(0, completed.returncode)
        self.assertIn("0", completed.stderr)
        self.assertFalse(cvt.exists())


class CvtProjectOutputWorkflowTest(unittest.TestCase):
    def test_the_converter_run_on_its_own_publishes_the_cvt_into_the_project(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory).resolve()
            config = root / "ztp/config"
            p2p_dir = config / "nvos/template/P2P"
            p2p_dir.mkdir(parents=True)
            shutil.copy2(CONVERTER, p2p_dir / "p2p-to-validation.py")
            for name in ("topology_rules.py", "ib_topology_provenance.py"):
                shutil.copy2(ROOT / "ztp/config" / name, config / name)
            project = root / "DAY0-Prepare/fabric-a"
            output = project / "99-output-p2p"
            output.mkdir(parents=True)
            write_p2p(project / "fabric-blue.xlsx")
            (project / "p2p.xlsx").symlink_to("fabric-blue.xlsx")
            # the links setup creates in the converter's own directory
            (p2p_dir / "p2p.xlsx").symlink_to("../../../../../DAY0-Prepare/fabric-a/p2p.xlsx")
            (p2p_dir / "output-p2p").symlink_to("../../../../../DAY0-Prepare/fabric-a/99-output-p2p")
            for name, text in (("01-inventory.log", "[IB-SW]\n*leaf*\n\n[server]\n*server*\n"),
                               ("02-port-mapping.log", ""), ("03-splitter.log", "")):
                (p2p_dir / name).write_text(text, encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, "-B", str(p2p_dir / "p2p-to-validation.py")],
                cwd=p2p_dir, text=True, capture_output=True, timeout=120, check=False,
                env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
            self.assertEqual(0, completed.returncode, completed.stdout + completed.stderr)
            cvt = output / "fabric-blue-cvt.xlsx"
            self.assertTrue(cvt.is_file(), completed.stdout)
            self.assertFalse(cvt.is_symlink())
            provenance = json.loads(Path(f"{cvt}.provenance.json").read_text(encoding="utf-8"))
            self.assertEqual(str(project / "fabric-blue.xlsx"), provenance["sources"]["p2p"]["resolved_path"])
            # a second run replaces it through the same link
            again = subprocess.run(
                [sys.executable, "-B", str(p2p_dir / "p2p-to-validation.py")],
                cwd=p2p_dir, text=True, capture_output=True, timeout=120, check=False,
                env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
            self.assertEqual(0, again.returncode, again.stdout + again.stderr)

            # load binds the run to the P2P bytes it validated (--expect-p2p-sha256)
            published = cvt.read_bytes()
            wrong = subprocess.run(
                [sys.executable, "-B", str(p2p_dir / "p2p-to-validation.py"), "--expect-p2p-sha256", "0" * 64],
                cwd=p2p_dir, text=True, capture_output=True, timeout=120, check=False,
                env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
            self.assertNotEqual(0, wrong.returncode)
            self.assertIn("P2P", wrong.stderr)
            self.assertEqual(published, cvt.read_bytes(), "nothing is published for another P2P")
            right = subprocess.run(
                [sys.executable, "-B", str(p2p_dir / "p2p-to-validation.py"), "--expect-p2p-sha256",
                 hashlib.sha256((project / "fabric-blue.xlsx").read_bytes()).hexdigest()],
                cwd=p2p_dir, text=True, capture_output=True, timeout=120, check=False,
                env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
            self.assertEqual(0, right.returncode, right.stdout + right.stderr)


if __name__ == "__main__":
    unittest.main()
