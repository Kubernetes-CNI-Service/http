#!/usr/bin/env python3
"""Literal P2P row-to-LLDPQ source authority, never a qualified cycle."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import tempfile
import unittest
import warnings
import zipfile

from openpyxl import Workbook, load_workbook


EXPECTED = (
    ("TAN-Links", 3, "leaf-a", "swp1", "leaf-z", "swp2"),
    ("TAN-Links", 4, "leaf-b", "swp3", "leaf-c", "swp4"),
)
DOT_LINES = (
    '"leaf-a":"swp1" -- "leaf-z":"swp2"',
    '"leaf-b":"swp3" -- "leaf-c":"swp4"',
)
ROLES = ("topology", "inventory", "port_mapping", "lldpq", "splitter_profiles")


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def literal_fixture(root: Path) -> dict[str, Path]:
    """Independent source rows: expected endpoints are written above, not parsed back."""
    workbook = root / "selected-fabric.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "TAN-Links"
    sheet.append(["Source", "", "Dest", ""])
    sheet.append(["name", "port", "name", "port"])
    sheet.append(["leaf-a", "swp1", "leaf-z", "swp2"])
    sheet.append(["leaf-b", "swp3", "leaf-c", "swp4"])
    sheet.append(["Empty", "Empty", "leaf-d", "swp5"])
    book.save(workbook)
    book.close()

    inventory = root / "01-inventory.log"
    inventory.write_text("[Eth-SW]\nleaf-*\n", encoding="utf-8")
    port_mapping = root / "02-port-mapping.log"
    port_mapping.write_text(
        "[Eth-SW]\nswp1,swp1\nswp2,swp2\nswp3,swp3\n"
        "swp4,swp4\nswp5,swp5\n", encoding="utf-8",
    )
    lldpq = root / "selected-fabric-lldpq.dot"
    lldpq.write_text("graph fabric {\n" + "\n".join(DOT_LINES) + "\n}\n",
                     encoding="utf-8")
    splitter_profiles = root / "selected-fabric-splitter-profiles.json"
    document = {
        "schema_version": 1,
        "source_workbook": workbook.name,
        "workbook_sha256": _sha(workbook),
        "inventory_sha256": _sha(inventory),
        "port_mapping_sha256": _sha(port_mapping),
        "lldpq_sha256": _sha(lldpq),
        "profiles": [],
    }
    splitter_profiles.write_text(json.dumps(document, sort_keys=True) + "\n",
                                 encoding="utf-8")
    return dict(topology=workbook, inventory=inventory,
                port_mapping=port_mapping, lldpq=lldpq,
                splitter_profiles=splitter_profiles)


def _verify(paths: dict[str, Path], *, expected: dict[str, str] | None = None,
            selected_name: str = "selected-fabric.xlsx"):
    from monitor.issue_tracker_topology_derivation import verify_lldpq_derivation

    return verify_lldpq_derivation(
        sources=paths,
        expected_sha256=expected or {role: _sha(paths[role]) for role in ROLES},
        selected_workbook_name=selected_name,
        legacy_columns=False,
    )


def _revise_sidecar(paths: dict[str, Path], **changes: object) -> None:
    document = json.loads(paths["splitter_profiles"].read_text(encoding="utf-8"))
    document.update(changes)
    paths["splitter_profiles"].write_text(
        json.dumps(document, sort_keys=True) + "\n", encoding="utf-8",
    )


class TopologyDerivationDirectTests(unittest.TestCase):
    def test_literal_two_edges_keep_unique_source_rows_and_exclude_one_sided_intent(self):
        with tempfile.TemporaryDirectory() as directory:
            witness = _verify(literal_fixture(Path(directory)))
        self.assertEqual(EXPECTED, tuple(
            (edge.sheet, edge.row, edge.a_node, edge.a_port,
             edge.z_node, edge.z_port) for edge in witness.edges
        ))

    def test_self_consistent_sidecar_hashes_cannot_launder_changed_p2p_z(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            book = load_workbook(paths["topology"])
            book["TAN-Links"].cell(3, 3, "leaf-foreign")
            book.save(paths["topology"])
            book.close()
            _revise_sidecar(paths, workbook_sha256=_sha(paths["topology"]))
            with self.assertRaises(TopologyDerivationHoldError):
                _verify(paths)

    def test_changed_port_mapping_and_matching_sidecar_hash_cannot_keep_old_dot(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            mapping = paths["port_mapping"]
            mapping.write_text(mapping.read_text().replace("swp2,swp2", "swp2,swp99"))
            _revise_sidecar(paths, port_mapping_sha256=_sha(mapping))
            with self.assertRaises(TopologyDerivationHoldError):
                _verify(paths)

    def test_extra_missing_and_duplicate_dot_edges_hold_even_when_hash_is_reissued(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        for mutation in ("extra", "missing", "duplicate"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as directory:
                paths = literal_fixture(Path(directory))
                if mutation == "extra":
                    lines = (*DOT_LINES, '"leaf-x":"swp8" -- "leaf-y":"swp9"')
                elif mutation == "missing":
                    lines = DOT_LINES[:1]
                else:
                    lines = (*DOT_LINES, DOT_LINES[0])
                paths["lldpq"].write_text(
                    "graph fabric {\n" + "\n".join(lines) + "\n}\n",
                    encoding="utf-8",
                )
                _revise_sidecar(paths, lldpq_sha256=_sha(paths["lldpq"]))
                with self.assertRaises(TopologyDerivationHoldError):
                    _verify(paths)

    def test_same_bytes_under_foreign_workbook_basename_hold(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            paths["topology"] = paths["topology"].rename(Path(directory) / "foreign.xlsx")
            _revise_sidecar(paths, source_workbook="foreign.xlsx")
            with self.assertRaises(TopologyDerivationHoldError):
                _verify(paths)

    def test_duplicate_workbook_physical_edge_cannot_gain_second_row_witness(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            book = load_workbook(paths["topology"])
            sheet = book["TAN-Links"]
            for column, value in enumerate(EXPECTED[0][2:], start=1):
                sheet.cell(4, column, value)
            book.save(paths["topology"])
            book.close()
            paths["lldpq"].write_text(
                "graph fabric {\n" + "\n".join((DOT_LINES[0], DOT_LINES[0])) + "\n}\n",
                encoding="utf-8",
            )
            _revise_sidecar(paths, workbook_sha256=_sha(paths["topology"]),
                            lldpq_sha256=_sha(paths["lldpq"]))
            with self.assertRaises(TopologyDerivationHoldError):
                _verify(paths)

    def test_self_consistent_hashes_do_not_authorize_uninferred_breakout_profile(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            _revise_sidecar(paths, profiles=[{
                "device": "leaf-a", "parent": "swp1", "profile": "1to4",
            }])
            with self.assertRaises(TopologyDerivationHoldError):
                _verify(paths)

    def test_literal_partial_1to8_lane_derives_one_physical_edge_and_one_intent(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            book = load_workbook(paths["topology"])
            sheet = book["TAN-Links"]
            for row_number, fields in (
                (3, ("leaf-a", "19/1/1", "leaf-z", "swp2")),
                (4, ("leaf-a", "19/1/3", "Empty", "Empty")),
                (5, ("Empty", "Empty", "Empty", "Empty")),
            ):
                for column, value in enumerate(fields, 1):
                    sheet.cell(row_number, column, value)
            book.save(paths["topology"])
            book.close()
            paths["port_mapping"].write_text(
                "[Eth-SW]\n"
                "1to8#/1/1,swp#s0\n1to8#/1/3,swp#s2\n"
                "swp2,swp2\n", encoding="utf-8",
            )
            paths["lldpq"].write_text(
                'graph fabric {\n"leaf-a":"swp19s0" -- "leaf-z":"swp2"\n}\n',
                encoding="utf-8",
            )
            _revise_sidecar(
                paths, workbook_sha256=_sha(paths["topology"]),
                port_mapping_sha256=_sha(paths["port_mapping"]),
                lldpq_sha256=_sha(paths["lldpq"]),
                profiles=[{"device": "leaf-a", "parent": "swp19", "profile": "1to8"}],
            )
            witness = _verify(paths)
            self.assertEqual(
                (("TAN-Links", 3, "leaf-a", "19/1/1", "leaf-z", "swp2"),),
                tuple((edge.sheet, edge.row, edge.a_node, edge.a_port,
                       edge.z_node, edge.z_port) for edge in witness.edges),
            )

    def test_missing_mapping_uses_explicit_original_port_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            mapping = paths["port_mapping"]
            mapping.write_text(mapping.read_text(encoding="utf-8").replace(
                "swp2,swp2\n", "",
            ), encoding="utf-8")
            _revise_sidecar(paths, port_mapping_sha256=_sha(mapping))
            witness = _verify(paths)
            self.assertEqual(2, len(witness.edges))

    def test_frozen_expected_digests_reject_later_source_rewrite(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            bound = {role: _sha(paths[role]) for role in ROLES}
            content = paths["topology"].read_bytes()
            paths["topology"].write_bytes(content + b"later source bytes")
            with self.assertRaises(TopologyDerivationHoldError):
                _verify(paths, expected=bound)

    def test_duplicate_xlsx_member_holds_even_when_sidecar_hash_is_reissued(self):
        from monitor.issue_tracker_topology_derivation import TopologyDerivationHoldError

        with tempfile.TemporaryDirectory() as directory:
            paths = literal_fixture(Path(directory))
            with zipfile.ZipFile(paths["topology"], "a") as archive:
                sheet = archive.read("xl/worksheets/sheet1.xml")
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore", UserWarning)
                    archive.writestr("xl/worksheets/sheet1.xml", sheet)
            _revise_sidecar(paths, workbook_sha256=_sha(paths["topology"]))
            with self.assertRaises(TopologyDerivationHoldError):
                _verify(paths)


if __name__ == "__main__":
    unittest.main()
