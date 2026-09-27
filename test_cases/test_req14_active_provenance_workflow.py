#!/usr/bin/env python3
"""REQ-14 synthetic 0915-layout load -> P2P -> NVUE provenance workflow.

The four endpoint groups below are an independently specified transcription
fixture.  This is *not* the private 0915 workbook or evidence that the active
project has published its matching sidecar and LLDPQ DOT.
"""

from __future__ import annotations

import contextlib
import csv
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook
import yaml

from test_cases.test_mlag_evpn_generation import border_globals
from test_cases.test_qos_evpn_uplink_generation import ROOT
from test_cases.test_splitter_profile_breakout import device_for_ports
from test_cases.test_splitter_profile_breakout_workflow import GENERATOR, P2P_DIR
from test_cases.test_xlsx_zero_row_fail_closed import invoke_main, prepare_runtime
from test_cases.test_xlsx_zero_row_workflow import LOAD
from test_cases.module_loader import load_script


PUBLISHER = load_script(
    "req14_active_private_publisher",
    ROOT / "ztp/config/cumulus/d-hostname2mac.py",
)


# (sheet, first spreadsheet row, device, parent, number of physically linked
# lanes).  The lane notation is specified here, not read from a generated DOT.
SEED = (
    ("OOB LF-SP-CR-BL", 764, "EXAMPLE-OOB-CORE01", 9, 4),
    ("OOB LF-SP-CR-BL", 804, "EXAMPLE-OOB-CORE02", 9, 4),
    ("TAN OBJ-LF", 75, "EXAMPLE-TAN-OBJ-LEAF01", 19, 1),
    ("TAN OBJ-LF", 155, "EXAMPLE-TAN-OBJ-LEAF02", 19, 1),
)
EXPECTED_PROFILES = {
    ("example-oob-core01", "swp9", "1to8"),
    ("example-oob-core02", "swp9", "1to8"),
    ("example-tan-obj-leaf01", "swp19", "1to8"),
    ("example-tan-obj-leaf02", "swp19", "1to8"),
}
EXPECTED_KEYS = {
    "schema_version", "source_workbook", "workbook_sha256",
    "inventory_sha256", "port_mapping_sha256", "lldpq_sha256", "profiles",
}


class _ReachedGenerator(RuntimeError):
    """Stop before 11-load's publisher, DHCP install or service operations."""


def _write_fixture_workbook(path: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    for sheet_name, first_row, hostname, parent, physical_count in SEED:
        if sheet_name not in book.sheetnames:
            sheet = book.create_sheet(sheet_name)
            if sheet_name.startswith("OOB"):
                # Deliberately lacks detectable endpoint headers: the real
                # private 0915 OOB sheets need explicit legacy-column opt-in.
                sheet.cell(2, 18, "legacy diagram")
            else:
                sheet.cell(1, 5, "Source")
                sheet.cell(1, 11, "Dest")
                for column, label in ((7, "name"), (8, "HCA/port"),
                                      (13, "name"), (14, "port")):
                    sheet.cell(2, column, label)
                sheet.cell(2, 23)
        sheet = book[sheet_name]
        for lane in range(8):
            notation = f"{parent}/{1 if lane < 4 else 2}/{lane % 4 + 1}"
            if lane < physical_count:
                peer = f"EXAMPLE-PEER-{parent}-{hostname[-2:]}-{lane + 1}"
                peer_port = f"P{lane + 1}"
            else:
                peer = peer_port = "empty"
            for column, value in zip((7, 8, 13, 14),
                                     (peer, peer_port, hostname, notation)):
                sheet.cell(first_row + lane, column, value)
    book.save(path)
    book.close()


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Req14ActiveProvenanceWorkflow(unittest.TestCase):
    def test_real_environment_card_distinguishes_private_positive_from_acceptance(self):
        card = (ROOT / "test_cases/REAL_ENVIRONMENT.md").read_text(encoding="utf-8")
        section = card.split(
            "## TC-REAL-REQ14-ACTIVE-PROVENANCE-001", 1,
        )[1].split("\n## ", 1)[0]
        self.assertIn("PARTIAL / PRIVATE REAL-INPUT POSITIVE + NINE NEGATIVE PASS", section)
        self.assertIn("SOURCE-YAML FULL-PARENT PRIVATE PASS; LIVE RUNTIME OPEN", section)
        self.assertIn("bfc4144da8a6b049fc5c2e501acfeb6bcb2400ac21658ad1f3394f2717fb06a1", section)
        self.assertIn("9336a3bc779708ad11504784986f0b50b07b60747c8f9247995d0a22ac65386a", section)
        self.assertIn("72985e000ab0a6f2862f5d84ed9486d1752e15e1efa7eef7408711f04377ccd9", section)
        self.assertIn("不证明 AIR/prod", section)

    def _fixture(self, root: Path):
        paths = prepare_runtime(root, _write_fixture_workbook)
        old = root / "rack-links.xlsx"
        workbook = root / "fixture-0915-transcription.xlsx"
        old.rename(workbook)
        link = paths["p2p_dir"] / "p2p.xlsx"
        link.unlink()
        link.symlink_to(workbook)
        for key, filename in (("inventory", "01-inventory.log"),
                              ("port_map", "02-port-mapping.log")):
            paths[key].write_bytes((P2P_DIR / filename).read_bytes())
        policy = root / "03-air-topology-policy.json"
        policy.write_text("{}\n", encoding="utf-8")
        paths.update(workbook=workbook, policy=policy)
        return paths

    def _devices(self, *, bare_child: bool = False):
        return {
            (name.removeprefix("EXAMPLE-") if bare_child else name):
            device_for_ports(
                name.removeprefix("EXAMPLE-") if bare_child else name,
                [f"swp{parent}s0"], {},
                role="bond" if parent == 19 else "bgp",
            )
            for _, _, name, parent, _ in SEED
        }

    def _run_three_scripts(self, root: Path, paths, *, legacy: bool,
                           output: Path | None = None,
                           bare_child: bool = False):
        commands = []
        output = output or root / "generated"
        devices = self._devices(bare_child=bare_child)

        def dispatch(command, *, cwd, **_kwargs):
            commands.append(list(command))
            script = command[1]
            if script == "b-xlsx_to_dot.py":
                self.assertEqual(paths["p2p_dir"], Path(cwd))
                caught, stdout, stderr = invoke_main(paths, list(command[2:]))
                if caught is not None and not (isinstance(caught, SystemExit)
                                               and caught.code == 0):
                    if not legacy:
                        self.assertIn("表头自动检测失败", stdout + stderr)
                        self.assertIn("OOB LF-SP-CR-BL", stdout + stderr)
                    raise caught
                return None
            if script == "c1-generate_dhcp.py":
                # DHCP is a separate real script, not part of this three-script
                # REQ-14 proof; never install or publish it here.
                return None
            self.assertEqual("90-c2-generate_configs.py", script)
            self.assertEqual(root / "ztp/config/cumulus/template", Path(cwd))
            with mock.patch.multiple(
                GENERATOR, P2P_INPUT_DIR=str(paths["p2p_dir"]),
                P2P_OUTPUT_DIR=str(paths["output"]), OUTPUT_DIR=str(output),
            ), mock.patch.object(GENERATOR, "load_devices", return_value=(
                border_globals(), devices,
            )), contextlib.redirect_stdout(io.StringIO()):
                GENERATOR.generate_all()
            raise _ReachedGenerator("real producer and real generator completed")

        with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), mock.patch.object(
            LOAD, "run", side_effect=dispatch,
        ), mock.patch.object(LOAD, "_device_types_after_dhcp",
                             return_value=frozenset({"eth"})):
            with self.assertRaises(_ReachedGenerator if legacy else SystemExit):
                LOAD.generate_configs(
                    frozenset({"eth"}), install_dhcp=False,
                    deployment_scope="prod", p2p_legacy_columns=legacy,
                    air_topology_policy=paths["policy"],
                )
        return commands, output, devices

    def _private_parent_inputs(self, root: Path, paths, *, bare_child: bool = False):
        """Keep every publisher/parent input and output inside one temp tree."""
        project = root / "DAY0-Prepare/req14-synthetic"
        project.mkdir(parents=True)
        workbook = project / paths["workbook"].name
        paths["workbook"].rename(workbook)
        paths["workbook"] = workbook
        (paths["p2p_dir"] / "p2p.xlsx").unlink()
        (paths["p2p_dir"] / "p2p.xlsx").symlink_to(workbook)
        (project / "p2p.xlsx").symlink_to(workbook.name)

        global_file = project / "01-global.yaml"
        global_file.write_text(
            "schema_version: 1\ncommon: {}\nswitches:\n"
            "- eth:\n    version: 5.18.1\n"
            "    system:\n      date-time:\n        timezone: Etc/UTC\n",
            encoding="utf-8",
        )
        service = root / "ztp/config/cumulus"
        template = service / "template"
        (template / "01-global.yaml").symlink_to(global_file)
        (service / "default_5.18.1.yaml").write_text(
            "- set:\n    system:\n      date-time:\n"
            "        timezone: Etc/UTC\n", encoding="utf-8",
        )
        devices_file = project / "02-devices_config.csv"
        with devices_file.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerow((
                "hostname", "type", "template", "eth0_ip", "netmask",
                "eth0_gw", "eth0_mac", "eth1_ip", "netmask", "eth1_gw",
                "eth1_mac",
            ))
            for index, (_, _, hostname, _, _) in enumerate(SEED, 1):
                if bare_child:
                    hostname = hostname.removeprefix("EXAMPLE-")
                writer.writerow((
                    hostname, "eth", "leaf", "192.0.2.10", "24",
                    "192.0.2.1", f"02:00:00:00:14:{index:02x}", "", "", "", "",
                ))
        paths["devices"].unlink()
        paths["devices"].symlink_to(devices_file)
        subnet_file = project / "03-subnet.csv"
        subnet_file.write_text(
            "shared_network,subnet\nnet,192.0.2.0\n", encoding="utf-8",
        )
        published_root = project / "99-output-eth"
        published_root.mkdir()
        (template / "99-output").symlink_to(published_root, target_is_directory=True)
        settings = LOAD.GlobalSettings(
            dhcp_enabled=False, dhcp_package="isc-dhcp-server",
            http_enabled=True, http_package="apache2", http_root=root,
            ztp_enabled=True, ztp_prefix="/ztp", ztp_ips={},
            versions={"cumulus": "5.18.1"},
        )
        inputs = LOAD.ProjectInputs(
            global_file=global_file, devices_file=devices_file,
            subnet_file=subnet_file, p2p_file=workbook,
            device_types=frozenset({"eth"}), pubkeys=(), settings=settings,
            deployment_scope="prod", switch_scope="eth",
        )
        return project, template, devices_file, inputs

    def _assert_published_breakouts(self, release: Path, *, bare_child: bool = False):
        for _, _, hostname, parent, _ in SEED:
            if bare_child:
                hostname = hostname.removeprefix("EXAMPLE-")
            rendered = yaml.safe_load((release / f"{hostname}.yaml").read_text())
            interfaces = rendered[0]["set"]["interface"]
            self.assertEqual(
                {"8x": {"lanes-per-port": "1"}},
                interfaces[f"swp{parent}"]["link"]["breakout"],
                hostname,
            )
            self.assertEqual({"type": "swp"}, interfaces[f"swp{parent}s7"])

    def test_site_prefixed_sidecar_binds_bare_child_and_rejects_ambiguity(self):
        """The real producer's site key must bind one matching child hostname."""
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            paths = self._fixture(root)
            project, template, devices_file, inputs = self._private_parent_inputs(
                root, paths, bare_child=True,
            )
            production = template / "99-output/20260915_120000"
            self._run_three_scripts(
                root, paths, legacy=True, output=production, bare_child=True,
            )
            context = PUBLISHER._cumulus_dir_context(str(production))
            with mock.patch.object(PUBLISHER, "_AUTO_YES", True), \
                    contextlib.redirect_stdout(io.StringIO()):
                self.assertTrue(PUBLISHER._publish_production_cumulus(
                    context, PUBLISHER.load_csv(devices_file), switch_scope="eth",
                ))
            release = (root / "ztp/config/cumulus/latest_yaml").resolve(strict=True)
            self._assert_published_breakouts(release, bare_child=True)
            child = json.loads((release / "release-manifest.json").read_text())
            self.assertEqual(
                {name.removeprefix("EXAMPLE-") for _, _, name, _, _ in SEED},
                {item["hostname"] for item in child["devices"]},
            )
            sidecar = paths["output"] / (
                f"{paths['workbook'].stem}-splitter-profiles.json"
            )
            sidecar_bytes = sidecar.read_bytes()
            with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), \
                    mock.patch.object(LOAD, "run", side_effect=AssertionError(
                        "parent receipt must not dispatch services or subprocesses"
                    )):
                parent = LOAD.validate_and_publish_release(project, inputs)
                self.assertEqual(
                    _sha(sidecar),
                    parent["components"]["cumulus"]["splitter_provenance"]
                    ["splitter_sidecar_sha256"],
                )
                receipt = project / "99-output-ztp/current-release.json"
                last_good = receipt.read_bytes()
                document = json.loads(sidecar_bytes)
                duplicate = dict(document["profiles"][0])
                duplicate["device"] = "OTHER-OOB-CORE01"
                document["profiles"].append(duplicate)
                sidecar.write_text(json.dumps(document), encoding="utf-8")
                try:
                    with self.assertRaisesRegex(LOAD.LoadError, "ambiguous"):
                        LOAD.validate_and_publish_release(project, inputs)
                    self.assertEqual(last_good, receipt.read_bytes())
                finally:
                    sidecar.write_bytes(sidecar_bytes)

    def test_synthetic_sidecar_reaches_real_private_publisher_and_parent(self):
        """Real child/parent release binding, not active 0915 site acceptance."""
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            paths = self._fixture(root)
            project, template, devices_file, inputs = self._private_parent_inputs(
                root, paths,
            )
            timestamp = "20260915_120000"
            production = template / "99-output" / timestamp
            commands, output, _ = self._run_three_scripts(
                root, paths, legacy=True, output=production,
            )
            self.assertEqual(3, len(commands))
            sidecar = paths["output"] / (
                f"{paths['workbook'].stem}-splitter-profiles.json"
            )
            document = json.loads(sidecar.read_text(encoding="utf-8"))
            self.assertEqual(EXPECTED_KEYS, set(document))
            self.assertEqual(EXPECTED_PROFILES, {
                (entry["device"], entry["parent"], entry["profile"])
                for entry in document["profiles"]
            })
            for field, source in (
                ("workbook_sha256", paths["workbook"]),
                ("inventory_sha256", paths["inventory"]),
                ("port_mapping_sha256", paths["port_map"]),
                ("lldpq_sha256", paths["output"] /
                 f"{paths['workbook'].stem}-lldpq.dot"),
            ):
                self.assertEqual(_sha(source), document[field], field)
            self._assert_published_breakouts(output)

            context = PUBLISHER._cumulus_dir_context(str(production))
            self.assertIsNotNone(context)
            self.assertIsNone(PUBLISHER._day0_project_for_cumulus_publish(
                context["combine_dir"],
            ))
            self.assertTrue(Path(context["combine_dir"]).resolve().is_relative_to(
                root.resolve(),
            ))
            with mock.patch.object(PUBLISHER, "_AUTO_YES", True), \
                    contextlib.redirect_stdout(io.StringIO()):
                self.assertTrue(PUBLISHER._publish_production_cumulus(
                    context, PUBLISHER.load_csv(devices_file), switch_scope="eth",
                ))
            release = (root / "ztp/config/cumulus/latest_yaml").resolve(strict=True)
            self.assertTrue(release.is_relative_to(root.resolve()))
            self._assert_published_breakouts(release)
            child = json.loads((release / "release-manifest.json").read_text())
            self.assertEqual("prod", child["deployment_scope"])
            self.assertEqual("eth", child["switch_scope"])
            self.assertEqual({name for _, _, name, _, _ in SEED}, {
                item["hostname"] for item in child["devices"]
            })
            for item in child["devices"]:
                self.assertEqual(_sha(release / item["config"]),
                                 item["config_sha256"])

            with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), \
                    mock.patch.object(LOAD, "run", side_effect=AssertionError(
                        "parent receipt must not dispatch services or subprocesses"
                    )):
                parent = LOAD.validate_and_publish_release(project, inputs)
            receipt = project / "99-output-ztp/current-release.json"
            self.assertEqual(parent, json.loads(receipt.read_text(encoding="utf-8")))
            self.assertEqual("disabled", parent["dhcp_status"])
            self.assertEqual(_sha(paths["workbook"]), parent["inputs"]["p2p"])
            self.assertEqual(2, parent["schema_version"])
            self.assertEqual({
                "path": paths["workbook"].name,
                "sha256": _sha(paths["workbook"]),
            }, parent["input_sources"]["p2p"])
            self.assertEqual(_sha(release / "release-manifest.json"),
                             parent["components"]["cumulus"]["manifest_sha256"])

            # An equally valid XLSX with identical bytes is still a different
            # setup-selected source.  A parent receipt must not claim the old
            # workbook after the project's relative selector is retargeted.
            alternate = project / "alternate-source.xlsx"
            alternate.write_bytes(paths["workbook"].read_bytes())
            selector = project / "p2p.xlsx"
            committed = receipt.read_bytes()
            selector.unlink()
            selector.symlink_to(alternate.name)
            try:
                with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), \
                        mock.patch.object(LOAD, "run", side_effect=AssertionError(
                            "source mismatch must fail before subprocesses"
                        )), self.assertRaises(LOAD.LoadError):
                    LOAD.validate_and_publish_release(project, inputs)
                self.assertEqual(committed, receipt.read_bytes())
            finally:
                selector.unlink()
                selector.symlink_to(paths["workbook"].name)

            victim = release / f"{SEED[0][2]}.yaml"
            original = victim.read_bytes()
            lossy = yaml.safe_load(original)
            lossy[0]["set"]["interface"]["swp9"]["link"]["breakout"] = {
                "4x": {"lanes-per-port": "1"},
            }
            victim.write_text(yaml.safe_dump(lossy), encoding="utf-8")
            try:
                with self.assertRaises(AssertionError):
                    self._assert_published_breakouts(release)
                with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), \
                        self.assertRaises(LOAD.LoadError):
                    LOAD.validate_and_publish_release(project, inputs)
                self.assertEqual(parent, json.loads(receipt.read_text(encoding="utf-8")))
            finally:
                victim.write_bytes(original)

    def test_parent_rejects_post_generation_sidecar_and_dot_replacement(self):
        """The parent cannot publish 8x YAML against changed source authority."""
        for mutation in ("profile-1to4", "missing-lldpq", "changed-lldpq"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as name:
                root = Path(name)
                paths = self._fixture(root)
                project, template, devices_file, inputs = self._private_parent_inputs(
                    root, paths,
                )
                production = template / "99-output/20260915_120000"
                self._run_three_scripts(root, paths, legacy=True, output=production)
                sidecar = paths["output"] / (
                    f"{paths['workbook'].stem}-splitter-profiles.json"
                )
                lldpq = paths["output"] / (
                    f"{paths['workbook'].stem}-lldpq.dot"
                )
                context = PUBLISHER._cumulus_dir_context(str(production))
                with mock.patch.object(PUBLISHER, "_AUTO_YES", True), \
                        contextlib.redirect_stdout(io.StringIO()):
                    self.assertTrue(PUBLISHER._publish_production_cumulus(
                        context, PUBLISHER.load_csv(devices_file),
                        switch_scope="eth",
                    ))
                release = (root / "ztp/config/cumulus/latest_yaml").resolve(strict=True)
                self._assert_published_breakouts(release)
                with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"):
                    LOAD.validate_and_publish_release(project, inputs)
                receipt = project / "99-output-ztp/current-release.json"
                last_good = receipt.read_bytes()
                child_hashes = {
                    path.name: _sha(path) for path in release.glob("*.yaml")
                }

                if mutation == "profile-1to4":
                    document = json.loads(sidecar.read_text(encoding="utf-8"))
                    self.assertEqual("1to8", document["profiles"][0]["profile"])
                    document["profiles"][0]["profile"] = "1to4"
                    replacement = sidecar.with_name(".replacement-profiles.json")
                    replacement.write_text(json.dumps(document, sort_keys=True) + "\n",
                                           encoding="utf-8")
                    replacement.replace(sidecar)
                    self.assertEqual(1, sidecar.stat().st_nlink)
                elif mutation == "missing-lldpq":
                    lldpq.unlink()
                else:
                    with lldpq.open("ab") as stream:
                        stream.write(b"\nchanged-after-generation")

                with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), \
                        mock.patch.object(LOAD, "run", side_effect=AssertionError(
                            "parent must not dispatch subprocesses"
                        )), self.assertRaises(LOAD.LoadError):
                    LOAD.validate_and_publish_release(project, inputs)
                self.assertEqual(last_good, receipt.read_bytes())
                self.assertEqual(child_hashes, {
                    path.name: _sha(path) for path in release.glob("*.yaml")
                })

    def test_parent_rechecks_splitter_authority_at_atomic_commit(self):
        """Validation cannot bless a sidecar or LLDPQ changed before commit."""
        for mutation in ("sidecar", "lldpq"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as name:
                root = Path(name)
                paths = self._fixture(root)
                project, template, devices_file, inputs = self._private_parent_inputs(
                    root, paths,
                )
                production = template / "99-output/20260915_120000"
                self._run_three_scripts(root, paths, legacy=True, output=production)
                context = PUBLISHER._cumulus_dir_context(str(production))
                with mock.patch.object(PUBLISHER, "_AUTO_YES", True), \
                        contextlib.redirect_stdout(io.StringIO()):
                    self.assertTrue(PUBLISHER._publish_production_cumulus(
                        context, PUBLISHER.load_csv(devices_file),
                        switch_scope="eth",
                    ))
                with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"):
                    parent = LOAD.validate_and_publish_release(
                        project, inputs, publish=False,
                    )
                    self.assertIn("splitter_provenance", parent["components"]["cumulus"])
                    candidate = LOAD.prepare_current_release(project, parent)
                    try:
                        stem = paths["workbook"].stem
                        target = paths["output"] / (
                            f"{stem}-splitter-profiles.json" if mutation == "sidecar"
                            else f"{stem}-lldpq.dot"
                        )
                        with target.open("ab") as stream:
                            stream.write(b"\nchanged-after-parent-prepare")
                        with self.assertRaises(LOAD.LoadError):
                            LOAD.commit_prepared_release(candidate)
                        self.assertFalse(candidate.committed)
                        self.assertFalse(candidate.destination.exists())
                    finally:
                        LOAD.discard_prepared_release(candidate)

    def test_explicit_legacy_flag_reaches_bound_sidecar_and_four_real_renderers(self):
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            paths = self._fixture(root)
            commands, output, _ = self._run_three_scripts(root, paths, legacy=True)
            self.assertEqual(1, commands[0].count("--legacy-columns"))
            self.assertEqual(["b-xlsx_to_dot.py", "c1-generate_dhcp.py",
                              "90-c2-generate_configs.py"],
                             [command[1] for command in commands])

            stem = paths["workbook"].stem
            sidecar = paths["output"] / f"{stem}-splitter-profiles.json"
            lldpq = paths["output"] / f"{stem}-lldpq.dot"
            self.assertTrue(sidecar.is_file())
            self.assertTrue(lldpq.is_file())
            self.assertEqual(1, sidecar.stat().st_nlink)
            document = json.loads(sidecar.read_text(encoding="utf-8"))
            self.assertEqual(EXPECTED_KEYS, set(document))
            self.assertEqual(1, document["schema_version"])
            self.assertEqual(paths["workbook"].name, document["source_workbook"])
            for field, path in (
                ("workbook_sha256", paths["workbook"]),
                ("inventory_sha256", paths["inventory"]),
                ("port_mapping_sha256", paths["port_map"]),
                ("lldpq_sha256", lldpq),
            ):
                self.assertEqual(_sha(path), document[field], field)
            self.assertEqual(EXPECTED_PROFILES, {
                (entry["device"], entry["parent"], entry["profile"])
                for entry in document["profiles"]
            })
            for _, _, hostname, parent, _ in SEED:
                rendered = yaml.safe_load((output / f"{hostname}.yaml").read_text())
                interfaces = rendered[0]["set"]["interface"]
                self.assertEqual({"8x": {"lanes-per-port": "1"}},
                    interfaces[f"swp{parent}"]["link"]["breakout"])
                self.assertEqual({"type": "swp"}, interfaces[f"swp{parent}s7"])

    def test_default_without_legacy_flag_fails_before_sidecar_and_renderer(self):
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            paths = self._fixture(root)
            commands, output, _ = self._run_three_scripts(root, paths, legacy=False)
            self.assertEqual(0, commands[0].count("--legacy-columns"))
            self.assertEqual(["b-xlsx_to_dot.py"], [command[1] for command in commands])
            self.assertFalse((paths["output"] /
                f"{paths['workbook'].stem}-splitter-profiles.json").exists())
            self.assertFalse(output.exists())

    def test_missing_wrong_stem_and_stale_bound_inputs_fail_before_new_yaml(self):
        for mutation in ("missing", "wrong-stem", "source-workbook", "workbook",
                         "inventory", "port-map", "missing-lldpq", "stale-lldpq"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as name:
                root = Path(name)
                paths = self._fixture(root)
                _, output, devices = self._run_three_scripts(root, paths, legacy=True)
                stem = paths["workbook"].stem
                sidecar = paths["output"] / f"{stem}-splitter-profiles.json"
                lldpq = paths["output"] / f"{stem}-lldpq.dot"
                before = {path.name: path.read_bytes() for path in output.glob("*.yaml")}
                if mutation == "missing":
                    sidecar.unlink()
                elif mutation == "wrong-stem":
                    sidecar.rename(paths["output"] / "newer-splitter-profiles.json")
                elif mutation == "source-workbook":
                    document = json.loads(sidecar.read_text(encoding="utf-8"))
                    document["source_workbook"] = "another-source.xlsx"
                    sidecar.write_text(json.dumps(document), encoding="utf-8")
                elif mutation in {"workbook", "inventory", "port-map"}:
                    path = {"workbook": paths["workbook"],
                            "inventory": paths["inventory"],
                            "port-map": paths["port_map"]}[mutation]
                    with path.open("ab") as stream:
                        stream.write(b"\nchanged-after-publication")
                elif mutation == "missing-lldpq":
                    lldpq.unlink()
                else:
                    with lldpq.open("ab") as stream:
                        stream.write(b"\nchanged-after-publication")

                new_output = root / "should-not-be-published"
                with mock.patch.multiple(
                    GENERATOR, P2P_INPUT_DIR=str(paths["p2p_dir"]),
                    P2P_OUTPUT_DIR=str(paths["output"]), OUTPUT_DIR=str(new_output),
                ), mock.patch.object(GENERATOR, "load_devices", return_value=(
                    border_globals(), devices,
                )), self.assertRaisesRegex(ValueError, "splitter profiles unavailable/invalid"):
                    GENERATOR.generate_all()
                self.assertFalse(new_output.exists())
                self.assertEqual(before, {
                    path.name: path.read_bytes() for path in output.glob("*.yaml")
                })


if __name__ == "__main__":
    unittest.main()
