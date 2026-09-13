#!/usr/bin/env python3
"""Direct fail-closed contracts for the P2P workbook extractor entrypoint."""

from __future__ import annotations

from contextlib import nullcontext, redirect_stderr, redirect_stdout
import importlib.util
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "ztp/config/cumulus/template/P2P/b-xlsx_to_dot.py"


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    previous = sys.modules.get(name)
    sys.modules[name] = module
    try:
        spec.loader.exec_module(module)
    finally:
        if previous is None:
            sys.modules.pop(name, None)
        else:
            sys.modules[name] = previous
    return module


P2P = load_module("xlsx_zero_row_direct", SCRIPT)


class LegacyReached(RuntimeError):
    """Sentinel proving extraction advanced into topology resolution."""


def _detected_headers(sheet) -> None:
    sheet.append(["Source", "", "Dest", ""])
    sheet.append(["name", "port", "name", "port"])


def make_zero_row_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "TAN-Links"
    _detected_headers(first)
    second = workbook.create_sheet("OOB-Links")
    _detected_headers(second)
    workbook.save(path)
    workbook.close()


def make_legacy_column_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "TAN-Legacy"
    sheet.append(["unrecognised section labels"])
    sheet.append(["unrecognised column labels"])
    row = [None] * 14
    row[6], row[7] = "leaf01", "swp1"
    row[12], row[13] = "leaf02", "swp1"
    sheet.append(row)
    workbook.save(path)
    workbook.close()


def prepare_runtime(root: Path, workbook_writer) -> dict[str, Path]:
    p2p_dir = root / "ztp/config/cumulus/template/P2P"
    p2p_dir.mkdir(parents=True)
    source = root / "rack-links.xlsx"
    workbook_writer(source)
    (p2p_dir / "p2p.xlsx").symlink_to(source)
    inventory = p2p_dir / "01-inventory.log"
    inventory.write_text("[Eth-SW]\n*leaf*\n", encoding="utf-8")
    port_map = p2p_dir / "02-port-mapping.log"
    port_map.write_text("[Eth-SW]\nswp1,swp1\n", encoding="utf-8")
    devices = p2p_dir.parent / "02-devices_config.csv"
    devices.write_text(
        "hostname,type,template,eth0_ip,netmask,eth0_gw,eth0_mac,"
        "eth1_ip,netmask,eth1_gw,eth1_mac\n",
        encoding="utf-8",
    )
    air_template = p2p_dir / "air-template-no-oob.json"
    air_template.write_text('{"content": {"nodes": {}, "links": {}}}\n', encoding="utf-8")
    output = p2p_dir / "output-p2p"
    output.mkdir()
    return {
        "p2p_dir": p2p_dir,
        "inventory": inventory,
        "port_map": port_map,
        "devices": devices,
        "air_template": air_template,
        "output": output,
    }


def stale_artifacts(paths: dict[str, Path]) -> tuple[Path, Path, list[Path]]:
    output = paths["output"]
    exact_dot = output / "rack-links-lldpq.dot"
    exact_intent = output / "rack-links-description-intent.json"
    unrelated = [
        output / "other-lldpq.dot",
        output / "rack-links-air.dot",
        output / "rack-links-air.json",
        output / "other-description-intent.json",
    ]
    for path in (exact_dot, exact_intent, *unrelated):
        path.write_text(f"stale:{path.name}\n", encoding="utf-8")
    return exact_dot, exact_intent, unrelated


def invoke_main(
    paths: dict[str, Path], args: list[str], *,
    load_inventory=None, extractor=None,
):
    stdout = io.StringIO()
    stderr = io.StringIO()
    inventory_behavior = load_inventory or P2P.load_inventory
    patches = mock.patch.multiple(
        P2P,
        SCRIPT_DIR=str(paths["p2p_dir"]),
        DEFAULT_INV=str(paths["inventory"]),
        DEFAULT_PORT_MAP=str(paths["port_map"]),
        DEVICES_CONFIG=str(paths["devices"]),
        AIR_JSON_TEMPLATE=str(paths["air_template"]),
        HTTP_ROOT=paths["p2p_dir"].parents[4],
        deployment_lock=lambda _root: nullcontext(),
        stop_native_ztp_monitors=mock.Mock(),
        load_inventory=inventory_behavior,
        _extract_xlsx_rows=extractor or P2P._extract_xlsx_rows,
    )
    caught = None
    with patches, mock.patch.object(sys, "argv", ["b-xlsx_to_dot.py", *args]), \
            redirect_stdout(stdout), redirect_stderr(stderr):
        try:
            P2P.main()
        except BaseException as exc:  # preserve SystemExit and the sentinel
            caught = exc
    return caught, stdout.getvalue(), stderr.getvalue()


class XlsxFailClosedDirectTests(unittest.TestCase):
    def test_exact_stale_pair_is_absent_before_workbook_extraction_begins(self):
        with tempfile.TemporaryDirectory() as name:
            paths = prepare_runtime(Path(name), make_zero_row_workbook)
            exact_dot, exact_intent, unrelated = stale_artifacts(paths)

            def observe_cleanup(_xlsx_path):
                self.assertFalse(exact_dot.exists())
                self.assertFalse(exact_intent.exists())
                self.assertTrue(all(path.is_file() for path in unrelated), unrelated)
                raise SystemExit(17)

            caught, _stdout, _stderr = invoke_main(
                paths, ["-y"], extractor=observe_cleanup,
            )
            self.assertIsInstance(caught, SystemExit)
            self.assertEqual(17, caught.code)

    def test_zero_rows_exit_one_name_all_scanned_sheets_and_remove_only_exact_stale_pair(self):
        with tempfile.TemporaryDirectory() as name:
            paths = prepare_runtime(Path(name), make_zero_row_workbook)
            exact_dot, exact_intent, unrelated = stale_artifacts(paths)
            caught, stdout, stderr = invoke_main(paths, ["-y"])

            self.assertIsInstance(caught, SystemExit)
            self.assertEqual(1, caught.code)
            diagnostic = stdout + stderr
            self.assertIn("TAN-Links", diagnostic)
            self.assertIn("OOB-Links", diagnostic)
            self.assertIn("0", diagnostic)
            self.assertFalse(exact_dot.exists())
            self.assertFalse(exact_intent.exists())
            self.assertTrue(all(path.is_file() for path in unrelated), unrelated)

    def test_header_auto_detection_failure_is_fatal_before_topology_resolution(self):
        with tempfile.TemporaryDirectory() as name:
            paths = prepare_runtime(Path(name), make_legacy_column_workbook)
            exact_dot, exact_intent, unrelated = stale_artifacts(paths)
            resolver = mock.Mock(side_effect=LegacyReached("must not resolve"))
            caught, stdout, stderr = invoke_main(
                paths, ["-y"], load_inventory=resolver,
            )

            self.assertIsInstance(caught, SystemExit)
            self.assertEqual(1, caught.code)
            resolver.assert_not_called()
            self.assertIn("TAN-Legacy", stdout + stderr)
            self.assertRegex(stdout + stderr, r"(?i)header|表头")
            self.assertFalse(exact_dot.exists())
            self.assertFalse(exact_intent.exists())
            self.assertTrue(all(path.is_file() for path in unrelated), unrelated)

    def test_explicit_legacy_columns_allows_only_the_legacy_fallback_extraction(self):
        with tempfile.TemporaryDirectory() as name:
            paths = prepare_runtime(Path(name), make_legacy_column_workbook)
            exact_dot, exact_intent, unrelated = stale_artifacts(paths)
            resolver = mock.Mock(side_effect=LegacyReached("legacy extracted"))
            caught, _stdout, _stderr = invoke_main(
                paths, ["-y", "--legacy-columns"], load_inventory=resolver,
            )

            self.assertIsInstance(caught, LegacyReached)
            resolver.assert_called_once_with(str(paths["inventory"]))
            self.assertFalse(exact_dot.exists())
            self.assertFalse(exact_intent.exists())
            self.assertTrue(all(path.is_file() for path in unrelated), unrelated)

    def test_legacy_columns_must_not_be_repeated(self):
        with tempfile.TemporaryDirectory() as name:
            paths = prepare_runtime(Path(name), make_legacy_column_workbook)
            caught, stdout, stderr = invoke_main(
                paths, ["-y", "--legacy-columns", "--legacy-columns"],
            )
            self.assertIsInstance(caught, SystemExit)
            self.assertEqual(1, caught.code)
            self.assertRegex(stdout + stderr, r"(?i)duplicate.*legacy-columns")


if __name__ == "__main__":
    unittest.main()
