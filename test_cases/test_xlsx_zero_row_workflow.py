#!/usr/bin/env python3
"""Workflow contract: DAY0 generation stops at a malformed P2P workbook."""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock


ROOT = Path(__file__).resolve().parents[1]


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    previous = sys.modules.get(name)
    sys.modules[name] = module
    try:
        spec.loader.exec_module(module)
    finally:
        if previous is None:
            sys.modules.pop(name, None)
        else:
            sys.modules[name] = previous
    return module


DIRECT = load_module(
    "xlsx_zero_row_direct_helpers",
    ROOT / "test_cases/test_xlsx_zero_row_fail_closed.py",
)
LOAD = load_module("xlsx_zero_row_day0_load", ROOT / "DAY0-Prepare/11-load.py")


class XlsxFailClosedWorkflowTests(unittest.TestCase):
    def test_explicit_load_legacy_flag_reaches_real_producer_once_for_mixed_workbook(self):
        import openpyxl

        def write_mixed_workbook(path):
            book = openpyxl.Workbook()
            book.active.title = "OOB Fabric"
            book.active.cell(2, 18)  # no endpoint header: a drawing sheet
            plan = book.create_sheet("OOB Plan")
            plan.cell(1, 2, "Core POD")
            plan.cell(2, 1, "sum")
            plan.cell(2, 4, "HPS(Weka BMC)")
            link = book.create_sheet("TAN OBJ-LF")
            link.cell(1, 5, "Source")
            link.cell(1, 11, "Dest")
            for col, label in ((7, "name"), (8, "HCA/port"), (13, "name"), (14, "port")):
                link.cell(2, col, label)
            link.cell(2, 23)
            for col, value in zip((7, 8, 13, 14), ("leaf01", "swp1", "leaf02", "swp1")):
                link.cell(3, col, value)
            book.save(path)
            book.close()

        for flags, succeeds in (([], False), (["--p2p-legacy-columns"] * 2, True)):
            with self.subTest(flags=flags), tempfile.TemporaryDirectory() as name:
                root = Path(name)
                paths = DIRECT.prepare_runtime(root, write_mixed_workbook)
                policy = root / "03-air-topology-policy.json"
                policy.write_text('{}\n')
                args = LOAD.parse_args(["example-project", *flags])
                commands = []

                def run_producer(command, *, cwd, **_kwargs):
                    commands.append(list(command))
                    if command[1] != "b-xlsx_to_dot.py":
                        self.assertEqual("c1-generate_dhcp.py", command[1])
                        raise DIRECT.LegacyReached("DHCP reached after real P2P")
                    self.assertEqual(paths["p2p_dir"], Path(cwd))
                    caught, stdout, stderr = DIRECT.invoke_main(paths, list(command[2:]))
                    if caught is not None:
                        if not succeeds:
                            self.assertIn("sheet OOB Fabric: 表头自动检测失败", stderr)
                        raise caught

                with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), mock.patch.object(
                    LOAD, "run", side_effect=run_producer,
                ), self.assertRaises(DIRECT.LegacyReached if succeeds else SystemExit) as result:
                    LOAD.generate_configs(
                        frozenset({"eth"}), install_dhcp=False, deployment_scope="prod",
                        p2p_legacy_columns=args.p2p_legacy_columns, air_topology_policy=policy,
                    )
                self.assertEqual(1 if succeeds else 0, commands[0].count("--legacy-columns"))
                self.assertEqual(2 if succeeds else 1, len(commands))
                if succeeds:
                    doc = json.loads((paths["output"] / "rack-links-splitter-profiles.json").read_text())
                    self.assertEqual([], doc["profiles"])
                    self.assertIn('"leaf01":"swp1" -- "leaf02":"swp1"',
                                  (paths["output"] / "rack-links-lldpq.dot").read_text())
                else:
                    self.assertEqual(1, result.exception.code)
                    self.assertFalse((paths["output"] / "rack-links-splitter-profiles.json").exists())

    def test_generate_configs_stops_before_dhcp_generator_and_publication(self):
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            paths = DIRECT.prepare_runtime(root, DIRECT.make_legacy_column_workbook)
            exact_dot, exact_intent, unrelated = DIRECT.stale_artifacts(paths)
            commands: list[list[str]] = []
            resolver = mock.Mock(
                side_effect=DIRECT.LegacyReached("malformed workbook escaped"),
            )

            def run_real_p2p(command, *, cwd, **_kwargs):
                commands.append(list(command))
                self.assertEqual(paths["p2p_dir"], Path(cwd))
                caught, stdout, stderr = DIRECT.invoke_main(
                    paths, list(command[2:]), load_inventory=resolver,
                )
                if caught is not None:
                    raise caught
                self.fail("malformed workbook unexpectedly completed: " + stdout + stderr)

            with mock.patch.object(LOAD, "ZTP_DIR", root / "ztp"), \
                    mock.patch.object(LOAD, "run", side_effect=run_real_p2p):
                caught = None
                try:
                    LOAD.generate_configs(
                        frozenset({"eth"}), install_dhcp=False,
                        eth_version="5.16.4",
                    )
                except BaseException as exc:
                    caught = exc

            self.assertIsInstance(caught, SystemExit)
            self.assertEqual(1, caught.code)
            self.assertEqual(1, len(commands), commands)
            self.assertEqual("b-xlsx_to_dot.py", commands[0][1])
            self.assertFalse(any("c1-generate_dhcp.py" in command for command in commands))
            self.assertFalse(any("90-c2-generate_configs.py" in command for command in commands))
            self.assertFalse(any("d-hostname2mac.py" in command for command in commands))
            resolver.assert_not_called()
            self.assertFalse(exact_dot.exists())
            self.assertFalse(exact_intent.exists())
            self.assertTrue(all(path.is_file() for path in unrelated), unrelated)


if __name__ == "__main__":
    unittest.main()
